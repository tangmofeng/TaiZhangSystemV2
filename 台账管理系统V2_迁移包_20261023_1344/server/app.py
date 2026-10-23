# -*- coding: utf-8 -*-
"""
信息中心设备及备件台账管理系统 V2 - Flask 后端
模块：打印机台账/耗材库存与发放/费用统计/备品备件/办公电脑/电脑优化/
      AI 工作台/系统管理（账号权限、操作日志、自动备份）
V2 增强：随机持久化密钥、Session Cookie 加固、CSRF 同源校验、
        统一错误处理、访问日志、健康检查、登录限流
"""
import os
import io
import json
import sys
import glob
import time
import re
import sqlite3
import datetime
import threading
import urllib.request
import secrets
from urllib.parse import urlparse

from flask import Flask, request, session, jsonify, send_from_directory, send_file, g
from werkzeug.security import generate_password_hash, check_password_hash
from werkzeug.utils import secure_filename

# pythonw 无控制台环境下 stdout/stderr 可能为 None 或已关闭，需兜底，否则启动即崩
if sys.stdout is None or getattr(sys.stdout, 'closed', False):
    sys.stdout = open(os.devnull, 'w', encoding='utf-8')
if sys.stderr is None or getattr(sys.stderr, 'closed', False):
    sys.stderr = open(os.devnull, 'w', encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PUBLIC_DIR = os.path.join(BASE_DIR, 'public')
DATA_DIR = os.path.join(BASE_DIR, 'data')
DB_PATH = os.path.join(DATA_DIR, 'system.db')
BACKUP_DIR = os.path.join(DATA_DIR, 'backup')
ACCESS_LOG = os.path.join(DATA_DIR, 'access.log')
BACKUP_KEEP = 14

app = Flask(__name__, static_folder=None)


def _migrate_notes():
    """notes 表幂等迁移：回收站 deleted_at 列"""
    try:
        conn = sqlite3.connect(DB_PATH)
        cols = [r[1] for r in conn.execute("PRAGMA table_info(notes)").fetchall()]
        if 'deleted_at' not in cols:
            conn.execute("ALTER TABLE notes ADD COLUMN deleted_at TEXT DEFAULT ''")
            conn.commit()
        conn.close()
    except Exception:
        pass
_migrate_notes()


def _load_secret():
    """随机生成并持久化会话密钥（首次启动生成，之后复用，避免硬编码）"""
    key_file = os.path.join(DATA_DIR, '.secret_key')
    try:
        if os.path.exists(key_file):
            k = open(key_file, encoding='utf-8').read().strip()
            if k:
                return k
    except Exception:
        pass
    k = secrets.token_hex(32)
    try:
        os.makedirs(DATA_DIR, exist_ok=True)
        with open(key_file, 'w', encoding='utf-8') as f:
            f.write(k)
    except Exception:
        pass
    return k


app.secret_key = os.environ.get('TZ_SECRET') or _load_secret()
app.config['MAX_CONTENT_LENGTH'] = 16 * 1024 * 1024
# Session Cookie 加固
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE='Lax',
    SESSION_COOKIE_SECURE=False,  # 局域网 HTTP 场景保持 False；若将来上 HTTPS 改为 True
    PERMANENT_SESSION_LIFETIME=datetime.timedelta(hours=12),
)


@app.after_request
def add_security_headers(resp):
    resp.headers.setdefault('X-Frame-Options', 'SAMEORIGIN')
    resp.headers.setdefault('X-Content-Type-Options', 'nosniff')
    resp.headers.setdefault('Referrer-Policy', 'strict-origin-when-cross-origin')
    resp.headers.setdefault('Content-Security-Policy',
        "default-src 'self'; script-src 'self' 'unsafe-inline'; style-src 'self' 'unsafe-inline'; "
        "img-src 'self' data: blob:; font-src 'self' data:; connect-src 'self'")
    # 页面与前端脚本/样式禁用缓存：避免浏览器沿用旧版代码导致功能不生效
    if request.path == '/' or request.path.endswith('.html') or request.path.endswith('.js') or request.path.endswith('.css'):
        resp.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0'
        resp.headers['Pragma'] = 'no-cache'
        resp.headers['Expires'] = '0'
    return resp


@app.before_request
def csrf_protect():
    """CSRF 同源校验：非 GET 的 /api/ 请求校验 Origin/Referer 与 Host 一致"""
    if request.method == 'GET' or not request.path.startswith('/api/'):
        return
    origin = request.headers.get('Origin') or request.headers.get('Referer') or ''
    if origin:
        try:
            o = urlparse(origin)
            if o.netloc and o.netloc != request.host:
                return jsonify({'ok': False, 'error': '请求来源不合法'}), 403
        except Exception:
            return jsonify({'ok': False, 'error': '请求来源不合法'}), 403


@app.before_request
def req_timer():
    g._start = time.time()


@app.after_request
def access_log(resp):
    """访问日志：时间 方法 路径 状态 耗时 来源IP"""
    try:
        dur = int((time.time() - getattr(g, '_start', time.time())) * 1000)
        line = '%s %s %s %s %dms %s' % (
            now_str(), request.method, request.path, resp.status_code, dur, request.remote_addr or '-')
        with open(ACCESS_LOG, 'a', encoding='utf-8') as f:
            f.write(line + '\n')
    except Exception:
        pass
    return resp


@app.errorhandler(404)
def e404(e):
    return jsonify({'ok': False, 'error': '接口不存在'}), 404


@app.errorhandler(405)
def e405(e):
    return jsonify({'ok': False, 'error': '请求方法不允许'}), 405


@app.errorhandler(413)
def e413(e):
    return jsonify({'ok': False, 'error': '上传内容超过限制'}), 413


@app.errorhandler(Exception)
def e500(e):
    try:
        with open(os.path.join(DATA_DIR, 'error.log'), 'a', encoding='utf-8') as f:
            f.write('%s %s %s\n%s\n' % (now_str(), request.method, request.path, repr(e)))
    except Exception:
        pass
    return jsonify({'ok': False, 'error': '服务器内部错误'}), 500


DEFAULT_ADMIN = ('admin', 'admin123')


def get_conn():
    conn = sqlite3.connect(DB_PATH, timeout=30)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA busy_timeout = 5000")
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def now_str():
    return datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')


def num(v, default=0):
    try:
        if v is None or str(v).strip() == '':
            return default
        return float(v)
    except (ValueError, TypeError):
        return default


def log(user, action, detail=''):
    try:
        conn = get_conn()
        conn.execute("INSERT INTO sys_logs(time,user,action,detail) VALUES(?,?,?,?)",
                     (now_str(), user, action, detail[:500]))
        conn.commit()
        conn.close()
    except Exception:
        pass


# ---------------- 认证 ----------------
import threading as _threading
_login_fail = {}
_login_lock = _threading.Lock()


def _login_throttle(ip):
    """返回 (是否放行, 剩余锁定秒数)"""
    now = time.time()
    with _login_lock:
        ent = _login_fail.get(ip)
        if ent and ent.get('lock_until', 0) > now:
            return False, int(ent['lock_until'] - now)
        if not ent or now - ent.get('ts', 0) > 300:
            ent = {'fails': 0, 'ts': now, 'lock_until': 0}
            _login_fail[ip] = ent
        return True, 0


def _login_fail_record(ip):
    now = time.time()
    with _login_lock:
        ent = _login_fail.get(ip) or {'fails': 0, 'ts': now, 'lock_until': 0}
        if now - ent.get('ts', 0) > 300:
            ent = {'fails': 0, 'ts': now, 'lock_until': 0}
        ent['fails'] += 1
        ent['ts'] = now
        if ent['fails'] >= 5:
            ent['lock_until'] = now + 30
        _login_fail[ip] = ent


def _login_clear(ip):
    with _login_lock:
        _login_fail.pop(ip, None)


def require_login(fn):
    def wrapper(*args, **kwargs):
        if not session.get('uid'):
            return jsonify({'ok': False, 'error': '未登录'}), 401
        return fn(*args, **kwargs)
    wrapper.__name__ = fn.__name__
    return wrapper


# ---------------- 权限控制（账号权限分配管理） ----------------
PERM_MODULES = ['printers', 'computers', 'consumables', 'stats', 'spares', 'pcopt', 'notes', 'ai', 'system']
PERM_MAP = [
    ('/api/printers', 'printers'), ('/api/repairs', 'printers'),
    ('/api/computers', 'computers'), ('/api/computer-repairs', 'computers'),
    ('/api/pc-upgrades', 'pcopt'), ('/api/pc-parts', 'pcopt'), ('/api/pc-flows', 'pcopt'), ('/api/pc-adjusts', 'pcopt'),
    ('/api/consumables', 'consumables'), ('/api/maintenance', 'consumables'),
    ('/api/issues', 'consumables'),
    ('/api/statistics', 'stats'), ('/api/budgets', 'stats'),
    ('/api/items', 'spares'), ('/api/stocks', 'spares'), ('/api/flows', 'spares'),
    ('/api/spare/adjusts', 'spares'), ('/api/stocktakes', 'spares'),
    ('/api/ai', 'ai'),
    ('/api/logs', 'system'), ('/api/backup', 'system'), ('/api/backups', 'system'),
]


def get_perms(uid=None):
    """取用户权限字典；admin 返回 {'*': 'edit'}；空/旧账号默认全部只读"""
    if uid is None:
        uid = session.get('uid')
    if not uid:
        return {}
    conn = get_conn()
    row = conn.execute("SELECT role, perms FROM users WHERE id=?", (uid,)).fetchone()
    conn.close()
    if not row:
        return {}
    if row['role'] == 'admin':
        return {'*': 'edit'}
    try:
        p = json.loads(row['perms']) if row['perms'] else {}
    except Exception:
        p = {}
    if not p:
        return {m: 'view' for _, m in PERM_MAP}
    return p


def _ids_arg():
    """解析 ?ids=1,2,3 为 int 列表（用于导出选中数据）"""
    try:
        return [int(x) for x in request.args.get('ids', '').split(',') if x.strip().isdigit()]
    except Exception:
        return []


def has_perm(mod, action='view'):
    """判断当前用户对模块是否有 view/edit 权限"""
    p = get_perms()
    if '*' in p:
        return True
    lv = p.get(mod, '')
    if lv == 'edit':
        return True
    if lv == 'view':
        return action == 'view'
    return False


@app.before_request
def authz_check():
    p = request.path
    # 健康检查免登录（守护进程/监控探活）
    if p == '/api/health':
        return
    if p == '/api/login' or p == '/api/logout' or not p.startswith('/api/'):
        return
    if not session.get('uid'):
        return jsonify({'ok': False, 'error': '未登录'}), 401
    if p == '/api/me' or p == '/api/dashboard' or p == '/api/password' or p == '/api/health':
        return
    # 用户/账号管理仅管理员可用
    if p.startswith('/api/users'):
        if session.get('role') != 'admin':
            return jsonify({'ok': False, 'error': '仅管理员可管理账号'}), 403
        return
    mod = None
    for prefix, m in PERM_MAP:
        if p.startswith(prefix):
            mod = m
            break
    if mod is None:
        return
    if not has_perm(mod, 'view'):
        return jsonify({'ok': False, 'error': '无权访问该模块'}), 403
    if request.method != 'GET' and not has_perm(mod, 'edit'):
        return jsonify({'ok': False, 'error': '该账号为只读权限，无法执行此操作'}), 403


@app.route('/api/health', methods=['GET'])
def health():
    """健康检查：供守护进程/监控使用（无需登录）"""
    ok = os.path.exists(DB_PATH)
    return jsonify({'ok': ok, 'status': 'running', 'db': os.path.basename(DB_PATH)})


@app.route('/api/version', methods=['GET'])
@require_login
def version():
    return jsonify({'ok': True, 'version': '2.0.0', 'name': '信息中心设备及备件台账管理系统'})


@app.route('/api/login', methods=['POST'])
def login():
    ip = request.remote_addr or 'unknown'
    allow, wait = _login_throttle(ip)
    if not allow:
        return jsonify({'ok': False, 'error': '登录尝试过于频繁，请 %d 秒后重试' % wait}), 429
    data = request.get_json(force=True, silent=True) or {}
    username = str(data.get('username', '')).strip()
    password = str(data.get('password', ''))
    conn = get_conn()
    row = conn.execute("SELECT * FROM users WHERE username=? AND active=1", (username,)).fetchone()
    conn.close()
    if row and check_password_hash(row['password_hash'], password):
        session['uid'] = row['id']
        session['username'] = row['username']
        session['real_name'] = row['real_name'] or row['username']
        session['role'] = row['role']
        log(row['username'], '登录')
        _login_clear(ip)
        return jsonify({'ok': True, 'user': {
            'id': row['id'], 'username': row['username'],
            'real_name': row['real_name'], 'role': row['role'],
            'perms': get_perms(row['id'])}})
    _login_fail_record(ip)
    return jsonify({'ok': False, 'error': '用户名或密码错误'})


@app.route('/api/logout', methods=['POST'])
def logout():
    if session.get('username'):
        log(session['username'], '退出登录')
    session.clear()
    return jsonify({'ok': True})


@app.route('/api/me', methods=['GET'])
@require_login
def me():
    return jsonify({'ok': True, 'user': {
        'id': session['uid'], 'username': session['username'],
        'real_name': session['real_name'], 'role': session['role'],
        'perms': get_perms()}})


# ---------------- 用户级个性化设置（列顺序/列显隐等，按账号隔离） ----------------
@app.route('/api/user-settings', methods=['GET'])
@require_login
def user_settings_get():
    conn = get_conn()
    rows = conn.execute("SELECT skey, sval FROM user_settings WHERE username=?", (session.get('username'),)).fetchall()
    conn.close()
    settings = {}
    for r in rows:
        try:
            settings[r['skey']] = json.loads(r['sval']) if r['sval'] else None
        except Exception:
            settings[r['skey']] = None
    return jsonify({'ok': True, 'settings': settings})


@app.route('/api/user-settings', methods=['PUT'])
@require_login
def user_settings_put():
    d = request.get_json(silent=True) or {}
    skey = str(d.get('skey') or '').strip()
    if not skey or len(skey) > 120:
        return jsonify({'ok': False, 'error': '设置键无效'}), 400
    sval = d.get('sval')
    username = session.get('username') or ''
    conn = get_conn()
    try:
        conn.execute(
            "INSERT INTO user_settings(username, skey, sval, updated_at) VALUES(?,?,?,?) "
            "ON CONFLICT(username, skey) DO UPDATE SET sval=excluded.sval, updated_at=excluded.updated_at",
            (username, skey, json.dumps(sval, ensure_ascii=False), datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')))
        conn.commit()
        return jsonify({'ok': True})
    except Exception as e:
        return jsonify({'ok': False, 'error': '保存失败：' + str(e)}), 500
    finally:
        conn.close()


@app.route('/api/user-settings', methods=['DELETE'])
@require_login
def user_settings_delete():
    skey = (request.args.get('skey') or '').strip()
    if not skey:
        return jsonify({'ok': False, 'error': '缺少 skey'}), 400
    conn = get_conn()
    conn.execute("DELETE FROM user_settings WHERE username=? AND skey=?", (session.get('username') or '', skey))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})


# ---------------- 全局列配置（列顺序/列显隐）：admin 保存后所有用户同步一致 ----------------
@app.route('/api/col-settings', methods=['GET'])
@require_login
def col_settings_get():
    conn = get_conn()
    rows = conn.execute("SELECT id, value FROM col_settings").fetchall()
    conn.close()
    settings = {}
    for r in rows:
        try:
            settings[r['id']] = json.loads(r['value']) if r['value'] else None
        except Exception:
            settings[r['id']] = None
    return jsonify({'ok': True, 'settings': settings})


@app.route('/api/col-settings', methods=['PUT'])
@require_login
def col_settings_put():
    # 全局列配置仅 admin 可写，普通用户只读跟随
    if session.get('role') != 'admin':
        return jsonify({'ok': False, 'error': '列配置由管理员统一管理，普通用户只读'}), 403
    d = request.get_json(silent=True) or {}
    cid = str(d.get('id') or '').strip()
    if not cid or len(cid) > 120:
        return jsonify({'ok': False, 'error': '配置键无效'}), 400
    val = d.get('value')
    conn = get_conn()
    try:
        conn.execute(
            "INSERT INTO col_settings(id, value, updated_at) VALUES(?,?,?) "
            "ON CONFLICT(id) DO UPDATE SET value=excluded.value, updated_at=excluded.updated_at",
            (cid, json.dumps(val, ensure_ascii=False), datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')))
        conn.commit()
        return jsonify({'ok': True})
    except Exception as e:
        return jsonify({'ok': False, 'error': '保存失败：' + str(e)}), 500
    finally:
        conn.close()


@app.route('/api/col-settings/<path:cid>', methods=['DELETE'])
@require_login
def col_settings_delete(cid):
    # 全局列配置仅 admin 可写，普通用户只读跟随
    if session.get('role') != 'admin':
        return jsonify({'ok': False, 'error': '列配置由管理员统一管理，普通用户只读'}), 403
    conn = get_conn()
    conn.execute("DELETE FROM col_settings WHERE id=?", (cid,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})


# ---------------- 仪表盘 ----------------
@app.route('/api/dashboard', methods=['GET'])
@require_login
def dashboard():
    conn = get_conn()
    printers_total = conn.execute("SELECT COUNT(*) c FROM printers").fetchone()['c']
    printers_active = conn.execute("SELECT COUNT(*) c FROM printers WHERE status='在用'").fetchone()['c']
    computers_total = conn.execute("SELECT COUNT(*) c FROM computers").fetchone()['c']
    issues_month = conn.execute(
        "SELECT COUNT(*) c, COALESCE(SUM(qty),0) q FROM consumable_issues "
        "WHERE issue_time >= ?", (datetime.date.today().strftime('%Y-%m-01'),)).fetchone()
    repairs_pending = conn.execute(
        "SELECT COUNT(*) c FROM printer_repairs WHERE status IN ('待维修','维修中')").fetchone()['c']
    # 耗材低库存/零库存预警
    low_consumables = conn.execute(
        "SELECT * FROM consumables WHERE stock <= low_threshold AND low_threshold > 0 "
        "OR stock <= 0 ORDER BY stock ASC LIMIT 20").fetchall()
    # 备件零库存/低库存
    low_stocks = conn.execute(
        "SELECT s.*, i.brand, i.name, i.spec, i.unit FROM spare_stocks s "
        "LEFT JOIN spare_items i ON i.id = s.item_id "
        "WHERE (s.qty <= s.low_threshold AND s.low_threshold > 0) OR s.qty <= 0 "
        "ORDER BY s.qty ASC LIMIT 20").fetchall()
    recent_issues = conn.execute(
        "SELECT * FROM consumable_issues ORDER BY issue_time DESC, id DESC LIMIT 6").fetchall()
    recent_flows = conn.execute(
        "SELECT * FROM spare_flows ORDER BY record_time DESC, id DESC LIMIT 6").fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': {
        'printers_total': printers_total,
        'printers_active': printers_active,
        'computers_total': computers_total,
        'issues_month': dict(issues_month) if issues_month else {'c': 0, 'q': 0},
        'repairs_pending': repairs_pending,
        'low_consumables': [dict(r) for r in low_consumables],
        'low_stocks': [dict(r) for r in low_stocks],
        'recent_issues': [dict(r) for r in recent_issues],
        'recent_flows': [dict(r) for r in recent_flows],
    }})


# ---------------- 打印机台账 ----------------
@app.route('/api/printers', methods=['GET'])
@require_login
def list_printers():
    q = request.args.get('q', '').strip()
    department = request.args.get('department', '').strip()
    status = request.args.get('status', '').strip()
    sql = "SELECT * FROM printers WHERE 1=1"
    args = []
    if q:
        sql += " AND (printer_no LIKE ? OR department LIKE ? OR model LIKE ? OR location LIKE ? OR ip LIKE ?)"
        like = f"%{q}%"
        args += [like, like, like, like, like]
    if department:
        sql += " AND department=?"
        args.append(department)
    if status:
        sql += " AND status=?"
        args.append(status)
    sql += " ORDER BY CAST(printer_no AS INTEGER) ASC"
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows]})


@app.route('/api/printers/<int:pid>', methods=['GET'])
@require_login
def get_printer(pid):
    conn = get_conn()
    p = conn.execute("SELECT * FROM printers WHERE id=?", (pid,)).fetchone()
    repairs = conn.execute("SELECT * FROM printer_repairs WHERE printer_id=? ORDER BY report_time DESC, id DESC", (pid,)).fetchall()
    conn.close()
    if not p:
        return jsonify({'ok': False, 'error': '打印机不存在'}), 404
    return jsonify({'ok': True, 'data': dict(p), 'repairs': [dict(r) for r in repairs]})


@app.route('/api/printers', methods=['POST'])
@require_login
def add_printer():
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    cur = conn.execute(
        "INSERT INTO printers(printer_no,department,contact,location,model,ip,remark,status,created_at,updated_at) "
        "VALUES(?,?,?,?,?,?,?,?,?,?)",
        (str(d.get('printer_no', '')).strip(), str(d.get('department', '')).strip(),
         str(d.get('contact', '')).strip(), str(d.get('location', '')).strip(),
         str(d.get('model', '')).strip(), str(d.get('ip', '')).strip(),
         str(d.get('remark', '')).strip(), d.get('status') or '在用', now_str(), now_str()))
    conn.commit()
    conn.close()
    log(session['username'], '新增打印机', str(cur.lastrowid))
    return jsonify({'ok': True, 'id': cur.lastrowid})


@app.route('/api/printers/<int:pid>', methods=['PUT'])
@require_login
def update_printer(pid):
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    if not conn.execute("SELECT id FROM printers WHERE id=?", (pid,)).fetchone():
        conn.close()
        return jsonify({'ok': False, 'error': '打印机不存在'}), 404
    conn.execute(
        "UPDATE printers SET printer_no=?,department=?,contact=?,location=?,model=?,ip=?,remark=?,status=?,updated_at=? WHERE id=?",
        (str(d.get('printer_no', '')).strip(), str(d.get('department', '')).strip(),
         str(d.get('contact', '')).strip(), str(d.get('location', '')).strip(),
         str(d.get('model', '')).strip(), str(d.get('ip', '')).strip(),
         str(d.get('remark', '')).strip(), d.get('status') or '在用', now_str(), pid))
    conn.commit()
    conn.close()
    log(session['username'], '修改打印机', str(pid))
    return jsonify({'ok': True})


@app.route('/api/printers/<int:pid>', methods=['DELETE'])
@require_login
def delete_printer(pid):
    conn = get_conn()
    conn.execute("DELETE FROM printers WHERE id=?", (pid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除打印机', str(pid))
    return jsonify({'ok': True})


# ---------------- 打印机 模板导入/导出 ----------------
@app.route('/api/printers/template', methods=['GET'])
@require_login
def printer_template():
    """下载打印机导入模板（xlsx）"""
    import openpyxl, io as _io
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '打印机导入模板'
    headers = ['打印机编号*', '使用部门', '联系人及电话', '部署位置', '型号', 'IP地址', '状态', '备注']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='D9E2F3')
        cell.alignment = Alignment(horizontal='center')
    ws.append(['303', '检验计量中心', '张三 13800000000', '二号岛辅楼101', 'HP LaserJet 1020', '10.1.1.1', '在用', ''])
    widths = [14, 16, 22, 26, 22, 16, 10, 18]
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='打印机导入模板.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/api/printers/import', methods=['POST'])
@require_login
def import_printers():
    """批量导入打印机：逐行校验（编号必填/状态合法/编号查重），全部通过后事务插入"""
    f = request.files.get('file')
    if not f or not f.filename:
        return jsonify({'ok': False, 'error': '请选择要导入的 Excel 文件'})
    if not (f.filename.endswith('.xlsx') or f.filename.endswith('.xls')):
        return jsonify({'ok': False, 'error': '仅支持 .xlsx / .xls 文件'})
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(f.read()), data_only=True)
    except Exception as e:
        return jsonify({'ok': False, 'error': '文件解析失败：%s' % e})
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    conn = get_conn()
    prepared = []
    errors = []
    seen_no = set()
    for idx, row in enumerate(rows, start=2):
        if row is None or all(v is None or str(v).strip() == '' for v in row):
            continue
        def gv(i):
            return str(row[i] or '').strip() if len(row) > i and row[i] is not None else ''
        printer_no = gv(0)
        if not printer_no:
            errors.append('第%d行：打印机编号必填' % idx)
            continue
        if printer_no in seen_no:
            errors.append('第%d行：打印机编号「%s」在文件中重复' % (idx, printer_no))
            continue
        seen_no.add(printer_no)
        status = gv(6) or '在用'
        if status not in ('在用', '停用', '报废'):
            errors.append('第%d行：状态必须为 在用/停用/报废' % idx)
            continue
        dup = conn.execute("SELECT id FROM printers WHERE printer_no=?", (printer_no,)).fetchone()
        if dup:
            errors.append('第%d行：打印机编号「%s」已存在（库内），请改用编辑' % (idx, printer_no))
            continue
        prepared.append((printer_no, gv(1), gv(2), gv(3), gv(4), gv(5), gv(7), status))
    if errors:
        conn.close()
        return jsonify({'ok': False, 'error': '存在 %d 行校验失败，已整体取消导入' % len(errors),
                        'errors': errors[:20]})
    if not prepared:
        conn.close()
        return jsonify({'ok': False, 'error': '模板中没有可导入的数据行'})
    try:
        now = now_str()
        for p in prepared:
            conn.execute(
                "INSERT INTO printers(printer_no,department,contact,location,model,ip,remark,status,created_at,updated_at) "
                "VALUES(?,?,?,?,?,?,?,?,?,?)",
                (p[0], p[1], p[2], p[3], p[4], p[5], p[6], p[7], now, now))
        conn.commit()
        conn.close()
        log(session['username'], '模板导入打印机', '共%d台' % len(prepared))
        return jsonify({'ok': True, 'added': len(prepared)})
    except Exception as e:
        conn.rollback()
        conn.close()
        return jsonify({'ok': False, 'error': '导入执行失败：%s' % e})


@app.route('/api/printers/export', methods=['GET'])
@require_login
def export_printers():
    """导出全部/选中/按筛选条件打印机（xlsx）"""
    q = request.args.get('q', '').strip()
    department = request.args.get('department', '').strip()
    status = request.args.get('status', '').strip()
    conn = get_conn()
    ids = _ids_arg()
    if ids:
        ph = ','.join('?' * len(ids))
        sql = "SELECT * FROM printers WHERE id IN (%s)" % ph
        args = list(ids)
    else:
        sql = "SELECT * FROM printers WHERE 1=1"
        args = []
    if q:
        sql += " AND (printer_no LIKE ? OR department LIKE ? OR model LIKE ? OR location LIKE ? OR ip LIKE ?)"
        like = f"%{q}%"
        args += [like, like, like, like, like]
    if department:
        sql += " AND department=?"
        args.append(department)
    if status:
        sql += " AND status=?"
        args.append(status)
    sql += " ORDER BY CAST(printer_no AS INTEGER) ASC"
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    import openpyxl, io as _io
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '打印机台账'
    headers = ['打印机编号', '使用部门', '联系人及电话', '部署位置', '型号', 'IP地址', '状态', '备注']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='D9E2F3')
        cell.alignment = Alignment(horizontal='center')
    for r in rows:
        ws.append([r['printer_no'], r['department'], r['contact'], r['location'], r['model'], r['ip'], r['status'], r['remark']])
    widths = [14, 16, 22, 26, 22, 16, 10, 18]
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='打印机台账导出.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ---------------- 耗材库存 ----------------
def _con_recalc(conn, cid):
    """重算种类合计库存（权威=consumable_stocks，consumables.stock 为冗余合计）"""
    s = conn.execute("SELECT COALESCE(SUM(stock),0) FROM consumable_stocks WHERE consumable_id=?", (cid,)).fetchone()[0]
    conn.execute("UPDATE consumables SET stock=? WHERE id=?", (s, cid))

def _con_stock_row(conn, cid, region):
    """取某种类在指定区域的库存行；不存在返回 None"""
    return conn.execute("SELECT * FROM consumable_stocks WHERE consumable_id=? AND region=?",
                        (cid, region or '办公楼')).fetchone()

def _con_refund_row(conn, cid, region):
    """回补目标行：region 有效→该区域行；region 空→唯一行或办公楼优先"""
    if region in ('中控', '办公楼'):
        return _con_stock_row(conn, cid, region)
    rows = conn.execute(
        "SELECT * FROM consumable_stocks WHERE consumable_id=? "
        "ORDER BY CASE region WHEN '办公楼' THEN 0 ELSE 1 END, id", (cid,)).fetchall()
    return rows[0] if rows else None

@app.route('/api/consumables', methods=['GET'])
@require_login
def list_consumables():
    q = request.args.get('q', '').strip()
    region = request.args.get('region', '').strip()
    sql = ("SELECT c.*, s.id AS stock_id, s.region, s.storage, s.stock AS region_stock, "
           "(SELECT COALESCE(SUM(stock),0) FROM consumable_stocks WHERE consumable_id=c.id) AS total_stock "
           "FROM consumables c LEFT JOIN consumable_stocks s ON s.consumable_id=c.id WHERE 1=1")
    args = []
    if q:
        sql += " AND (c.model LIKE ? OR c.category LIKE ? OR c.material_code LIKE ?)"
        like = f"%{q}%"
        args += [like, like, like]
    if region:
        sql += " AND s.region = ?"
        args.append(region)
    sql += " ORDER BY c.model ASC, c.category ASC, s.region ASC"
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows]})


@app.route('/api/consumables', methods=['POST'])
@require_login
def add_consumable():
    """新增耗材种类（种类唯一：同 型号+品类 已存在则拒绝）+ 可选初始区域库存"""
    d = request.get_json(force=True, silent=True) or {}
    model = str(d.get('model', '')).strip()
    category = str(d.get('category', '')).strip()
    if not model or not category:
        return jsonify({'ok': False, 'error': '打印机型号与耗材品类必填'})
    conn = get_conn()
    dup = conn.execute("SELECT id FROM consumables WHERE model=? AND category=?",
                       (model, category)).fetchone()
    if dup:
        conn.close()
        return jsonify({'ok': False, 'error': '该耗材种类「%s · %s」已存在（种类唯一，勿重复建档），如需补货请直接调整/入库' % (model, category)})
    region = str(d.get('region', '')).strip()
    region = region if region in ('中控', '办公楼') else '办公楼'
    storage = str(d.get('storage', '')).strip()
    stock = max(0, num(d.get('stock')) if d.get('stock') is not None else 0)
    cur = conn.execute(
        "INSERT INTO consumables(model,category,material_code,low_threshold,remark,price,printer_models) VALUES(?,?,?,?,?,?,?)",
        (model, category, str(d.get('material_code', '')).strip(),
         num(d.get('low_threshold')), str(d.get('remark', '')).strip(), num(d.get('price')),
         str(d.get('printer_models', '')).strip()))
    cid = cur.lastrowid
    conn.execute("INSERT INTO consumable_stocks(consumable_id, region, storage, stock) VALUES(?,?,?,?)",
                 (cid, region, storage, stock))
    _con_recalc(conn, cid)
    conn.commit()
    conn.close()
    log(session['username'], '新增耗材种类', f"id={cid} region={region} storage={storage} stock={stock}")
    return jsonify({'ok': True, 'id': cid, 'stock_id': cur.lastrowid if False else None})


@app.route('/api/consumables/<int:cid>', methods=['PUT'])
@require_login
def update_consumable(cid):
    """编辑耗材种类信息（型号/品类/编码/单价/阈值/备注/兼容机型）；库存与区域走调整/入库"""
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    model = str(d.get('model', '')).strip()
    category = str(d.get('category', '')).strip()
    if not model or not category:
        conn.close()
        return jsonify({'ok': False, 'error': '打印机型号与耗材品类必填'})
    dup = conn.execute("SELECT id FROM consumables WHERE model=? AND category=? AND id<>?",
                       (model, category, cid)).fetchone()
    if dup:
        conn.close()
        return jsonify({'ok': False, 'error': '已存在同名耗材种类「%s · %s」，请勿重复' % (model, category)})
    conn.execute(
        "UPDATE consumables SET model=?,category=?,material_code=?,low_threshold=?,remark=?,price=?,printer_models=? WHERE id=?",
        (model, category, str(d.get('material_code', '')).strip(),
         num(d.get('low_threshold')), str(d.get('remark', '')).strip(), num(d.get('price')),
         str(d.get('printer_models', '')).strip(), cid))
    conn.commit()
    conn.close()
    log(session['username'], '修改耗材', str(cid))
    return jsonify({'ok': True})


@app.route('/api/consumables/<int:cid>', methods=['DELETE'])
@require_login
def delete_consumable(cid):
    conn = get_conn()
    row = conn.execute("SELECT * FROM consumables WHERE id=?", (cid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '耗材不存在'}), 404
    conn.execute("DELETE FROM consumable_stocks WHERE consumable_id=?", (cid,))
    conn.execute("DELETE FROM consumables WHERE id=?", (cid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除耗材', str(cid))
    return jsonify({'ok': True})


@app.route('/api/consumables/<int:cid>/adjust', methods=['POST'])
@require_login
def adjust_consumable(cid):
    """手动调整库存（入库/盘盈盘亏等）：按 区域 调整 consumable_stocks 对应行，
    delta 可为正负，写入调整记录（含区域/库位）"""
    d = request.get_json(force=True, silent=True) or {}
    delta = num(d.get('delta'))
    if delta == 0:
        return jsonify({'ok': False, 'error': '调整数量不能为 0'})
    atype = str(d.get('type', '手工调整')).strip() or '手工调整'
    reason = str(d.get('reason', '')).strip()
    region = str(d.get('region', '')).strip()
    region = region if region in ('中控', '办公楼') else '办公楼'
    storage = str(d.get('storage', '')).strip()
    conn = get_conn()
    row = conn.execute("SELECT * FROM consumables WHERE id=?", (cid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '耗材不存在'}), 404
    srow = _con_stock_row(conn, cid, region)
    if not srow:
        conn.execute("INSERT INTO consumable_stocks(consumable_id, region, storage, stock) VALUES(?,?,?,0)",
                     (cid, region, storage))
        srow = _con_stock_row(conn, cid, region)
    before = num(srow['stock'])
    new_stock = before + delta
    if new_stock < 0:
        conn.close()
        return jsonify({'ok': False, 'error': '库存不足，%s区域当前库存 %s' % (region, before)})
    conn.execute("UPDATE consumable_stocks SET stock=?, storage=? WHERE id=?",
                 (new_stock, storage or srow['storage'] or '', srow['id']))
    _con_recalc(conn, cid)
    now = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    conn.execute(
        "INSERT INTO consumable_adjusts(consumable_id,model,category,type,delta,before_stock,after_stock,reason,operator,adjust_time,created_at,region,storage) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (cid, row['model'], row['category'], atype, delta, before, new_stock, reason,
         session['username'], now, now, region, storage or srow['storage'] or ''))
    adj_id = conn.execute("SELECT last_insert_rowid() AS i").fetchone()['i']
    conn.commit()
    conn.close()
    log(session['username'], '调整耗材库存', f"{cid} region={region} delta={delta} type={atype}")
    return jsonify({'ok': True, 'stock': new_stock, 'region': region, 'total_stock': num(row['stock']) + delta,
                    'adjust_id': adj_id})


@app.route('/api/consumables/<int:cid>/transfer', methods=['POST'])
@require_login
def transfer_consumable(cid):
    """移库：把某 区域+库位 库存行（from_stock_id）的 qty 转移到目标 区域+库位（upsert），
    源行扣减（扣至 0 自动删除该行），写两条调整流水（移库转出/移库转入）"""
    d = request.get_json(force=True, silent=True) or {}
    from_stock_id = d.get('from_stock_id')
    qty = num(d.get('qty'))
    to_region = str(d.get('to_region', '')).strip()
    to_region = to_region if to_region in ('中控', '办公楼') else '办公楼'
    to_storage = str(d.get('to_storage', '')).strip()
    if not from_stock_id:
        return jsonify({'ok': False, 'error': '请选择源库存行'})
    if not qty or qty <= 0:
        return jsonify({'ok': False, 'error': '转移数量应为正数'})
    if not to_storage:
        return jsonify({'ok': False, 'error': '目标库位必填'})
    conn = get_conn()
    row = conn.execute("SELECT * FROM consumables WHERE id=?", (cid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '耗材不存在'}), 404
    src = conn.execute("SELECT * FROM consumable_stocks WHERE id=?", (from_stock_id,)).fetchone()
    if not src or src['consumable_id'] != cid:
        conn.close()
        return jsonify({'ok': False, 'error': '源库存行不存在'}), 404
    src_stock = num(src['stock'])
    if src_stock < qty:
        conn.close()
        return jsonify({'ok': False, 'error': '%s %s库存不足（当前 %s）' % (src['region'], src['storage'] or '(未设库位)', src_stock)})
    if src['region'] == to_region and (src['storage'] or '') == to_storage:
        conn.close()
        return jsonify({'ok': False, 'error': '源与目标为同一 区域+库位，无需移库'})
    dst = conn.execute("SELECT * FROM consumable_stocks WHERE consumable_id=? AND region=? AND storage=?",
                       (cid, to_region, to_storage)).fetchone()
    dst_before = num(dst['stock']) if dst else 0
    dst_after = dst_before + qty
    if dst:
        conn.execute("UPDATE consumable_stocks SET stock=? WHERE id=?", (dst_after, dst['id']))
        dst_id = dst['id']
    else:
        cur = conn.execute("INSERT INTO consumable_stocks(consumable_id,region,storage,stock) VALUES(?,?,?,?)",
                           (cid, to_region, to_storage, qty))
        dst_id = cur.lastrowid
    new_src = src_stock - qty
    if new_src <= 0:
        conn.execute("DELETE FROM consumable_stocks WHERE id=?", (src['id'],))
    else:
        conn.execute("UPDATE consumable_stocks SET stock=? WHERE id=?", (new_src, src['id']))
    _con_recalc(conn, cid)
    now = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    src_label = src['region'] + (('/' + src['storage']) if src['storage'] else '(未设库位)')
    dst_label = to_region + (('/' + to_storage) if to_storage else '')
    conn.execute(
        "INSERT INTO consumable_adjusts(consumable_id,model,category,type,delta,before_stock,after_stock,reason,operator,adjust_time,created_at,region,storage) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (cid, row['model'], row['category'], '移库(转出)', -qty, src_stock, new_src, '移库 → ' + dst_label,
         session['username'], now, now, src['region'], src['storage'] or ''))
    conn.execute(
        "INSERT INTO consumable_adjusts(consumable_id,model,category,type,delta,before_stock,after_stock,reason,operator,adjust_time,created_at,region,storage) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (cid, row['model'], row['category'], '移库(转入)', qty, dst_before, dst_after, '自 ' + src_label + ' 移入',
         session['username'], now, now, to_region, to_storage))
    conn.commit()
    conn.close()
    log(session['username'], '耗材移库', f"{cid} {src_label} -{qty} -> {dst_label}")
    return jsonify({'ok': True, 'from_stock': new_src, 'to_stock': dst_after, 'total_stock': num(row['stock'])})


@app.route('/api/consumables/adjusts', methods=['GET'])
@require_login
def list_consumable_adjusts():
    """耗材调整记录查询：q(型号/品类)、type、operator、start/end(日期)、page/size"""
    q = request.args.get('q', '').strip()
    atype = request.args.get('type', '').strip()
    operator = request.args.get('operator', '').strip()
    start = request.args.get('start', '').strip()
    end = request.args.get('end', '').strip()
    try:
        page = max(1, int(request.args.get('page', 1)))
        size = max(1, min(100000, int(request.args.get('size', 20))))
    except Exception:
        page, size = 1, 20
    where = " WHERE 1=1"
    args = []
    if q:
        like = f"%{q}%"
        where += " AND (model LIKE ? OR category LIKE ?)"
        args += [like, like]
    if atype:
        where += " AND type=?"
        args.append(atype)
    if operator:
        where += " AND operator LIKE ?"
        args.append(f"%{operator}%")
    if start:
        where += " AND adjust_time >= ?"
        args.append(start + ' 00:00:00')
    if end:
        where += " AND adjust_time <= ?"
        args.append(end + ' 23:59:59')
    conn = get_conn()
    total = conn.execute("SELECT COUNT(*) c FROM consumable_adjusts" + where, args).fetchone()['c']
    rows = conn.execute(
        "SELECT * FROM consumable_adjusts" + where + " ORDER BY adjust_time DESC, id DESC LIMIT ? OFFSET ?",
        args + [size, (page - 1) * size]).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows],
                    'total': total, 'page': page, 'size': size})


@app.route('/api/consumables/adjusts/<int:aid>', methods=['DELETE'])
@require_login
def delete_consumable_adjust(aid):
    """删除调整记录（仅删记录，不影响当前库存；修正库存请用调整功能反向调整）"""
    conn = get_conn()
    row = conn.execute("SELECT * FROM consumable_adjusts WHERE id=?", (aid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '调整记录不存在'}), 404
    conn.execute("DELETE FROM consumable_adjusts WHERE id=?", (aid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除调整记录', f"{aid} {row['model']} {row['type']} delta={row['delta']}")
    return jsonify({'ok': True})


# ---------------- 耗材入库记录（采购入库/补录，自动加库存） ----------------
@app.route('/api/consumable-stock-in', methods=['GET'])
@require_login
def list_consumable_stock_in():
    q = request.args.get('q', '').strip()
    start = request.args.get('start', '').strip()
    end = request.args.get('end', '').strip()
    operator = request.args.get('operator', '').strip()
    try:
        page = max(1, int(request.args.get('page', 1)))
        size = max(1, min(100000, int(request.args.get('size', 20))))
    except Exception:
        page, size = 1, 20
    where = " WHERE 1=1"
    args = []
    if q:
        like = f"%{q}%"
        where += " AND (model LIKE ? OR category LIKE ? OR note LIKE ?)"
        args += [like, like, like]
    if operator:
        where += " AND operator LIKE ?"
        args.append(f"%{operator}%")
    if start:
        where += " AND stock_in_time >= ?"
        args.append(start + ' 00:00:00')
    if end:
        where += " AND stock_in_time <= ?"
        args.append(end + ' 23:59:59')
    conn = get_conn()
    total = conn.execute("SELECT COUNT(*) c FROM consumable_stock_in" + where, args).fetchone()['c']
    rows = conn.execute(
        "SELECT * FROM consumable_stock_in" + where + " ORDER BY stock_in_time DESC, id DESC LIMIT ? OFFSET ?",
        args + [size, (page - 1) * size]).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows], 'total': total, 'page': page, 'size': size})


@app.route('/api/consumable-stock-in', methods=['POST'])
@require_login
def create_consumable_stock_in():
    """耗材入库登记：按 区域+库位 累加库存并生成入库记录（记录带区域/库位）"""
    d = request.get_json(force=True, silent=True) or {}
    cid = int(d.get('consumable_id') or 0)
    qty = num(d.get('qty'))
    if cid <= 0:
        return jsonify({'ok': False, 'error': '请选择耗材'})
    if qty <= 0:
        return jsonify({'ok': False, 'error': '入库数量必须大于 0'})
    region = str(d.get('region', '')).strip()
    region = region if region in ('中控', '办公楼') else '办公楼'
    storage = str(d.get('storage', '')).strip()
    price = num(d.get('price'))
    operator = str(d.get('operator') or '').strip() or (session.get('real_name') or session['username'])
    note = str(d.get('note') or '').strip()
    sin_time = str(d.get('stock_in_time') or '').strip() or now_str()[:10]
    conn = get_conn()
    row = conn.execute("SELECT * FROM consumables WHERE id=?", (cid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '耗材不存在'}), 404
    srow = _con_stock_row(conn, cid, region)
    if not srow:
        conn.execute("INSERT INTO consumable_stocks(consumable_id, region, storage, stock) VALUES(?,?,?,0)",
                     (cid, region, storage))
        srow = _con_stock_row(conn, cid, region)
    conn.execute("UPDATE consumable_stocks SET stock=?, storage=? WHERE id=?",
                 (num(srow['stock']) + qty, storage or srow['storage'] or '', srow['id']))
    _con_recalc(conn, cid)
    if price > 0:
        conn.execute("UPDATE consumables SET price=? WHERE id=?", (price, cid))
    now = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    conn.execute(
        "INSERT INTO consumable_stock_in(stock_in_time,consumable_id,model,category,qty,price,amount,operator,note,created_at,region,storage) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
        (sin_time, cid, row['model'], row['category'], qty, price, round(qty * price, 2), operator, note, now,
         region, storage))
    conn.commit()
    conn.close()
    log(session['username'], '耗材入库', f"{row['model']} {row['category']} {region} +{qty} 经办人={operator}")
    return jsonify({'ok': True, 'stock': num(srow['stock']) + qty, 'region': region, 'total_stock': num(row['stock']) + qty})


@app.route('/api/consumable-stock-in/<int:sid>', methods=['DELETE'])
@require_login
def delete_consumable_stock_in(sid):
    """删除入库记录：按记录的区域回扣该区域库存"""
    conn = get_conn()
    row = conn.execute("SELECT * FROM consumable_stock_in WHERE id=?", (sid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '入库记录不存在'}), 404
    region = row['region'] if ('region' in row.keys() and row['region']) else '办公楼'
    srow = _con_stock_row(conn, row['consumable_id'], region)
    if srow and num(srow['stock']) - num(row['qty']) < 0:
        conn.close()
        return jsonify({'ok': False, 'error': '%s区域库存不足，无法删除该入库记录（库存 %s）' % (region, srow['stock'])})
    if srow:
        conn.execute("UPDATE consumable_stocks SET stock=stock-? WHERE id=?", (row['qty'], srow['id']))
        _con_recalc(conn, row['consumable_id'])
    conn.execute("DELETE FROM consumable_stock_in WHERE id=?", (sid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除耗材入库记录', f"{sid} {row['model']} {region} -{row['qty']}")
    return jsonify({'ok': True})


@app.route('/api/consumable-stock-in/export', methods=['GET'])
@require_login
def export_consumable_stock_in():
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    q = request.args.get('q', '').strip()
    start = request.args.get('start', '').strip()
    end = request.args.get('end', '').strip()
    where = " WHERE 1=1"
    args = []
    if q:
        like = f"%{q}%"
        where += " AND (model LIKE ? OR category LIKE ? OR note LIKE ?)"
        args += [like, like, like]
    if start:
        where += " AND stock_in_time >= ?"
        args.append(start + ' 00:00:00')
    if end:
        where += " AND stock_in_time <= ?"
        args.append(end + ' 23:59:59')
    conn = get_conn()
    rows = conn.execute("SELECT * FROM consumable_stock_in" + where + " ORDER BY stock_in_time DESC, id DESC", args).fetchall()
    conn.close()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '耗材入库记录'
    heads = ['入库时间', '区域', '库位', '打印机型号', '耗材品类', '数量', '单价(元)', '金额(元)', '经办人', '备注']
    ws.append(heads)
    for c in range(1, len(heads) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='E5E7EB')
        cell.alignment = Alignment(horizontal='center')
    for r in rows:
        ws.append([r['stock_in_time'],
                   r['region'] if ('region' in r.keys() and r['region']) else '',
                   r['storage'] if ('storage' in r.keys() and r['storage']) else '',
                   r['model'], r['category'], r['qty'], r['price'], r['amount'], r['operator'], r['note']])
    for i, w in enumerate([16, 8, 14, 22, 14, 8, 10, 12, 12, 24]):
        ws.column_dimensions[chr(65 + i)].width = w
    import io as _io
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='耗材入库记录.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ---------------- 耗材入库记录模板批量导入 ----------------
@app.route('/api/consumable-stock-in/template', methods=['GET'])
@require_login
def consumable_stock_in_template():
    """下载耗材入库记录导入模板（xlsx）：按种类预填现有型号+品类+单价，只需填写本次入库数量；
    区域（中控/办公楼）+库位可修改，同种类可在两个区域各入库一行（种类不重复）"""
    import openpyxl, io as _io
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    conn = get_conn()
    rows = conn.execute("SELECT model,category,price FROM consumables ORDER BY model,category").fetchall()
    conn.close()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '耗材入库导入'
    headers = ['入库时间', '区域', '库位', '打印机型号*', '耗材品类*', '本次入库数量*', '单价(元)', '经办人', '备注']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='D9E2F3')
        cell.alignment = Alignment(horizontal='center')
    today = datetime.datetime.now().strftime('%Y-%m-%d')
    if rows:
        for r in rows:
            ws.append([today, '办公楼', '', r['model'], r['category'], '', r['price'] or 0, '', ''])
        note = ('已预填全部 %d 种现有耗材：只需在「本次入库数量」列填写数量，导入后按 种类+区域 累加库存并生成入库记录；'
                '区域（中控/办公楼）与库位可修改，同一型号如需两个区域各入库，复制该行改「区域」即可；型号/品类请勿修改，未匹配的行会报错' % len(rows))
    else:
        ws.append([today, '办公楼', '', '惠普M254', '黑粉', '', 190, '', '示例：模板按现有库存自动生成，请先在耗材库存中添加耗材'])
        note = '当前暂无耗材库存数据，请在「耗材库存」中先添加耗材后再下载模板'
    ws.cell(ws.max_row + 2, 1, note).font = Font(color='888888', italic=True)
    for i, w in enumerate([12, 8, 14, 22, 14, 16, 10, 12, 24]):
        ws.column_dimensions[get_column_letter(i + 1)].width = w
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='耗材入库记录导入模板.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/api/consumable-stock-in/import', methods=['POST'])
@require_login
def import_consumable_stock_in():
    """批量导入耗材入库记录：按 型号+品类 匹配现有耗材（种类唯一），按 区域+库位 累加该区域库存并生成入库记录"""
    f = request.files.get('file')
    if not f or not f.filename:
        return jsonify({'ok': False, 'error': '请选择要导入的 Excel 文件'})
    if not (f.filename.endswith('.xlsx') or f.filename.endswith('.xls')):
        return jsonify({'ok': False, 'error': '仅支持 .xlsx / .xls 文件'})
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(f.read()), data_only=True)
    except Exception as e:
        return jsonify({'ok': False, 'error': '文件解析失败：%s' % e})
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    conn = get_conn()
    added = failed = 0
    errors = []

    def cell_num(v):
        if v is None:
            return None
        s = str(v).strip()
        if s == '':
            return None
        try:
            return float(s)
        except (ValueError, TypeError):
            return 'ERR'

    default_op = session.get('real_name') or session['username']
    today = datetime.datetime.now().strftime('%Y-%m-%d')
    for idx, row in enumerate(rows, start=2):
        if row is None or all(v is None or str(v).strip() == '' for v in row):
            continue
        region = str(row[1] or '').strip() if len(row) > 1 and row[1] is not None else ''
        storage = str(row[2] or '').strip() if len(row) > 2 and row[2] is not None else ''
        model = str(row[3] or '').strip() if len(row) > 3 else ''
        category = str(row[4] or '').strip() if len(row) > 4 else ''
        if not model or not category:
            failed += 1
            errors.append(f"第{idx}行：打印机型号与耗材品类必填")
            continue
        qty = cell_num(row[5]) if len(row) > 5 else None
        if qty == 'ERR':
            failed += 1
            errors.append(f"第{idx}行：本次入库数量不是数字")
            continue
        if qty is None or qty <= 0:
            failed += 1
            errors.append(f"第{idx}行：本次入库数量必须大于 0")
            continue
        price = cell_num(row[6]) if len(row) > 6 else None
        if price == 'ERR':
            failed += 1
            errors.append(f"第{idx}行：单价不是数字")
            continue
        sin_time = str(row[0] or '').strip() if len(row) > 0 else ''
        sin_time = sin_time[:10] if len(sin_time) >= 10 else today
        operator = str(row[7] or '').strip() if len(row) > 7 else ''
        note = str(row[8] or '').strip() if len(row) > 8 else ''
        region = region if region in ('中控', '办公楼') else '办公楼'
        con = conn.execute("SELECT * FROM consumables WHERE model=? AND category=?",
                           (model, category)).fetchone()
        if not con:
            failed += 1
            errors.append(f"第{idx}行：库存中不存在「{model} {category}」，请先在耗材库存中添加")
            continue
        try:
            srow = _con_stock_row(conn, con['id'], region)
            if not srow:
                conn.execute("INSERT INTO consumable_stocks(consumable_id, region, storage, stock) VALUES(?,?,?,0)",
                             (con['id'], region, storage))
                srow = _con_stock_row(conn, con['id'], region)
            conn.execute("UPDATE consumable_stocks SET stock=?, storage=? WHERE id=?",
                         (num(srow['stock']) + qty, storage or srow['storage'] or '', srow['id']))
            _con_recalc(conn, con['id'])
            eff_price = price if (price is not None and price > 0) else (num(con['price']) or 0)
            if price is not None and price > 0:
                conn.execute("UPDATE consumables SET price=? WHERE id=?", (price, con['id']))
            now = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            conn.execute(
                "INSERT INTO consumable_stock_in(stock_in_time,consumable_id,model,category,qty,price,amount,operator,note,created_at,region,storage) VALUES(?,?,?,?,?,?,?,?,?,?,?,?)",
                (sin_time, con['id'], model, category, qty, eff_price, round(qty * eff_price, 2),
                 operator or default_op, note, now, region, storage))
            added += 1
        except Exception as e:
            failed += 1
            errors.append(f"第{idx}行：{e}")
    conn.commit()
    conn.close()
    log(session['username'], '批量导入耗材入库记录', f"added={added} failed={failed}")
    return jsonify({'ok': True, 'added': added, 'failed': failed, 'errors': errors[:50]})


# ---------------- 耗材模板导入 ----------------
@app.route('/api/consumables/template', methods=['GET'])
@require_login
def consumable_template():
    """下载耗材批量导入模板（xlsx）：按种类预填各区域库存行，只需填写「当前库存」数量"""
    import openpyxl, io as _io
    from openpyxl.styles import Font, PatternFill, Alignment
    conn = get_conn()
    rows = conn.execute(
        "SELECT c.model,c.category,c.material_code,c.price,c.low_threshold,c.remark, s.region, s.storage "
        "FROM consumables c LEFT JOIN consumable_stocks s ON s.consumable_id=c.id "
        "WHERE s.id IS NOT NULL ORDER BY c.model, c.category, s.region"
    ).fetchall()
    conn.close()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '耗材库存盘点'
    headers = ['区域', '库位', '打印机型号*', '耗材品类*', '物料编码', '单价(元)', '当前库存(填数量)', '低库存预警阈值', '备注']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='D9E2F3')
        cell.alignment = Alignment(horizontal='center')
    if rows:
        for r in rows:
            ws.append([
                (r['region'] or '办公楼'), (r['storage'] or ''),
                r['model'], r['category'],
                (r['material_code'] or ''),
                (r['price'] or 0),
                '',  # 当前库存留空，由用户填写实际盘点数量
                (r['low_threshold'] or 0),
                (r['remark'] or ''),
            ])
        note = ('已预填全部 %d 个库存行（按 种类×区域，种类唯一不重复）：区域（中控/办公楼）与库位可修改；'
                '只需在「当前库存」列填写实际数量，留空的行不改变该区域库存；'
                '同一型号如需在中控、办公楼各备货，复制该行并改「区域」即可；如需新增耗材种类请在末尾追加新行' % len(rows))
    else:
        ws.append(['办公楼', '文件柜4-1', '惠普M254', '黑粉', 'S0440001', 190, 20, 5, '示例：请在末尾按同样格式追加耗材'])
        note = '当前暂无耗材数据，请按表头格式自行填写（区域可选 中控/办公楼）'
    ws.cell(ws.max_row + 2, 1, note).font = Font(color='888888', italic=True)
    widths = [10, 14, 18, 14, 14, 10, 20, 16, 40]
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='耗材库存盘点模板.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/api/consumables/import', methods=['POST'])
@require_login
def import_consumables():
    """批量导入耗材库存：种类唯一（型号+品类不重复建），库存按 区域+库位 upsert。
    填了数量的行更新该区域库存，留空保留原值；单价/阈值/备注/编码留空不覆盖；
    同型号可在中控/办公楼各建区域库存行（种类只建一次，种类数不增加）"""
    f = request.files.get('file')
    if not f or not f.filename:
        return jsonify({'ok': False, 'error': '请选择要导入的 Excel 文件'})
    if not (f.filename.endswith('.xlsx') or f.filename.endswith('.xls')):
        return jsonify({'ok': False, 'error': '仅支持 .xlsx / .xls 文件'})
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(f.read()), data_only=True)
    except Exception as e:
        return jsonify({'ok': False, 'error': '文件解析失败：%s' % e})
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    conn = get_conn()
    added = updated = failed = 0
    errors = []

    def cell_num(v):
        if v is None:
            return None
        s = str(v).strip()
        if s == '':
            return None
        try:
            return float(s)
        except (ValueError, TypeError):
            return 'ERR'

    for idx, row in enumerate(rows, start=2):
        if row is None or all(v is None or str(v).strip() == '' for v in row):
            continue
        region = str(row[0] or '').strip() if len(row) > 0 and row[0] is not None else ''
        storage = str(row[1] or '').strip() if len(row) > 1 and row[1] is not None else ''
        model = str(row[2] or '').strip() if len(row) > 2 else ''
        category = str(row[3] or '').strip() if len(row) > 3 else ''
        if not model or not category:
            failed += 1
            errors.append(f"第{idx}行：打印机型号与耗材品类必填")
            continue
        code = str(row[4] or '').strip() if len(row) > 4 and row[4] is not None else ''
        price = cell_num(row[5]) if len(row) > 5 else None
        stock = cell_num(row[6]) if len(row) > 6 else None
        low = cell_num(row[7]) if len(row) > 7 else None
        remark = str(row[8] or '').strip() if len(row) > 8 and row[8] is not None else None
        bad = []
        if price == 'ERR': bad.append('单价')
        if stock == 'ERR': bad.append('库存')
        if low == 'ERR': bad.append('阈值')
        if bad:
            failed += 1
            errors.append(f"第{idx}行：{'、'.join(bad)}不是数字")
            continue
        region = region if region in ('中控', '办公楼') else '办公楼'
        try:
            # 种类唯一：型号+品类
            con = conn.execute("SELECT * FROM consumables WHERE model=? AND category=?",
                               (model, category)).fetchone()
            if not con:
                conn.execute(
                    "INSERT INTO consumables(model,category,material_code,price,low_threshold,remark) VALUES(?,?,?,?,?,?)",
                    (model, category, code, price or 0, low or 0, remark or ''))
                con = conn.execute("SELECT * FROM consumables WHERE model=? AND category=?",
                                   (model, category)).fetchone()
                added += 1
            else:
                sets, args = [], []
                if code:
                    sets.append('material_code=?'); args.append(code)
                if price is not None:
                    sets.append('price=?'); args.append(price)
                if low is not None:
                    sets.append('low_threshold=?'); args.append(low)
                if remark:
                    sets.append('remark=?'); args.append(remark)
                if sets:
                    args.append(con['id'])
                    conn.execute("UPDATE consumables SET %s WHERE id=?" % ','.join(sets), args)
            # 区域库存行 upsert（同种类可中控/办公楼各一行）
            srow = _con_stock_row(conn, con['id'], region)
            if not srow:
                conn.execute(
                    "INSERT INTO consumable_stocks(consumable_id, region, storage, stock) VALUES(?,?,?,?)",
                    (con['id'], region, storage, stock if stock is not None else 0))
            else:
                sets2, args2 = [], []
                if stock is not None:
                    sets2.append('stock=?'); args2.append(stock)
                if storage:
                    sets2.append('storage=?'); args2.append(storage)
                if sets2:
                    args2.append(srow['id'])
                    conn.execute("UPDATE consumable_stocks SET %s WHERE id=?" % ','.join(sets2), args2)
            _con_recalc(conn, con['id'])
            updated += 1
        except Exception as e:
            failed += 1
            errors.append(f"第{idx}行：{e}")
    conn.commit()
    conn.close()
    log(session['username'], '批量导入耗材库存', f"added={added} updated={updated} failed={failed}")
    return jsonify({'ok': True, 'added': added, 'updated': updated, 'failed': failed,
                    'errors': errors[:20]})


# ---------------- 耗材维护记录 ----------------
@app.route('/api/maintenance', methods=['GET'])
@require_login
def list_maintenance():
    q = request.args.get('q', '').strip()
    start = request.args.get('start', '').strip()
    end = request.args.get('end', '').strip()
    sql = "SELECT * FROM consumable_maintenance WHERE 1=1"
    args = []
    if q:
        sql += " AND (model LIKE ? OR category LIKE ? OR content LIKE ? OR maintainer LIKE ? OR remark LIKE ?)"
        like = f"%{q}%"
        args += [like] * 5
    if start:
        sql += " AND maintain_time >= ?"
        args.append(start)
    if end:
        sql += " AND maintain_time <= ?"
        args.append(end)
    sql += " ORDER BY maintain_time DESC, id DESC"
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows]})


@app.route('/api/maintenance', methods=['POST'])
@require_login
def add_maintenance():
    d = request.get_json(force=True, silent=True) or {}
    model = str(d.get('model', '')).strip()
    category = str(d.get('category', '')).strip()
    content = str(d.get('content', '')).strip()
    if not model or not content:
        return jsonify({'ok': False, 'error': '打印机型号与维护内容必填'})
    maintain_time = str(d.get('maintain_time', '')).strip() or now_str()[:10]
    try:
        cost = num(d.get('cost'))
    except (TypeError, ValueError):
        cost = 0
    conn = get_conn()
    cur = conn.execute(
        "INSERT INTO consumable_maintenance(maintain_time,model,category,content,cost,maintainer,remark,created_at) "
        "VALUES(?,?,?,?,?,?,?,?)",
        (maintain_time, model, category, content, cost,
         str(d.get('maintainer', '')).strip() or (session['real_name'] or session['username']),
         str(d.get('remark', '')).strip(), now_str()))
    conn.commit()
    conn.close()
    log(session['username'], '登记耗材维护', f"{model} {content}")
    return jsonify({'ok': True, 'id': cur.lastrowid})


@app.route('/api/maintenance/<int:mid>', methods=['DELETE'])
@require_login
def delete_maintenance(mid):
    conn = get_conn()
    conn.execute("DELETE FROM consumable_maintenance WHERE id=?", (mid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除维护记录', str(mid))
    return jsonify({'ok': True})


# ---------------- 耗材发放 ----------------
@app.route('/api/issues', methods=['GET'])
@require_login
def list_issues():
    q = request.args.get('q', '').strip()
    start = request.args.get('start', '').strip()
    end = request.args.get('end', '').strip()
    page = max(1, int(request.args.get('page', 1)))
    per = min(max(1, int(request.args.get('per', 30))), 100000)
    sql = "SELECT * FROM consumable_issues WHERE 1=1"
    args = []
    if q:
        sql += " AND (oa_no LIKE ? OR department LIKE ? OR applicant LIKE ? OR model LIKE ? OR consumable LIKE ? OR location LIKE ?)"
        like = f"%{q}%"
        args += [like] * 6
    if start:
        sql += " AND issue_time >= ?"
        args.append(start)
    if end:
        sql += " AND issue_time <= ?"
        args.append(end)
    total = get_conn().execute(sql.replace("SELECT *", "SELECT COUNT(*) c"), args).fetchone()['c']
    sql += " ORDER BY issue_time DESC, id DESC LIMIT ? OFFSET ?"
    args += [per, (page - 1) * per]
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows], 'total': total, 'page': page, 'per': per})


@app.route('/api/issues', methods=['POST'])
@require_login
def add_issue():
    """发放耗材：事务内按区域扣减库存并登记记录"""
    d = request.get_json(force=True, silent=True) or {}
    cid = d.get('consumable_id')
    qty = num(d.get('qty'))
    if not cid or qty <= 0:
        return jsonify({'ok': False, 'error': '请选择耗材并填写正确的数量'})
    region = str(d.get('region', '')).strip()
    region = region if region in ('中控', '办公楼') else '办公楼'
    try:
        issue_time = d.get('issue_time') or now_str()[:10]
        dt = datetime.datetime.strptime(issue_time[:10], '%Y-%m-%d')
    except ValueError:
        dt = datetime.datetime.now()
        issue_time = dt.strftime('%Y-%m-%d')
    conn = get_conn()
    row = conn.execute("SELECT * FROM consumables WHERE id=?", (cid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '耗材不存在'}), 404
    srow = _con_stock_row(conn, cid, region)
    if not srow:
        conn.close()
        return jsonify({'ok': False, 'error': '该耗材在%s区域未建立库存，请先在该区域入库/调整' % region})
    if num(srow['stock']) < qty:
        conn.close()
        return jsonify({'ok': False, 'error': '库存不足！%s区域当前库存 %s，本次发放 %s' % (region, srow['stock'], qty)})
    oa = str(d.get('oa_no', '')).strip()
    if not oa:
        conn.close()
        return jsonify({'ok': False, 'error': '请填写OA流程号'})
    department = str(d.get('department', '')).strip()
    if not department:
        conn.close()
        return jsonify({'ok': False, 'error': '领用部门必填'})
    operator = str(d.get('operator', '')).strip() or (session['real_name'] or session['username'])
    conn.execute("UPDATE consumable_stocks SET stock=? WHERE id=?", (num(srow['stock']) - qty, srow['id']))
    _con_recalc(conn, cid)
    conn.execute(
        "INSERT INTO consumable_issues(oa_no,location,model,consumable,qty,department,applicant,operator,issue_time,created_at,price,consumable_id,region) "
        "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (oa, str(d.get('location', '')).strip(), str(d.get('model', '')).strip(),
         str(d.get('consumable', '')).strip(), qty, department,
         str(d.get('applicant', '')).strip(), operator,
         issue_time, now_str(), num(row['price']), cid, region))
    conn.commit()
    conn.close()
    log(session['username'], '耗材发放', f"consumable={cid} region={region} qty={qty} oa={oa}")
    return jsonify({'ok': True, 'oa_no': oa})


@app.route('/api/issues/<int:iid>', methods=['DELETE'])
@require_login
def delete_issue(iid):
    """删除发放记录（按记录的区域回补库存）"""
    conn = get_conn()
    row = conn.execute("SELECT * FROM consumable_issues WHERE id=?", (iid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '记录不存在'}), 404
    qty = num(row['qty'])
    restored = False
    # 1) 优先用记录中的耗材ID + 区域精准回补
    cid = row['consumable_id'] if ('consumable_id' in row.keys() and row['consumable_id']) else None
    old_region = row['region'] if ('region' in row.keys() and row['region']) else ''
    con = conn.execute("SELECT * FROM consumables WHERE id=?", (cid,)).fetchone() if cid else None
    # 2) 历史记录无耗材ID：按 型号+品类 唯一匹配兜底
    if not con:
        model = str(row['model'] or '').strip()
        cat = str(row['consumable'] or '').strip()
        if model or cat:
            rows = conn.execute(
                "SELECT * FROM consumables WHERE model=? AND category=?",
                (model, cat)).fetchall()
            if len(rows) == 1:
                con = rows[0]
                cid = con['id']
    if con and cid and qty > 0:
        srow = _con_refund_row(conn, cid, old_region)
        if srow:
            conn.execute("UPDATE consumable_stocks SET stock=? WHERE id=?",
                         (num(srow['stock']) + qty, srow['id']))
            _con_recalc(conn, cid)
            restored = True
    conn.execute("DELETE FROM consumable_issues WHERE id=?", (iid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除发放记录', f"{iid} qty={qty} 回补={restored}")
    return jsonify({'ok': True, 'restored': restored})


@app.route('/api/issues/<int:iid>', methods=['PUT'])
@require_login
def update_issue(iid):
    """编辑发放记录：支持修改数量/耗材/区域等字段，库存按区域自动联动
    同耗材同区域改数量：按差值补扣/回补；换耗材或换区域：旧回补原数量 + 新扣减新数量"""
    d = request.get_json(force=True, silent=True) or {}
    cid = d.get('consumable_id')
    qty = num(d.get('qty'))
    if not cid or qty <= 0:
        return jsonify({'ok': False, 'error': '请选择耗材并填写正确的数量'})
    new_region = str(d.get('region', '')).strip()
    new_region = new_region if new_region in ('中控', '办公楼') else ''
    conn = get_conn()
    row = conn.execute("SELECT * FROM consumable_issues WHERE id=?", (iid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '发放记录不存在'}), 404
    new_con = conn.execute("SELECT * FROM consumables WHERE id=?", (cid,)).fetchone()
    if not new_con:
        conn.close()
        return jsonify({'ok': False, 'error': '耗材不存在'}), 404
    oa = str(d.get('oa_no', '')).strip()
    if not oa:
        conn.close()
        return jsonify({'ok': False, 'error': '请填写OA流程号'})
    department = str(d.get('department', '')).strip()
    if not department:
        conn.close()
        return jsonify({'ok': False, 'error': '领用部门必填'})
    try:
        issue_time = d.get('issue_time') or now_str()[:10]
        dt = datetime.datetime.strptime(issue_time[:10], '%Y-%m-%d')
        issue_time = dt.strftime('%Y-%m-%d')
    except ValueError:
        issue_time = now_str()[:10]
    old_qty = num(row['qty'])
    old_cid = row['consumable_id'] if ('consumable_id' in row.keys() and row['consumable_id']) else None
    old_region = row['region'] if ('region' in row.keys() and row['region']) else ''
    old_con = conn.execute("SELECT * FROM consumables WHERE id=?", (old_cid,)).fetchone() if old_cid else None
    # 历史记录无耗材ID：按 型号+品类 唯一匹配兜底
    if not old_con:
        model0 = str(row['model'] or '').strip()
        cat0 = str(row['consumable'] or '').strip()
        if model0 or cat0:
            cands = conn.execute("SELECT * FROM consumables WHERE model=? AND category=?",
                                 (model0, cat0)).fetchall()
            if len(cands) == 1:
                old_con = cands[0]
                old_cid = old_con['id']
    same_con = old_con and old_con['id'] == cid
    same_region = bool(old_region) and old_region == new_region
    # 先统一校验最终库存是否够（不落任何中间写）
    if same_con and same_region:
        srow = _con_stock_row(conn, cid, new_region or old_region)
        cur_stock = num(srow['stock']) if srow else 0
        final_stock = cur_stock - qty + old_qty
        if final_stock < 0:
            conn.close()
            return jsonify({'ok': False, 'error': '库存不足！%s区域调整后库存将为 %s，请检查数量' % (new_region or old_region, final_stock)})
    else:
        if new_region:
            nsrow = _con_stock_row(conn, cid, new_region)
            if not nsrow:
                conn.close()
                return jsonify({'ok': False, 'error': '该耗材在%s区域未建立库存，请先在该区域入库/调整' % new_region})
            if num(nsrow['stock']) < qty:
                conn.close()
                return jsonify({'ok': False, 'error': '库存不足！%s区域当前库存 %s，本次发放 %s' % (new_region, nsrow['stock'], qty)})
    # 联动库存（按区域）
    if same_con and same_region:
        diff = qty - old_qty
        if diff:
            srow = _con_stock_row(conn, cid, old_region)
            if srow:
                conn.execute("UPDATE consumable_stocks SET stock=? WHERE id=?",
                             (num(srow['stock']) - diff, srow['id']))
                _con_recalc(conn, cid)
    else:
        # 旧记录回补
        if old_con and old_qty > 0:
            osrow = _con_refund_row(conn, old_con['id'], old_region)
            if osrow:
                conn.execute("UPDATE consumable_stocks SET stock=? WHERE id=?",
                             (num(osrow['stock']) + old_qty, osrow['id']))
                _con_recalc(conn, old_con['id'])
        # 新记录扣减
        if new_region:
            nsrow = _con_stock_row(conn, cid, new_region)
            if nsrow:
                conn.execute("UPDATE consumable_stocks SET stock=? WHERE id=?",
                             (num(nsrow['stock']) - qty, nsrow['id']))
                _con_recalc(conn, cid)
    operator = str(d.get('operator', '')).strip() or (session.get('real_name') or session['username'])
    price = num(new_con['price']) if (not same_con) else num(row['price'])
    conn.execute(
        "UPDATE consumable_issues SET consumable_id=?, model=?, consumable=?, qty=?, price=?, oa_no=?, location=?, department=?, applicant=?, operator=?, issue_time=?, region=? WHERE id=?",
        (cid, str(d.get('model', '')).strip(), str(d.get('consumable', '')).strip(), qty, price,
         oa, str(d.get('location', '')).strip(), department,
         str(d.get('applicant', '')).strip(), operator, issue_time, new_region or old_region, iid))
    conn.commit()
    conn.close()
    log(session['username'], '编辑发放记录', f"issue={iid} qty={old_qty}->{qty} cid={cid} region={new_region or old_region}")
    return jsonify({'ok': True})


# ---------------- 批量发放 / 导出 / 费用统计 / 预算 ----------------
@app.route('/api/issues/batch', methods=['POST'])
@require_login
def batch_add_issues():
    """批量发放多种耗材：事务内逐项校验库存并扣减、生成发放记录；任一失败整体回滚"""
    d = request.get_json(force=True, silent=True) or {}
    items = d.get('items') or []
    if not items:
        return jsonify({'ok': False, 'error': '请至少选择一种耗材'})
    department = str(d.get('department', '')).strip()
    applicant = str(d.get('applicant', '')).strip()
    location = str(d.get('location', '')).strip()
    operator = str(d.get('operator', '')).strip() or (session['real_name'] or session['username'])
    oa = str(d.get('oa_no', '')).strip()
    if not oa:
        return jsonify({'ok': False, 'error': '请填写OA流程号'})
    if not department:
        return jsonify({'ok': False, 'error': '领用部门必填'})
    try:
        issue_time = d.get('issue_time') or now_str()[:10]
        dt = datetime.datetime.strptime(issue_time[:10], '%Y-%m-%d')
    except ValueError:
        dt = datetime.datetime.now()
        issue_time = dt.strftime('%Y-%m-%d')
    conn = get_conn()
    prepared = []
    for it in items:
        cid = it.get('consumable_id')
        qty = num(it.get('qty'))
        if not cid or qty <= 0:
            conn.close()
            return jsonify({'ok': False, 'error': '存在未选择耗材或数量无效的行，请检查'})
        region = str(it.get('region', '')).strip()
        region = region if region in ('中控', '办公楼') else '办公楼'
        row = conn.execute("SELECT * FROM consumables WHERE id=?", (cid,)).fetchone()
        if not row:
            conn.close()
            return jsonify({'ok': False, 'error': '耗材不存在: id=%s' % cid})
        srow = _con_stock_row(conn, cid, region)
        if not srow:
            conn.close()
            return jsonify({'ok': False, 'error': '【%s %s】%s区域未建立库存，请先在该区域入库' % (row['model'], row['category'], region)})
        if num(srow['stock']) < qty:
            conn.close()
            return jsonify({'ok': False, 'error': '【%s %s】%s区域库存不足！当前库存 %s，本次发放 %s' % (
                row['model'], row['category'], region, srow['stock'], qty)})
        prepared.append((cid, row, qty, oa, region, srow['id']))
    try:
        for cid, row, qty, oa, region, sid in prepared:
            conn.execute("UPDATE consumable_stocks SET stock=stock-? WHERE id=?", (qty, sid))
            _con_recalc(conn, cid)
            conn.execute(
                "INSERT INTO consumable_issues(oa_no,location,model,consumable,qty,department,applicant,operator,issue_time,created_at,price,consumable_id,region) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (oa, location, row['model'], row['category'], qty, department, applicant,
                 operator, issue_time, now_str(), num(row['price']), row['id'], region))
        conn.commit()
    except Exception as e:
        conn.rollback()
        conn.close()
        return jsonify({'ok': False, 'error': '批量发放失败: %s' % e})
    conn.close()
    log(session['username'], '批量耗材发放', f"items={len(prepared)} dept={department} applicant={applicant}")
    return jsonify({'ok': True, 'count': len(prepared), 'oa_nos': [x[3] for x in prepared]})


@app.route('/api/issues/export', methods=['GET'])
@require_login
def export_issues():
    """按条件导出发放记录为 Excel 文件"""
    q = request.args.get('q', '').strip()
    start = request.args.get('start', '').strip()
    end = request.args.get('end', '').strip()
    sql = "SELECT * FROM consumable_issues WHERE 1=1"
    args = []
    if q:
        sql += " AND (oa_no LIKE ? OR department LIKE ? OR applicant LIKE ? OR model LIKE ? OR consumable LIKE ? OR location LIKE ?)"
        like = f"%{q}%"
        args += [like] * 6
    if start:
        sql += " AND issue_time >= ?"
        args.append(start)
    if end:
        sql += " AND issue_time <= ?"
        args.append(end)
    ids = _ids_arg()
    if ids:
        sql += " AND id IN (%s)" % ','.join('?' * len(ids))
        args += ids
    sql += " ORDER BY issue_time DESC, id DESC"
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    import io as _io
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '发放记录'
    headers = ['OA流程号', '日期', '区域', '部门', '领用人', '位置', '打印机型号', '耗材', '数量', '单价(元)', '金额(元)', '登记人']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='D9E2F3')
        cell.alignment = Alignment(horizontal='center')
    total_amt = 0.0
    for r in rows:
        amt = num(r['qty']) * num(r['price'])
        total_amt += amt
        ws.append([r['oa_no'], r['issue_time'], r['region'] if ('region' in r.keys() and r['region']) else '',
                   r['department'], r['applicant'], r['location'],
                   r['model'], r['consumable'], num(r['qty']), num(r['price']), round(amt, 2), r['operator']])
    ws.append(['合计', '', '', '', '', '', '', '', '', '', round(total_amt, 2), ''])
    ws.cell(row=ws.max_row, column=11).font = Font(bold=True)
    widths = [16, 12, 8, 14, 10, 16, 15, 12, 8, 10, 12, 10]
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = '耗材发放记录_%s.xlsx' % datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    log(session['username'], '导出发放记录', f"rows={len(rows)}")
    return send_file(buf, as_attachment=True, download_name=fname,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ---------------- 发放记录导入（含关键词匹配） ----------------
CONSUME_BRANDS = ['施乐', '惠普', '佳能', '爱普生', '三星', '联想', '兄弟', '理光',
                  '京瓷', '震旦', '柯尼卡', '东芝', '富士', '利盟', '奔图']
# 耗材品类同义词：黑白机"粉=硒鼓"、彩机"色粉=色硒鼓"等
CONSUME_SYN = {
    '黑粉': ['黑硒鼓', '硒鼓', '黑鼓', '黑色墨粉'],
    '黑硒鼓': ['黑粉', '硒鼓', '黑鼓'],
    '红粉': ['红硒鼓'], '红硒鼓': ['红粉'],
    '青粉': ['青硒鼓'], '青硒鼓': ['青粉'],
    '黄粉': ['黄硒鼓'], '黄硒鼓': ['黄粉'],
    '硒鼓': ['黑粉', '黑硒鼓', '黑鼓'],
}
# 型号别名：耗材通用的机型（导入时按别名匹配库存条目）
CONSUME_ALIAS = {
    '施乐V2265': '施乐C2560',
}


def _norm_model(s):
    return re.sub(r'[\s\-/_（）()]', '', (s or '').upper())


def _brand_of(s):
    for b in CONSUME_BRANDS:
        if b in s:
            return b
    return ''


def _digits_of(s):
    return ''.join(re.findall(r'\d+', s or ''))


def _cat_ok(t, e):
    t, e = (t or '').strip(), (e or '').strip()
    if not t or not e:
        return False
    if t == e:
        return True
    if t in e or e in t:
        return True
    if e in CONSUME_SYN.get(t, []):
        return True
    if t in CONSUME_SYN.get(e, []):
        return True
    return False


def _find_best_con(conn, model, consumable):
    """关键词匹配耗材：品牌+数字近似 + 品类同义/包含 + 型号别名，返回最匹配条目或 None"""
    model = CONSUME_ALIAS.get((model or '').strip(), model or '')
    best, bs = None, 0
    for ex in conn.execute("SELECT * FROM consumables").fetchall():
        if not _cat_ok(consumable, ex['category']):
            continue
        t, e = _norm_model(model), _norm_model(ex['model'])
        if not t or not e:
            continue
        if t == e:
            s = 100
        elif t in e or e in t:
            s = 60 + min(len(t), len(e))
        else:
            tb, eb = _brand_of(model), _brand_of(ex['model'])
            td, ed = _digits_of(model), _digits_of(ex['model'])
            if tb and tb == eb:
                if td and td == ed:
                    s = 80
                elif td and ed and (td in ed or ed in td):
                    s = 60
                else:
                    s = 10
            else:
                s = 0
        if s > bs:
            bs, best = s, ex
    return best if best and bs >= 60 else None


@app.route('/api/issues/template', methods=['GET'])
@require_login
def issue_template():
    """下载发放记录批量导入模板（xlsx）"""
    import openpyxl, io as _io
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '发放记录导入模板'
    headers = ['OA流程号*', '发放时间', '区域', '打印机所处地', '打印机型号*', '耗材品类*', '数量*', '领用部门*', '领用人*', '登记人/经办人']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='D9E2F3')
        cell.alignment = Alignment(horizontal='center')
    ws.append(['IT02202609001', '2026-09-25', '办公楼', '中控室217办公室', '施乐C2560', '硒鼓', 1, '信息中心', '张三', '李龙帅'])
    ws.append(['IT02202609002', '', '中控', '5#岛B207', '惠普M254', '黑色墨盒', 2, '机电仪部', '李四', ''])
    widths = [18, 12, 8, 20, 14, 12, 8, 14, 12, 14]
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='发放记录导入模板.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/api/issues/import', methods=['POST'])
@require_login
def import_issues():
    """批量导入发放记录：逐行校验并自动扣减库存、写价格快照；失败行列原因，成功行入库"""
    f = request.files.get('file')
    if not f or not f.filename:
        return jsonify({'ok': False, 'error': '请选择要导入的 Excel 文件'})
    if not (f.filename.endswith('.xlsx') or f.filename.endswith('.xls')):
        return jsonify({'ok': False, 'error': '仅支持 .xlsx / .xls 文件'})
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(f.read()), data_only=True)
    except Exception as e:
        return jsonify({'ok': False, 'error': '文件解析失败：%s' % e})
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    conn = get_conn()
    added = failed = 0
    errors = []
    today = now_str()[:10]
    for idx, row in enumerate(rows, start=2):
        if row is None or all(v is None or str(v).strip() == '' for v in row):
            continue
        oa = str(row[0] or '').strip() if len(row) > 0 else ''
        issue_time = str(row[1] or '').strip() if len(row) > 1 and row[1] is not None else ''
        region = str(row[2] or '').strip() if len(row) > 2 and row[2] is not None else ''
        location = str(row[3] or '').strip() if len(row) > 3 and row[3] is not None else ''
        model = str(row[4] or '').strip() if len(row) > 4 and row[4] is not None else ''
        consumable = str(row[5] or '').strip() if len(row) > 5 and row[5] is not None else ''
        try:
            qty = num(row[6]) if len(row) > 6 and row[6] is not None else 0
        except (TypeError, ValueError):
            qty = 0
        department = str(row[7] or '').strip() if len(row) > 7 and row[7] is not None else ''
        applicant = str(row[8] or '').strip() if len(row) > 8 and row[8] is not None else ''
        operator = str(row[9] or '').strip() if len(row) > 9 and row[9] is not None else ''
        operator = operator or (session['real_name'] or session['username'])
        region = region if region in ('中控', '办公楼') else '办公楼'
        if not oa:
            # OA 流程号允许留空（补充历史记录时可不填）；非空时才查重
            oa = ''
        if not model or not consumable:
            failed += 1
            errors.append(f"第{idx}行：打印机型号与耗材品类必填")
            continue
        if not qty or qty <= 0:
            failed += 1
            errors.append(f"第{idx}行：数量无效（应为正数）")
            continue
        if not department or not applicant:
            failed += 1
            errors.append(f"第{idx}行：领用部门与领用人必填")
            continue
        con = conn.execute("SELECT * FROM consumables WHERE model=? AND category=?",
                           (model, consumable)).fetchone()
        if not con:
            con = _find_best_con(conn, model, consumable)  # 关键词匹配兜底
        if not con:
            failed += 1
            errors.append(f"第{idx}行：耗材库中无「{model} {consumable}」，请先在耗材库存中导入/新增")
            continue
        if oa:
            # 同一 OA 可含多行不同耗材（一次申请多种）；查重须用入库规范名（匹配后），
            # 仅当 OA+型号+品类+区域 完全相同才算重复
            dup = conn.execute(
                "SELECT id FROM consumable_issues WHERE oa_no=? AND model=? AND consumable=? AND region=?",
                (oa, con['model'], con['category'], region)).fetchone()
            if dup:
                failed += 1
                errors.append(f"第{idx}行：该记录已存在（{oa} {con['model']} {con['category']} {region}），跳过避免重复")
                continue
        else:
            # OA 留空时按 型号+品类+区域+地点+数量+部门+领用人 判重，防止重复导入
            dup = conn.execute(
                "SELECT id FROM consumable_issues WHERE oa_no='' AND model=? AND consumable=? AND region=? AND location=? AND qty=? AND department=? AND applicant=?",
                (con['model'], con['category'], region, location, qty, department, applicant)).fetchone()
            if dup:
                failed += 1
                errors.append(f"第{idx}行：该记录已存在（无OA {con['model']} {con['category']} {region}），跳过避免重复")
                continue
        srow = _con_stock_row(conn, con['id'], region)
        if not srow or num(srow['stock']) < qty:
            failed += 1
            errors.append(f"第{idx}行：【{model} {consumable}】{region}区域库存不足（当前库存 %s）" % (srow['stock'] if srow else 0))
            continue
        try:
            conn.execute("UPDATE consumable_stocks SET stock=? WHERE id=?",
                         (num(srow['stock']) - qty, srow['id']))
            _con_recalc(conn, con['id'])
            conn.execute(
                "INSERT INTO consumable_issues(oa_no,location,model,consumable,qty,department,applicant,operator,issue_time,created_at,price,consumable_id,region) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (oa, location, con['model'], con['category'], qty, department, applicant,
                 operator, issue_time or today, now_str(), num(con['price']), con['id'], region))
            added += 1
        except Exception as e:
            failed += 1
            errors.append(f"第{idx}行：{e}")
    conn.commit()
    conn.close()
    log(session['username'], '批量导入发放记录', f"added={added} failed={failed}")
    return jsonify({'ok': True, 'added': added, 'failed': failed, 'errors': errors[:20]})


@app.route('/api/statistics', methods=['GET'])
@require_login
def statistics():
    """耗材费用统计：月度趋势 / 部门分布 / 耗材TOP / 区间汇总 / 预算"""
    start = request.args.get('start', '').strip()
    end = request.args.get('end', '').strip()
    conn = get_conn()
    where = "1=1"
    args = []
    if start:
        where += " AND issue_time >= ?"
        args.append(start)
    if end:
        where += " AND issue_time <= ?"
        args.append(end)
    # 区间汇总
    total = conn.execute(
        "SELECT COUNT(*) c, COALESCE(SUM(qty),0) q, COALESCE(SUM(qty*price),0) a "
        "FROM consumable_issues WHERE %s" % where, args).fetchone()
    # 月度趋势（近12个月）
    months = []
    now = datetime.date.today()
    for i in range(11, -1, -1):
        y, m = now.year, now.month - i
        while m <= 0:
            m += 12
            y -= 1
        months.append('%04d-%02d' % (y, m))
    trend = []
    for mm in months:
        row = conn.execute(
            "SELECT COUNT(*) c, COALESCE(SUM(qty),0) q, COALESCE(SUM(qty*price),0) a "
            "FROM consumable_issues WHERE issue_time LIKE ?", (mm + '%',)).fetchone()
        trend.append({'month': mm, 'count': row['c'], 'qty': row['q'], 'amount': row['a']})
    # 部门分布（区间内，按金额排序，全量返回）
    depts = conn.execute(
        "SELECT department, COUNT(*) c, COALESCE(SUM(qty),0) q, COALESCE(SUM(qty*price),0) a "
        "FROM consumable_issues WHERE %s AND department != '' GROUP BY department ORDER BY a DESC" % where,
        args).fetchall()
    # 耗材 TOP（区间内，全量返回）
    tops = conn.execute(
        "SELECT model, consumable, COUNT(*) c, COALESCE(SUM(qty),0) q, COALESCE(SUM(qty*price),0) a "
        "FROM consumable_issues WHERE %s GROUP BY model, consumable ORDER BY a DESC" % where,
        args).fetchall()
    # 经办人统计
    operators = conn.execute(
        "SELECT operator, COUNT(*) c, COALESCE(SUM(qty),0) q, COALESCE(SUM(qty*price),0) a "
        "FROM consumable_issues WHERE %s AND operator != '' GROUP BY operator ORDER BY a DESC" % where,
        args).fetchall()
    # 单价为 0/空的记录数（提醒补价）
    missing_price = conn.execute(
        "SELECT COUNT(*) c FROM consumable_issues WHERE %s AND (price IS NULL OR price=0)" % where, args).fetchone()['c']
    # 发放时间为空的记录数
    missing_time = conn.execute(
        "SELECT COUNT(*) c FROM consumable_issues WHERE %s AND (issue_time IS NULL OR issue_time='')" % where, args).fetchone()['c']
    # 预算
    budgets = conn.execute("SELECT * FROM budgets ORDER BY year, month").fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': {
        'total': dict(total), 'trend': trend,
        'departments': [dict(r) for r in depts],
        'tops': [dict(r) for r in tops],
        'operators': [dict(r) for r in operators],
        'missing_price': missing_price,
        'missing_time': missing_time,
        'budgets': [dict(r) for r in budgets],
    }})


@app.route('/api/statistics/export', methods=['GET'])
@require_login
def export_statistics():
    """导出费用统计：发放明细 / 部门汇总 / 耗材TOP / 月度趋势（按 start~end 区间）"""
    start = request.args.get('start', '').strip()
    end = request.args.get('end', '').strip()
    where = "1=1"
    args = []
    if start:
        where += " AND issue_time >= ?"
        args.append(start)
    if end:
        where += " AND issue_time <= ?"
        args.append(end)
    conn = get_conn()
    rows = conn.execute(
        "SELECT * FROM consumable_issues WHERE %s ORDER BY issue_time DESC, id DESC" % where, args).fetchall()
    depts = conn.execute(
        "SELECT department, COUNT(*) c, COALESCE(SUM(qty),0) q, COALESCE(SUM(qty*price),0) a "
        "FROM consumable_issues WHERE %s AND department != '' GROUP BY department ORDER BY a DESC" % where,
        args).fetchall()
    tops = conn.execute(
        "SELECT model, consumable, COUNT(*) c, COALESCE(SUM(qty),0) q, COALESCE(SUM(qty*price),0) a "
        "FROM consumable_issues WHERE %s GROUP BY model, consumable ORDER BY a DESC" % where, args).fetchall()
    months = []
    now = datetime.date.today()
    for i in range(11, -1, -1):
        y, m = now.year, now.month - i
        while m <= 0:
            m += 12
            y -= 1
        months.append('%04d-%02d' % (y, m))
    trend = []
    for mm in months:
        row = conn.execute(
            "SELECT COUNT(*) c, COALESCE(SUM(qty),0) q, COALESCE(SUM(qty*price),0) a "
            "FROM consumable_issues WHERE issue_time LIKE ?", (mm + '%',)).fetchone()
        trend.append([mm, row['c'], row['q'], row['a']])
    conn.close()

    import openpyxl
    from openpyxl.styles import Font, Alignment, PatternFill
    import io as _io
    from openpyxl.utils import get_column_letter
    wb = openpyxl.Workbook()

    def style_header(ws):
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill('solid', fgColor='D9E2F3')
            cell.alignment = Alignment(horizontal='center')
        ws.freeze_panes = 'A2'

    def set_widths(ws, widths):
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

    # Sheet1 发放明细
    ws = wb.active
    ws.title = '发放明细'
    headers = ['OA流程号', '发放时间', '领用部门', '需求人', '登记人', '打印机所在地', '打印机型号', '耗材', '数量', '单价(元)', '金额(元)']
    ws.append(headers)
    style_header(ws)
    total_amt = 0.0
    for r in rows:
        amt = num(r['qty']) * num(r['price'])
        total_amt += amt
        ws.append([r['oa_no'], r['issue_time'], r['department'], r['applicant'], r['operator'],
                   r['location'], r['model'], r['consumable'], num(r['qty']), num(r['price']), round(amt, 2)])
    ws.append(['合计', '', '', '', '', '', '', '', '', '', round(total_amt, 2)])
    ws.cell(row=ws.max_row, column=11).font = Font(bold=True)
    set_widths(ws, [16, 13, 15, 10, 10, 16, 16, 14, 8, 10, 12])

    # Sheet2 部门汇总
    ws2 = wb.create_sheet('部门汇总')
    ws2.append(['部门', '发放次数', '数量(件)', '金额(元)'])
    style_header(ws2)
    for r in depts:
        ws2.append([r['department'], r['c'], r['q'], round(r['a'], 2)])
    set_widths(ws2, [20, 12, 12, 14])

    # Sheet3 耗材TOP
    ws3 = wb.create_sheet('耗材TOP')
    ws3.append(['打印机型号', '耗材', '次数', '数量(件)', '金额(元)'])
    style_header(ws3)
    for r in tops:
        ws3.append([r['model'], r['consumable'], r['c'], r['q'], round(r['a'], 2)])
    set_widths(ws3, [20, 18, 10, 12, 14])

    # Sheet4 月度趋势
    ws4 = wb.create_sheet('月度趋势')
    ws4.append(['月份', '发放次数', '数量(件)', '金额(元)'])
    style_header(ws4)
    for r in trend:
        ws4.append([r[0], r[1], r[2], round(r[3], 2)])
    set_widths(ws4, [12, 12, 12, 14])

    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = '费用统计_%s_%s.xlsx' % (start[:10] or '全部', datetime.datetime.now().strftime('%Y%m%d_%H%M%S'))
    log(session['username'], '导出费用统计', f"start={start} end={end} rows={len(rows)}")
    return send_file(buf, as_attachment=True, download_name=fname,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/api/budgets', methods=['GET'])
@require_login
def list_budgets():
    year = request.args.get('year', '').strip()
    sql = "SELECT * FROM budgets WHERE 1=1"
    args = []
    if year:
        sql += " AND year=?"
        args.append(int(year))
    sql += " ORDER BY year, month"
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows]})


@app.route('/api/budgets', methods=['POST'])
@require_login
def upsert_budget():
    d = request.get_json(force=True, silent=True) or {}
    try:
        year = int(d.get('year'))
        month = int(d.get('month'))
        amount = num(d.get('amount'))
    except (TypeError, ValueError):
        return jsonify({'ok': False, 'error': '年份/月份/金额无效'})
    if not (2000 <= year <= 2100 and 1 <= month <= 12):
        return jsonify({'ok': False, 'error': '年份/月份范围无效'})
    if amount < 0:
        return jsonify({'ok': False, 'error': '预算金额不能为负'})
    conn = get_conn()
    conn.execute(
        "INSERT INTO budgets(year,month,amount,remark,created_at,updated_at) VALUES(?,?,?,?,?,?) "
        "ON CONFLICT(year,month) DO UPDATE SET amount=excluded.amount, remark=excluded.remark, updated_at=excluded.updated_at",
        (year, month, amount, str(d.get('remark', '')).strip(), now_str(), now_str()))
    conn.commit()
    conn.close()
    log(session['username'], '设置预算', f"{year}-{month} = {amount}")
    return jsonify({'ok': True})


@app.route('/api/budgets/<int:bid>', methods=['DELETE'])
@require_login
def delete_budget(bid):
    conn = get_conn()
    conn.execute("DELETE FROM budgets WHERE id=?", (bid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除预算', str(bid))
    return jsonify({'ok': True})


# ---------------- 打印机报修/维修 ----------------
@app.route('/api/repairs', methods=['GET'])
@require_login
def list_repairs():
    pid = request.args.get('printer_id', '').strip()
    status = request.args.get('status', '').strip()
    q = request.args.get('q', '').strip()
    sql = ("SELECT r.*, p.printer_no, p.department, p.location, p.model AS printer_model "
           "FROM printer_repairs r LEFT JOIN printers p ON p.id = r.printer_id WHERE 1=1")
    args = []
    if pid:
        sql += " AND r.printer_id=?"
        args.append(pid)
    if status:
        sql += " AND r.status=?"
        args.append(status)
    if q:
        sql += " AND (r.fault_desc LIKE ? OR r.reported_by LIKE ? OR p.printer_no LIKE ?)"
        like = f"%{q}%"
        args += [like, like, like]
    sql += " ORDER BY r.report_time DESC, r.id DESC"
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows]})


@app.route('/api/repairs', methods=['POST'])
@require_login
def add_repair():
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    cur = conn.execute(
        "INSERT INTO printer_repairs(printer_id,fault_desc,report_time,reported_by,repair_result,"
        "repair_time,repair_company,cost,status,remark,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
        (d.get('printer_id'), str(d.get('fault_desc', '')).strip(),
         d.get('report_time') or now_str()[:10], str(d.get('reported_by', '')).strip(),
         str(d.get('repair_result', '')).strip(), d.get('repair_time') or '',
         str(d.get('repair_company', '')).strip(), num(d.get('cost')),
         d.get('status') or '待维修', str(d.get('remark', '')).strip(), now_str()))
    conn.commit()
    conn.close()
    log(session['username'], '新增报修', str(cur.lastrowid))
    return jsonify({'ok': True, 'id': cur.lastrowid})


@app.route('/api/repairs/<int:rid>', methods=['PUT'])
@require_login
def update_repair(rid):
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    conn.execute(
        "UPDATE printer_repairs SET printer_id=?,fault_desc=?,report_time=?,reported_by=?,repair_result=?,"
        "repair_time=?,repair_company=?,cost=?,status=?,remark=? WHERE id=?",
        (d.get('printer_id'), str(d.get('fault_desc', '')).strip(),
         d.get('report_time') or now_str()[:10], str(d.get('reported_by', '')).strip(),
         str(d.get('repair_result', '')).strip(), d.get('repair_time') or '',
         str(d.get('repair_company', '')).strip(), num(d.get('cost')),
         d.get('status') or '待维修', str(d.get('remark', '')).strip(), rid))
    conn.commit()
    conn.close()
    log(session['username'], '更新报修/维修', str(rid))
    return jsonify({'ok': True})


@app.route('/api/repairs/<int:rid>', methods=['DELETE'])
@require_login
def delete_repair(rid):
    conn = get_conn()
    conn.execute("DELETE FROM printer_repairs WHERE id=?", (rid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除报修记录', str(rid))
    return jsonify({'ok': True})


# ---------------- 备品备件 ----------------
@app.route('/api/items', methods=['GET'])
@require_login
def list_items():
    q = request.args.get('q', '').strip()
    sql = ("SELECT i.*, COALESCE(SUM(s.qty),0) AS total_qty, COUNT(s.id) AS slots "
           "FROM spare_items i LEFT JOIN spare_stocks s ON s.item_id = i.id WHERE 1=1")
    args = []
    if q:
        sql += " AND (i.brand LIKE ? OR i.name LIKE ? OR i.spec LIKE ?)"
        like = f"%{q}%"
        args += [like, like, like]
    sql += " GROUP BY i.id ORDER BY i.name ASC"
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows]})


@app.route('/api/items', methods=['POST'])
@require_login
def add_item():
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    cur = conn.execute("INSERT INTO spare_items(brand,name,spec,unit,remark) VALUES(?,?,?,?,?)",
                       (str(d.get('brand', '')).strip(), str(d.get('name', '')).strip(),
                        str(d.get('spec', '')).strip(), str(d.get('unit', '')).strip(),
                        str(d.get('remark', '')).strip()))
    conn.commit()
    conn.close()
    log(session['username'], '新增备件物资', str(cur.lastrowid))
    return jsonify({'ok': True, 'id': cur.lastrowid})


@app.route('/api/items/<int:iid>', methods=['PUT'])
@require_login
def update_item(iid):
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    conn.execute("UPDATE spare_items SET brand=?,name=?,spec=?,unit=?,remark=? WHERE id=?",
                 (str(d.get('brand', '')).strip(), str(d.get('name', '')).strip(),
                  str(d.get('spec', '')).strip(), str(d.get('unit', '')).strip(),
                  str(d.get('remark', '')).strip(), iid))
    conn.commit()
    conn.close()
    log(session['username'], '修改备件物资', str(iid))
    return jsonify({'ok': True})


@app.route('/api/items/<int:iid>', methods=['DELETE'])
@require_login
def delete_item(iid):
    conn = get_conn()
    # 有库存或流水时禁止直接删除
    n = conn.execute("SELECT COUNT(*) c FROM spare_stocks WHERE item_id=?", (iid,)).fetchone()['c']
    f = conn.execute("SELECT COUNT(*) c FROM spare_flows WHERE item_id=?", (iid,)).fetchone()['c']
    if n > 0 or f > 0:
        conn.close()
        return jsonify({'ok': False, 'error': f'该物资存在库存({n})或流水({f})，不能删除，可改为停用'})
    conn.execute("DELETE FROM spare_items WHERE id=?", (iid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除备件物资', str(iid))
    return jsonify({'ok': True})


@app.route('/api/stocks', methods=['GET'])
@require_login
def list_stocks():
    q = request.args.get('q', '').strip()
    warehouse = request.args.get('warehouse', '').strip()
    region = request.args.get('region', '').strip()
    low_only = request.args.get('low_only', '0') == '1'
    sql = ("SELECT s.*, i.brand, i.name, i.spec, i.unit "
           "FROM spare_stocks s LEFT JOIN spare_items i ON i.id = s.item_id WHERE 1=1")
    args = []
    if q:
        sql += " AND (i.brand LIKE ? OR i.name LIKE ? OR i.spec LIKE ? OR s.location LIKE ?)"
        like = f"%{q}%"
        args += [like, like, like, like]
    if warehouse:
        sql += " AND s.warehouse=?"
        args.append(warehouse)
    if region:
        sql += " AND s.region=?"
        args.append(region)
    if low_only:
        sql += " AND ((s.qty <= s.low_threshold AND s.low_threshold > 0) OR s.qty <= 0)"
    sql += " ORDER BY i.name ASC, s.warehouse ASC, s.location ASC"
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    warehouses = conn.execute(
        "SELECT DISTINCT warehouse FROM spare_stocks WHERE warehouse != '' ORDER BY warehouse").fetchall()
    regions = conn.execute(
        "SELECT DISTINCT region FROM spare_stocks WHERE region != '' ORDER BY region").fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows],
                    'warehouses': [r['warehouse'] for r in warehouses],
                    'regions': [r['region'] for r in regions]})


@app.route('/api/stocks', methods=['POST'])
@require_login
def add_stock():
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    cur = conn.execute(
        "INSERT INTO spare_stocks(item_id,region,warehouse,location,qty,owner,remark,low_threshold) "
        "VALUES(?,?,?,?,?,?,?,?)",
        (d.get('item_id'), str(d.get('region', '')).strip(), str(d.get('warehouse', '')).strip(),
         str(d.get('location', '')).strip(), num(d.get('qty')), str(d.get('owner', '')).strip(),
         str(d.get('remark', '')).strip(), num(d.get('low_threshold'))))
    conn.commit()
    conn.close()
    log(session['username'], '新增库存位', str(cur.lastrowid))
    return jsonify({'ok': True, 'id': cur.lastrowid})


@app.route('/api/stocks/<int:sid>', methods=['PUT'])
@require_login
def update_stock(sid):
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    old = conn.execute("SELECT * FROM spare_stocks WHERE id=?", (sid,)).fetchone()
    if not old:
        conn.close()
        return jsonify({'ok': False, 'error': '库存位不存在'}), 404
    new_qty = num(d.get('qty'))
    conn.execute(
        "UPDATE spare_stocks SET item_id=?,region=?,warehouse=?,location=?,qty=?,owner=?,remark=?,low_threshold=? WHERE id=?",
        (d.get('item_id'), str(d.get('region', '')).strip(), str(d.get('warehouse', '')).strip(),
         str(d.get('location', '')).strip(), new_qty, str(d.get('owner', '')).strip(),
         str(d.get('remark', '')).strip(), num(d.get('low_threshold')), sid))
    # 数量变化写日志
    if num(old['qty']) != new_qty:
        item = conn.execute("SELECT * FROM spare_items WHERE id=?", (old['item_id'],)).fetchone()
        conn.execute(
            "INSERT INTO spare_adjusts(type,item_id,brand,name,spec,unit,warehouse,location,qty,reason,operator,adjust_time,created_at) "
            "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
            ('manual', old['item_id'], item['brand'] if item else '', item['name'] if item else '',
             item['spec'] if item else '', item['unit'] if item else '',
             str(d.get('warehouse', '')).strip(), str(d.get('location', '')).strip(),
             new_qty - num(old['qty']),
             '修改库存位：%s → %s（原 %s 件）' % (old['qty'], new_qty, old['qty']),
             session['real_name'] or session['username'], now_str()[:10], now_str()))
    conn.commit()
    conn.close()
    log(session['username'], '修改库存位', str(sid))
    return jsonify({'ok': True})


@app.route('/api/stocks/<int:sid>', methods=['DELETE'])
@require_login
def delete_stock(sid):
    conn = get_conn()
    old = conn.execute("SELECT * FROM spare_stocks WHERE id=?", (sid,)).fetchone()
    if old:
        item = conn.execute("SELECT * FROM spare_items WHERE id=?", (old['item_id'],)).fetchone()
        conn.execute(
            "INSERT INTO spare_adjusts(type,item_id,brand,name,spec,unit,warehouse,location,qty,reason,operator,adjust_time,created_at) "
            "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
            ('manual', old['item_id'], item['brand'] if item else '', item['name'] if item else '',
             item['spec'] if item else '', item['unit'] if item else '',
             old['warehouse'], old['location'], -num(old['qty']),
             '删除库存位（清空 %s 件）' % old['qty'],
             session['real_name'] or session['username'], now_str()[:10], now_str()))
    conn.execute("DELETE FROM spare_stocks WHERE id=?", (sid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除库存位', str(sid))
    return jsonify({'ok': True})


@app.route('/api/stocks/template', methods=['GET'])
@require_login
def stock_template():
    """下载备件库存导入模板（xlsx）：按 品牌/名称/规格/单位 匹配物资，库位已存在则覆盖库存"""
    import openpyxl, io as _io
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '库存导入模板'
    headers = ['品牌', '物资名称*', '规格型号', '数量*', '单位', '区域', '库房*', '库位*', '物资所有人', '预警阈值', '备注']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='D9E2F3')
        cell.alignment = Alignment(horizontal='center')
    ws.append(['惠普', '原装硒鼓', 'HP 2612A', 20, '个', '一号区域', '一号库', 'A-01', '', 5, ''])
    widths = [12, 18, 16, 10, 8, 12, 12, 10, 12, 10, 18]
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='备件库存导入模板.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/api/stocks/import', methods=['POST'])
@require_login
def import_stocks():
    """批量导入备件库存：逐行校验，全部通过后事务执行（物资/库位不存在自动创建，已存在则覆盖库存数量）"""
    f = request.files.get('file')
    if not f or not f.filename:
        return jsonify({'ok': False, 'error': '请选择要导入的 Excel 文件'})
    if not (f.filename.endswith('.xlsx') or f.filename.endswith('.xls')):
        return jsonify({'ok': False, 'error': '仅支持 .xlsx / .xls 文件'})
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(f.read()), data_only=True)
    except Exception as e:
        return jsonify({'ok': False, 'error': '文件解析失败：%s' % e})
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    conn = get_conn()
    prepared = []
    errors = []
    for idx, row in enumerate(rows, start=2):
        if row is None or all(v is None or str(v).strip() == '' for v in row):
            continue
        def gv(i):
            return str(row[i] or '').strip() if len(row) > i and row[i] is not None else ''
        brand, name, spec = gv(0), gv(1), gv(2)
        try:
            qty = num(row[3]) if len(row) > 3 and row[3] is not None else 0
        except (TypeError, ValueError):
            qty = 0
        unit = gv(4)
        region, warehouse, location = gv(5), gv(6), gv(7)
        owner, low_th, remark = gv(8), gv(9), gv(10)
        if not name:
            errors.append('第%d行：物资名称必填' % idx)
            continue
        if not warehouse or not location:
            errors.append('第%d行：库房与库位必填' % idx)
            continue
        if qty < 0:
            errors.append('第%d行：数量不能为负数' % idx)
            continue
        try:
            low_v = num(low_th) if low_th else 0
        except (TypeError, ValueError):
            low_v = 0
        prepared.append((brand, name, spec, unit, region, warehouse, location, qty, owner, low_v, remark))
    if errors:
        conn.close()
        return jsonify({'ok': False, 'error': '存在 %d 行校验失败，已整体取消导入' % len(errors),
                        'errors': errors[:20]})
    if not prepared:
        conn.close()
        return jsonify({'ok': False, 'error': '模板中没有可导入的数据行'})
    try:
        for pr in prepared:
            brand, name, spec, unit, region, warehouse, location, qty, owner, low_v, remark = pr
            item = conn.execute(
                "SELECT * FROM spare_items WHERE name=? AND brand=? AND spec=? AND unit=? LIMIT 1",
                (name, brand, spec, unit)).fetchone()
            if not item:
                cur = conn.execute("INSERT INTO spare_items(brand,name,spec,unit,remark) VALUES(?,?,?,?,?)",
                                   (brand, name, spec, unit, '库存导入自动创建'))
                item = {'id': cur.lastrowid}
            stock = conn.execute(
                "SELECT * FROM spare_stocks WHERE item_id=? AND warehouse=? AND location=? LIMIT 1",
                (item['id'], warehouse, location)).fetchone()
            if stock:
                conn.execute("UPDATE spare_stocks SET qty=?, region=?, owner=?, low_threshold=?, remark=? WHERE id=?",
                             (qty, region or stock['region'], owner or stock['owner'], low_v, remark, stock['id']))
            else:
                conn.execute(
                    "INSERT INTO spare_stocks(item_id,region,warehouse,location,qty,owner,remark,low_threshold) "
                    "VALUES(?,?,?,?,?,?,?,?)",
                    (item['id'], region, warehouse, location, qty, owner, remark, low_v))
        conn.commit()
        conn.close()
        log(session['username'], '模板导入备件库存', '共%d条' % len(prepared))
        return jsonify({'ok': True, 'added': len(prepared)})
    except Exception as e:
        conn.rollback()
        conn.close()
        return jsonify({'ok': False, 'error': '导入执行失败：%s' % e})


@app.route('/api/flows', methods=['GET'])
@require_login
def list_flows():
    q = request.args.get('q', '').strip()
    ftype = request.args.get('type', '').strip()
    start = request.args.get('start', '').strip()
    end = request.args.get('end', '').strip()
    page = max(1, int(request.args.get('page', 1)))
    per = min(max(1, int(request.args.get('per', 30))), 100000)
    sql = "SELECT * FROM spare_flows WHERE 1=1"
    args = []
    if q:
        sql += " AND (name LIKE ? OR spec LIKE ? OR brand LIKE ? OR applicant LIKE ? OR remark LIKE ?)"
        like = f"%{q}%"
        args += [like] * 5
    if ftype:
        sql += " AND type=?"
        args.append(ftype)
    if start:
        sql += " AND record_time >= ?"
        args.append(start)
    if end:
        sql += " AND record_time <= ?"
        args.append(end)
    total = get_conn().execute(sql.replace("SELECT *", "SELECT COUNT(*) c"), args).fetchone()['c']
    sql += " ORDER BY record_time DESC, id DESC LIMIT ? OFFSET ?"
    args += [per, (page - 1) * per]
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows], 'total': total, 'page': page, 'per': per})


@app.route('/api/flows', methods=['POST'])
@require_login
def add_flow():
    """备件出入库：事务内更新库存并记录流水
    入库：找到或新建目标库位库存行，数量增加
    出库：从指定库位扣减，库存不足则拒绝
    """
    d = request.get_json(force=True, silent=True) or {}
    ftype = d.get('type', 'in')
    if ftype not in ('in', 'out'):
        return jsonify({'ok': False, 'error': '类型必须为入库或出库'})
    qty = num(d.get('qty'))
    if qty <= 0:
        return jsonify({'ok': False, 'error': '数量必须大于 0'})
    item_id = d.get('item_id')
    region = str(d.get('region', '')).strip()
    warehouse = str(d.get('warehouse', '')).strip()
    location = str(d.get('location', '')).strip()
    conn = get_conn()
    # 校验物资
    if item_id:
        item = conn.execute("SELECT * FROM spare_items WHERE id=?", (item_id,)).fetchone()
        if not item:
            conn.close()
            return jsonify({'ok': False, 'error': '物资不存在'}), 404
    else:
        # 无 item_id：尝试按名称匹配，匹配不到则新建物资
        name = str(d.get('name', '')).strip()
        if not name:
            conn.close()
            return jsonify({'ok': False, 'error': '请选择或填写物资名称'})
        item = conn.execute(
            "SELECT * FROM spare_items WHERE name=? AND brand=? AND spec=? AND unit=? LIMIT 1",
            (name, str(d.get('brand', '')).strip(), str(d.get('spec', '')).strip(),
             str(d.get('unit', '')).strip())).fetchone()
        if not item:
            cur = conn.execute("INSERT INTO spare_items(brand,name,spec,unit,remark) VALUES(?,?,?,?,?)",
                               (str(d.get('brand', '')).strip(), name, str(d.get('spec', '')).strip(),
                                str(d.get('unit', '')).strip(), '流水自动创建'))
            item = {'id': cur.lastrowid, 'brand': d.get('brand', ''), 'name': name,
                    'spec': d.get('spec', ''), 'unit': d.get('unit', '')}
        item_id = item['id']
    # 查找目标库位库存行
    stock = conn.execute(
        "SELECT * FROM spare_stocks WHERE item_id=? AND warehouse=? AND location=? LIMIT 1",
        (item_id, warehouse, location)).fetchone()
    if ftype == 'in':
        if stock:
            conn.execute("UPDATE spare_stocks SET qty=qty+?, region=?, owner=? WHERE id=?",
                         (qty, region or stock['region'], d.get('owner', '') or stock['owner'], stock['id']))
        else:
            conn.execute(
                "INSERT INTO spare_stocks(item_id,region,warehouse,location,qty,owner,remark,low_threshold) "
                "VALUES(?,?,?,?,?,?,?,0)",
                (item_id, region, warehouse, location, qty, str(d.get('owner', '')).strip(), ''))
    else:
        if not stock or num(stock['qty']) < qty:
            conn.close()
            return jsonify({'ok': False, 'error': '该库位库存不足或不存在，当前库存 %s' % (stock['qty'] if stock else 0)})
        conn.execute("UPDATE spare_stocks SET qty=qty-? WHERE id=?", (qty, stock['id']))
    try:
        record_time = d.get('record_time') or now_str()[:10]
        datetime.datetime.strptime(record_time[:10], '%Y-%m-%d')
    except ValueError:
        record_time = datetime.datetime.now().strftime('%Y-%m-%d')
    cur = conn.execute(
        "INSERT INTO spare_flows(type,item_id,brand,name,spec,unit,qty,region,warehouse,location,"
        "applicant,reviewer,operator,record_time,remark,created_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (ftype, item_id, item['brand'], item['name'], item['spec'], item['unit'], qty,
         region, warehouse, location, str(d.get('applicant', '')).strip(),
         str(d.get('reviewer', '')).strip(), session['real_name'] or session['username'],
         record_time, str(d.get('remark', '')).strip(), now_str()))
    new_fid = cur.lastrowid
    conn.commit()
    conn.close()
    log(session['username'], '备件' + ('入库' if ftype == 'in' else '出库'), f"item={item_id} qty={qty}")
    return jsonify({'ok': True, 'id': new_fid})


@app.route('/api/flows/<int:fid>', methods=['DELETE'])
@require_login
def delete_flow(fid):
    """删除流水并反向回补库存：入库→库位-数量；出库→库位+数量；移库→源库位+数量、目标库位-数量"""
    conn = get_conn()
    row = conn.execute("SELECT * FROM spare_flows WHERE id=?", (fid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '流水不存在'}), 404
    ftype, qty, item_id = row['type'], num(row['qty']), row['item_id']
    warehouse, location = row['warehouse'], row['location']
    restored = False
    if ftype == 'in':
        stock = conn.execute("SELECT * FROM spare_stocks WHERE item_id=? AND warehouse=? AND location=? LIMIT 1",
                             (item_id, warehouse, location)).fetchone()
        if not stock or num(stock['qty']) < qty:
            conn.close()
            return jsonify({'ok': False, 'error': '无法删除：对应库位库存不足（当前 %s），请先调整库存' % (stock['qty'] if stock else 0)})
        conn.execute("UPDATE spare_stocks SET qty=qty-? WHERE id=?", (qty, stock['id']))
        restored = True
    elif ftype == 'out':
        stock = conn.execute("SELECT * FROM spare_stocks WHERE item_id=? AND warehouse=? AND location=? LIMIT 1",
                             (item_id, warehouse, location)).fetchone()
        if stock:
            conn.execute("UPDATE spare_stocks SET qty=qty+? WHERE id=?", (qty, stock['id']))
            restored = True
        else:
            conn.close()
            return jsonify({'ok': False, 'error': '无法删除：对应库位已不存在，请先确认库存状态'})
    elif ftype == 'move':
        # 流水 remark 形如「移库：源库房/源库位 → 目标库房/目标库位」
        src_desc = ''
        rm = str(row['remark'] or '')
        if '移库：' in rm and '→' in rm:
            src_desc = rm.split('移库：', 1)[1].split('→', 1)[0].strip()
        src_wh, src_loc = '', ''
        if '/' in src_desc:
            src_wh, src_loc = src_desc.split('/', 1)
        src_wh, src_loc = src_wh.strip(), src_loc.strip()
        if not src_wh or not src_loc:
            conn.close()
            return jsonify({'ok': False, 'error': '无法删除移库流水：源库位信息缺失'})
        src = conn.execute("SELECT * FROM spare_stocks WHERE item_id=? AND warehouse=? AND location=? LIMIT 1",
                           (item_id, src_wh, src_loc)).fetchone()
        dst = conn.execute("SELECT * FROM spare_stocks WHERE item_id=? AND warehouse=? AND location=? LIMIT 1",
                           (item_id, warehouse, location)).fetchone()
        if not src:
            conn.close()
            return jsonify({'ok': False, 'error': '无法删除：源库位「%s/%s」已不存在' % (src_wh, src_loc)})
        if not dst or num(dst['qty']) < qty:
            conn.close()
            return jsonify({'ok': False, 'error': '无法删除：目标库位库存不足（当前 %s）' % (dst['qty'] if dst else 0)})
        conn.execute("UPDATE spare_stocks SET qty=qty+? WHERE id=?", (qty, src['id']))
        conn.execute("UPDATE spare_stocks SET qty=qty-? WHERE id=?", (qty, dst['id']))
        restored = True
    else:
        conn.close()
        return jsonify({'ok': False, 'error': '未知流水类型，无法回补'})
    conn.execute("DELETE FROM spare_flows WHERE id=?", (fid,))
    # 写变更日志到 spare_adjusts
    conn.execute(
        "INSERT INTO spare_adjusts(type,item_id,brand,name,spec,unit,warehouse,location,qty,reason,operator,adjust_time,created_at) "
        "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (ftype, item_id, row['brand'], row['name'], row['spec'], row['unit'],
         warehouse, location, -qty,
         '删除%s流水并回补库存' % ('入库' if ftype == 'in' else '出库' if ftype == 'out' else '移库'),
         session['real_name'] or session['username'], row['record_time'] or now_str()[:10], now_str()))
    conn.commit()
    conn.close()
    log(session['username'], '删除备件流水', f"{fid} type={ftype} qty={qty} 回补={restored}")
    return jsonify({'ok': True, 'restored': restored})


# ---------------- 备件出入库纸质单生成（按模板填充 docx，用于打印） ----------------
_SLIP_TPL = {
    'in': os.path.join(os.path.dirname(os.path.abspath(__file__)), 'templates', '入库单模板.docx'),
    'out': os.path.join(os.path.dirname(os.path.abspath(__file__)), 'templates', '出库单模板.docx'),
}


def _slip_cell(cell, text):
    """保留单元格首 Run 格式地替换文本（模板占位文本被真实值替换）"""
    p = cell.paragraphs[0]
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ''
    else:
        p.add_run(text)


def _slip_uniq(cells):
    """合并单元格在 row.cells 中按网格列展开为重复引用，按对象去重取物理单元格"""
    seen, out = set(), []
    for c in cells:
        if id(c) not in seen:
            seen.add(id(c))
            out.append(c)
    return out


@app.route('/api/flows/<int:fid>/slip', methods=['GET'])
@require_login
def flow_slip(fid):
    """按备件出入库流水记录生成纸质出入库单（docx），用于打印
    入库单：入库人=经办人、入库时间=记录时间、明细=品牌/名称/规格/数量/位置/备注、审核人
    出库单：使用人=领用人、流程编号=备注、领用时间=记录时间、明细同上、审核人
    """
    conn = get_conn()
    row = conn.execute("SELECT * FROM spare_flows WHERE id=?", (fid,)).fetchone()
    conn.close()
    if not row:
        return jsonify({'ok': False, 'error': '流水不存在'}), 404
    ftype = str(row['type'])
    if ftype not in ('in', 'out'):
        return jsonify({'ok': False, 'error': '仅入库/出库流水可生成纸质单'})
    tpl = _SLIP_TPL[ftype]
    if not os.path.exists(tpl):
        return jsonify({'ok': False, 'error': '纸质单模板缺失：%s' % os.path.basename(tpl)}), 500
    try:
        from docx import Document
    except ImportError:
        return jsonify({'ok': False, 'error': '服务端未安装 python-docx，无法生成纸质单'}), 500
    doc = Document(tpl)
    if not doc.tables:
        return jsonify({'ok': False, 'error': '纸质单模板无表格，无法填充'}), 500
    table = doc.tables[0]
    # 信息行（标题之下第 1 行；合并单元格按物理格处理）
    info = _slip_uniq(table.rows[1].cells)
    date = str(row['record_time'] or '')[:10] or datetime.datetime.now().strftime('%Y-%m-%d')
    qty = row['qty']
    qty_s = str(int(qty)) if isinstance(qty, float) and qty.is_integer() else str(qty)
    if ftype == 'in':
        if len(info) >= 2:
            _slip_cell(info[0], '入库人：' + str(row['operator'] or '').strip())
            _slip_cell(info[1], '入库时间：' + date)
    else:
        if len(info) >= 3:
            _slip_cell(info[0], '使用人：' + str(row['applicant'] or '').strip())
            _slip_cell(info[1], '流程编号：' + str(row['remark'] or '').strip())
            _slip_cell(info[2], '领用时间：' + date)
    # 明细：填写第 1 个数据行（表头之下第 1 行；6 列无合并，直接按网格格填写）
    if len(table.rows) >= 4:
        drow = table.rows[3].cells
        pos = (str(row['warehouse'] or '') + '/' + str(row['location'] or '')).strip('/')
        vals = [str(row['brand'] or '').strip(), str(row['name'] or '').strip(),
                str(row['spec'] or '').strip(), qty_s, pos, '']
        for i, v in enumerate(vals):
            if i < len(drow):
                _slip_cell(drow[i], v)
    # 审核人 / 库管（表格外正文）
    reviewer = str(row['reviewer'] or '').strip()
    for p in doc.paragraphs:
        if '审核人：' in p.text and '库管' in p.text:
            if p.runs:
                p.runs[0].text = '审核人：' + (reviewer or '') + '                                库管：'
                for r in p.runs[1:]:
                    r.text = ''
            else:
                p.text = '审核人：' + (reviewer or '') + '                                库管：'
            break
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    if ftype == 'in':
        fname = '入库单_%s_%s.docx' % (str(row['name'] or '备件'), date)
    else:
        fname = '出库单_%s_%s.docx' % (str(row['applicant'] or row['name'] or '备件'), date)
    log(session['username'], '生成备件%s纸质单' % ('入库' if ftype == 'in' else '出库'), 'flow=%d' % fid)
    return send_file(buf, as_attachment=True, download_name=fname,
                     mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')


@app.route('/api/flows/slip-batch', methods=['GET'])
@require_login
def flow_slip_batch():
    """批量出入库流水合成一张纸质单（docx）：?type=in|out&ids=1,2,3
    信息行取第一条记录，明细按记录逐行填充（复制模板数据行格式插入）"""
    ftype = request.args.get('type', '').strip()
    ids = _ids_arg()
    if ftype not in ('in', 'out'):
        return jsonify({'ok': False, 'error': 'type 必须为 in 或 out'})
    if not ids:
        return jsonify({'ok': False, 'error': '请选择要打印的流水'})
    conn = get_conn()
    ph = ','.join('?' * len(ids))
    rows = conn.execute("SELECT * FROM spare_flows WHERE id IN (%s) AND type=? ORDER BY record_time ASC, id ASC" % ph, list(ids) + [ftype]).fetchall()
    conn.close()
    if not rows:
        return jsonify({'ok': False, 'error': '未找到符合条件的流水'}), 404
    tpl = _SLIP_TPL[ftype]
    if not os.path.exists(tpl):
        return jsonify({'ok': False, 'error': '纸质单模板缺失：%s' % os.path.basename(tpl)}), 500
    try:
        from docx import Document
    except ImportError:
        return jsonify({'ok': False, 'error': '服务端未安装 python-docx，无法生成纸质单'}), 500
    doc = Document(tpl)
    if not doc.tables:
        return jsonify({'ok': False, 'error': '纸质单模板无表格，无法填充'}), 500
    table = doc.tables[0]
    # 信息行（标题之下第 1 行）：取第一条记录
    info = _slip_uniq(table.rows[1].cells)
    first = rows[0]
    date = str(first['record_time'] or '')[:10] or datetime.datetime.now().strftime('%Y-%m-%d')
    if ftype == 'in':
        if len(info) >= 2:
            _slip_cell(info[0], '入库人：' + str(first['operator'] or '').strip())
            _slip_cell(info[1], '入库时间：' + date)
    else:
        if len(info) >= 3:
            _slip_cell(info[0], '使用人：' + str(first['applicant'] or '').strip())
            _slip_cell(info[1], '流程编号：' + str(first['remark'] or '').strip())
            _slip_cell(info[2], '领用时间：' + date)
    # 明细：第 1 条填模板数据行，其余复制该行格式依次插入（保持边框样式）
    if len(table.rows) >= 4:
        from copy import deepcopy

        def _fill_row_cells(cells, r):
            qty = r['qty']
            qty_s = str(int(qty)) if isinstance(qty, float) and qty.is_integer() else str(qty)
            pos = (str(r['warehouse'] or '') + '/' + str(r['location'] or '')).strip('/')
            vals = [str(r['brand'] or '').strip(), str(r['name'] or '').strip(),
                    str(r['spec'] or '').strip(), qty_s, pos, '']
            for i, v in enumerate(vals):
                if i < len(cells):
                    _slip_cell(cells[i], v)

        base_tr = table.rows[3]._tr
        _fill_row_cells(table.rows[3].cells, first)
        anchor = base_tr
        for i, r in enumerate(rows[1:], start=1):
            tr = deepcopy(base_tr)
            anchor.addnext(tr)
            anchor = tr
            _fill_row_cells(table.rows[3 + i].cells, r)
    # 审核人 / 库管（表格外正文）：取第一条记录
    reviewer = str(first['reviewer'] or '').strip()
    for p in doc.paragraphs:
        if '审核人：' in p.text and '库管' in p.text:
            if p.runs:
                p.runs[0].text = '审核人：' + (reviewer or '') + '                                库管：'
                for rr in p.runs[1:]:
                    rr.text = ''
            else:
                p.text = '审核人：' + (reviewer or '') + '                                库管：'
            break
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    if ftype == 'in':
        fname = '入库单_批量%d项_%s.docx' % (len(rows), date)
    else:
        fname = '出库单_批量%d项_%s.docx' % (len(rows), date)
    log(session['username'], '生成备件批量%s纸质单' % ('入库' if ftype == 'in' else '出库'), 'ids=%s' % ','.join(map(str, ids)))
    return send_file(buf, as_attachment=True, download_name=fname,
                     mimetype='application/vnd.openxmlformats-officedocument.wordprocessingml.document')


# ---------------- 备件批量出入库 / 移库 / 流水导出 / 盘库 ----------------
@app.route('/api/stocks/batch', methods=['POST'])
@require_login
def batch_stock_flow():
    """批量出入库：一次登记多种物资，事务内逐项校验并更新库存、生成流水；任一失败整体回滚"""
    d = request.get_json(force=True, silent=True) or {}
    ftype = d.get('type', 'in')
    if ftype not in ('in', 'out'):
        return jsonify({'ok': False, 'error': '类型必须为入库或出库'})
    items = d.get('items') or []
    if not items:
        return jsonify({'ok': False, 'error': '请至少填写一项物资'})
    conn = get_conn()
    prepared = []
    for it in items:
        qty = num(it.get('qty'))
        if qty <= 0:
            conn.close()
            return jsonify({'ok': False, 'error': '存在数量无效的行，请检查'})
        item_id = it.get('item_id')
        if not item_id:
            conn.close()
            return jsonify({'ok': False, 'error': '存在未选择物资的行，请检查'})
        item = conn.execute("SELECT * FROM spare_items WHERE id=?", (item_id,)).fetchone()
        if not item:
            conn.close()
            return jsonify({'ok': False, 'error': '物资不存在: id=%s' % item_id})
        warehouse = str(it.get('warehouse', '')).strip()
        location = str(it.get('location', '')).strip()
        if not warehouse or not location:
            conn.close()
            return jsonify({'ok': False, 'error': '库房与库位必填'})
        stock = conn.execute(
            "SELECT * FROM spare_stocks WHERE item_id=? AND warehouse=? AND location=? LIMIT 1",
            (item_id, warehouse, location)).fetchone()
        if ftype == 'out':
            if not stock or num(stock['qty']) < qty:
                conn.close()
                return jsonify({'ok': False, 'error': '【%s %s @ %s/%s】库存不足（当前 %s）' % (
                    item['brand'], item['name'], warehouse, location,
                    stock['qty'] if stock else 0)})
        prepared.append((it, item, stock))
    record_time = str(d.get('record_time', '')).strip() or now_str()[:10]
    for it, item, stock in prepared:
        qty = num(it.get('qty'))
        region = str(it.get('region', '')).strip()
        owner = str(it.get('owner', '')).strip()
        warehouse = str(it.get('warehouse', '')).strip()
        location = str(it.get('location', '')).strip()
        if ftype == 'in':
            if stock:
                conn.execute("UPDATE spare_stocks SET qty=qty+?, region=?, owner=? WHERE id=?",
                             (qty, region or stock['region'], owner or stock['owner'], stock['id']))
            else:
                conn.execute(
                    "INSERT INTO spare_stocks(item_id,region,warehouse,location,qty,owner,remark,low_threshold) VALUES(?,?,?,?,?,?,?,0)",
                    (item['id'], region, warehouse, location, qty, owner, ''))
        else:
            conn.execute("UPDATE spare_stocks SET qty=qty-? WHERE id=?", (qty, stock['id']))
        conn.execute(
            "INSERT INTO spare_flows(type,item_id,brand,name,spec,unit,qty,region,warehouse,location,applicant,reviewer,operator,record_time,remark,created_at) "
            "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            (ftype, item['id'], item['brand'], item['name'], item['spec'], item['unit'], qty,
             region, warehouse, location, str(d.get('applicant', '')).strip(),
             str(d.get('reviewer', '')).strip(), session['real_name'] or session['username'],
             record_time, str(d.get('remark', '')).strip(), now_str()))
    conn.commit()
    conn.close()
    log(session['username'], '批量' + ('入库' if ftype == 'in' else '出库'), f"items={len(prepared)}")
    return jsonify({'ok': True, 'count': len(prepared)})


@app.route('/api/stocks/flow-template', methods=['GET'])
@require_login
def stock_flow_template():
    """下载备件批量出入库导入模板（xlsx）"""
    import openpyxl, io as _io
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '备件出入库导入模板'
    headers = ['出入库类型*', '品牌', '物资名称*', '规格型号', '数量*', '单位', '库房*', '库位*', '领用人', '审核人', '出入库时间', '备注']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='D9E2F3')
        cell.alignment = Alignment(horizontal='center')
    ws.append(['入库', '惠普', '原装硒鼓', 'HP 2612A', 5, '个', '一号库', 'A-01', '', '', '2026-09-25', '采购入库'])
    ws.append(['出库', '佳能', '定影组件', 'Canon 475', 1, '套', '一号库', 'A-02', '张三', '李四', '', '领用'])
    widths = [12, 10, 18, 14, 8, 8, 12, 10, 10, 10, 14, 18]
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='备件批量出入库导入模板.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/api/stocks/flow-import', methods=['POST'])
@require_login
def import_stock_flows():
    """批量导入备件出入库：逐行校验（类型/物资/库房库位/数量/库存），全部通过后事务执行并生成流水"""
    f = request.files.get('file')
    if not f or not f.filename:
        return jsonify({'ok': False, 'error': '请选择要导入的 Excel 文件'})
    if not (f.filename.endswith('.xlsx') or f.filename.endswith('.xls')):
        return jsonify({'ok': False, 'error': '仅支持 .xlsx / .xls 文件'})
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(f.read()), data_only=True)
    except Exception as e:
        return jsonify({'ok': False, 'error': '文件解析失败：%s' % e})
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    conn = get_conn()
    today = now_str()[:10]
    prepared = []
    errors = []
    adjust = {}  # (item_id, warehouse, location) -> 本批次已累计的净变更（入库+ / 出库-）
    for idx, row in enumerate(rows, start=2):
        if row is None or all(v is None or str(v).strip() == '' for v in row):
            continue
        ftype_raw = str(row[0] or '').strip() if len(row) > 0 else ''
        ftype = 'in' if ftype_raw in ('入库', 'in', 'IN') else ('out' if ftype_raw in ('出库', 'out', 'OUT') else '')
        brand = str(row[1] or '').strip() if len(row) > 1 and row[1] is not None else ''
        name = str(row[2] or '').strip() if len(row) > 2 and row[2] is not None else ''
        spec = str(row[3] or '').strip() if len(row) > 3 and row[3] is not None else ''
        try:
            qty = num(row[4]) if len(row) > 4 and row[4] is not None else 0
        except (TypeError, ValueError):
            qty = 0
        unit = str(row[5] or '').strip() if len(row) > 5 and row[5] is not None else ''
        warehouse = str(row[6] or '').strip() if len(row) > 6 and row[6] is not None else ''
        location = str(row[7] or '').strip() if len(row) > 7 and row[7] is not None else ''
        applicant = str(row[8] or '').strip() if len(row) > 8 and row[8] is not None else ''
        reviewer = str(row[9] or '').strip() if len(row) > 9 and row[9] is not None else ''
        record_time = str(row[10] or '').strip() if len(row) > 10 and row[10] is not None else ''
        remark = str(row[11] or '').strip() if len(row) > 11 and row[11] is not None else ''
        if not ftype:
            errors.append(f"第{idx}行：出入库类型必须为「入库」或「出库」")
            continue
        if not name:
            errors.append(f"第{idx}行：物资名称必填")
            continue
        if not warehouse or not location:
            errors.append(f"第{idx}行：库房与库位必填")
            continue
        if not qty or qty <= 0:
            errors.append(f"第{idx}行：数量无效（应为正数）")
            continue
        if record_time:
            try:
                datetime.datetime.strptime(record_time[:10], '%Y-%m-%d')
            except ValueError:
                errors.append(f"第{idx}行：出入库时间格式应为 YYYY-MM-DD")
                continue
        # 物资匹配：按 名称+品牌+规格+单位；匹配不到时入库可自动新建，出库则拒绝
        item = conn.execute(
            "SELECT * FROM spare_items WHERE name=? AND brand=? AND spec=? AND unit=? LIMIT 1",
            (name, brand, spec, unit)).fetchone()
        if not item:
            if ftype == 'out':
                errors.append(f"第{idx}行：物资库中无「{brand} {name} {spec}」，出库请先在物资管理中新增")
                continue
            cur = conn.execute("INSERT INTO spare_items(brand,name,spec,unit,remark) VALUES(?,?,?,?,?)",
                               (brand, name, spec, unit, '模板导入自动创建'))
            item = {'id': cur.lastrowid, 'brand': brand, 'name': name, 'spec': spec, 'unit': unit}
        # 库位匹配
        stock = conn.execute(
            "SELECT * FROM spare_stocks WHERE item_id=? AND warehouse=? AND location=? LIMIT 1",
            (item['id'], warehouse, location)).fetchone()
        if ftype == 'out':
            key = (item['id'], warehouse, location)
            eff = num(stock['qty'] if stock else 0) + adjust.get(key, 0)
            if eff < qty:
                errors.append(f"第{idx}行：【{brand} {name} @ {warehouse}/{location}】库存不足（当前 {eff}）")
                continue
        prepared.append((idx, ftype, item, stock, qty, brand, name, spec, unit, warehouse, location,
                         applicant, reviewer, record_time or today, remark))
        adjust[(item['id'], warehouse, location)] = adjust.get((item['id'], warehouse, location), 0) + (
            qty if ftype == 'in' else -qty)
    if errors:
        conn.close()
        return jsonify({'ok': False, 'error': '存在 %d 行校验失败，已整体取消导入' % len(errors),
                        'errors': errors[:20]})
    if not prepared:
        conn.close()
        return jsonify({'ok': False, 'error': '模板中没有有效数据行'})
    # 全部通过：事务执行（执行期再做一次库存防御校验）
    try:
        for idx, ftype, item, stock, qty, brand, name, spec, unit, warehouse, location, applicant, reviewer, record_time, remark in prepared:
            # 执行时实时重查库位（校验阶段的缓存可能已被本批次前序行改变）
            cur = conn.execute(
                "SELECT * FROM spare_stocks WHERE item_id=? AND warehouse=? AND location=? LIMIT 1",
                (item['id'], warehouse, location)).fetchone()
            if ftype == 'in':
                if cur:
                    conn.execute("UPDATE spare_stocks SET qty=qty+? WHERE id=?", (qty, cur['id']))
                else:
                    conn.execute(
                        "INSERT INTO spare_stocks(item_id,region,warehouse,location,qty,owner,remark,low_threshold) VALUES(?,?,?,?,?,?,?,0)",
                        (item['id'], '', warehouse, location, qty, '', ''))
            else:
                if not cur or num(cur['qty']) < qty:
                    raise Exception(f"第{idx}行：【{brand} {name}】执行时库存不足，已回滚")
                conn.execute("UPDATE spare_stocks SET qty=qty-? WHERE id=?", (qty, cur['id']))
            conn.execute(
                "INSERT INTO spare_flows(type,item_id,brand,name,spec,unit,qty,region,warehouse,location,applicant,reviewer,operator,record_time,remark,created_at) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (ftype, item['id'], item['brand'], item['name'], item['spec'], item['unit'], qty,
                 '', warehouse, location, applicant, reviewer, session['real_name'] or session['username'],
                 record_time, remark, now_str()))
        conn.commit()
    except Exception as e:
        conn.rollback()
        conn.close()
        return jsonify({'ok': False, 'error': str(e)})
    conn.close()
    log(session['username'], '模板导入备件出入库', f"rows={len(prepared)}")
    return jsonify({'ok': True, 'added': len(prepared), 'failed': 0, 'errors': []})


@app.route('/api/stocks/move', methods=['POST'])
@require_login
def move_stock():
    """物资移库位：从源库存位移动到目标库房/库位，自动更新两处库存并生成移库流水"""
    d = request.get_json(force=True, silent=True) or {}
    sid = d.get('stock_id')
    qty = num(d.get('qty'))
    to_warehouse = str(d.get('to_warehouse', '')).strip()
    to_location = str(d.get('to_location', '')).strip()
    if not sid or qty <= 0:
        return jsonify({'ok': False, 'error': '请选择源库存位并填写正确的移库数量'})
    if not to_warehouse or not to_location:
        return jsonify({'ok': False, 'error': '目标库房与库位必填'})
    conn = get_conn()
    src = conn.execute("SELECT s.*, i.brand, i.name, i.spec, i.unit FROM spare_stocks s "
                       "LEFT JOIN spare_items i ON i.id = s.item_id WHERE s.id=?", (sid,)).fetchone()
    if not src:
        conn.close()
        return jsonify({'ok': False, 'error': '源库存位不存在'}), 404
    if num(src['qty']) < qty:
        conn.close()
        return jsonify({'ok': False, 'error': '源库位库存不足（当前 %s）' % src['qty']})
    # 目标库位：同物资+目标库房库位，存在则累加，否则新建
    dst = conn.execute(
        "SELECT * FROM spare_stocks WHERE item_id=? AND warehouse=? AND location=? LIMIT 1",
        (src['item_id'], to_warehouse, to_location)).fetchone()
    if dst:
        conn.execute("UPDATE spare_stocks SET qty=qty+? WHERE id=?", (qty, dst['id']))
    else:
        conn.execute(
            "INSERT INTO spare_stocks(item_id,region,warehouse,location,qty,owner,remark,low_threshold) VALUES(?,?,?,?,?,?,?,0)",
            (src['item_id'], str(d.get('to_region', '')).strip() or src['region'],
             to_warehouse, to_location, qty, src['owner'], ''))
    conn.execute("UPDATE spare_stocks SET qty=qty-? WHERE id=?", (qty, src['id']))
    record_time = str(d.get('record_time', '')).strip() or now_str()[:10]
    conn.execute(
        "INSERT INTO spare_flows(type,item_id,brand,name,spec,unit,qty,region,warehouse,location,applicant,reviewer,operator,record_time,remark,created_at) "
        "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        ('move', src['item_id'], src['brand'], src['name'], src['spec'], src['unit'], qty,
         str(d.get('to_region', '')).strip() or src['region'], to_warehouse, to_location,
         str(d.get('applicant', '')).strip(), '', session['real_name'] or session['username'],
         record_time, f"移库：{src['warehouse']}/{src['location']} → {to_warehouse}/{to_location}", now_str()))
    conn.commit()
    conn.close()
    log(session['username'], '物资移库', f"stock={sid} qty={qty} to={to_warehouse}/{to_location}")
    return jsonify({'ok': True})


@app.route('/api/spare/adjusts', methods=['GET'])
@require_login
def list_spare_adjusts():
    """备件库存变更日志：手动调整 + 出入库记录改删 + 库存改删"""
    q = request.args.get('q', '').strip()
    page = max(1, int(request.args.get('page', 1)))
    per = min(max(1, int(request.args.get('per', 30))), 100000)
    sql = "SELECT * FROM spare_adjusts WHERE 1=1"
    args = []
    if q:
        sql += " AND (brand LIKE ? OR name LIKE ? OR spec LIKE ? OR reason LIKE ? OR operator LIKE ?)"
        like = '%' + q + '%'
        args += [like] * 5
    total = get_conn().execute(sql.replace("SELECT *", "SELECT COUNT(*) c"), args).fetchone()['c']
    sql += " ORDER BY id DESC LIMIT ? OFFSET ?"
    args += [per, (page - 1) * per]
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows], 'total': total, 'page': page, 'per': per})


@app.route('/api/spare/adjusts', methods=['POST'])
@require_login
def add_spare_adjust():
    """新增备件库存调整：调整类型 in=增加 / out=减少，事务内更新库存并记录"""
    d = request.get_json(force=True, silent=True) or {}
    atype = d.get('type', 'in')
    if atype not in ('in', 'out'):
        return jsonify({'ok': False, 'error': '调整类型必须为 增加/减少'})
    qty = num(d.get('qty'))
    if qty <= 0:
        return jsonify({'ok': False, 'error': '数量必须大于 0'})
    warehouse = str(d.get('warehouse', '')).strip()
    location = str(d.get('location', '')).strip()
    item_id = d.get('item_id')
    conn = get_conn()
    if item_id:
        item = conn.execute("SELECT * FROM spare_items WHERE id=?", (item_id,)).fetchone()
        if not item:
            conn.close()
            return jsonify({'ok': False, 'error': '物资不存在'}), 404
    else:
        name = str(d.get('name', '')).strip()
        if not name:
            conn.close()
            return jsonify({'ok': False, 'error': '请选择物资'})
        item = conn.execute("SELECT * FROM spare_items WHERE name=? AND brand=? AND spec=? AND unit=? LIMIT 1",
                            (name, str(d.get('brand', '')).strip(), str(d.get('spec', '')).strip(),
                             str(d.get('unit', '')).strip())).fetchone()
        if not item:
            conn.close()
            return jsonify({'ok': False, 'error': '物资库中无「%s」，请先在物资管理中新增' % name})
        item_id = item['id']
    stock = conn.execute("SELECT * FROM spare_stocks WHERE item_id=? AND warehouse=? AND location=? LIMIT 1",
                         (item_id, warehouse, location)).fetchone()
    if not stock:
        conn.close()
        return jsonify({'ok': False, 'error': '库位「%s/%s」不存在，请先登记库存' % (warehouse, location)})
    if atype == 'out':
        if num(stock['qty']) < qty:
            conn.close()
            return jsonify({'ok': False, 'error': '库存不足！当前库存 %s' % stock['qty']})
        conn.execute("UPDATE spare_stocks SET qty=qty-? WHERE id=?", (qty, stock['id']))
    else:
        conn.execute("UPDATE spare_stocks SET qty=qty+? WHERE id=?", (qty, stock['id']))
    cur = conn.execute(
        "INSERT INTO spare_adjusts(type,item_id,brand,name,spec,unit,warehouse,location,qty,reason,operator,adjust_time,created_at) "
        "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (atype, item_id, item['brand'], item['name'], item['spec'], item['unit'],
         warehouse, location, qty, str(d.get('reason', '')).strip(),
         session['real_name'] or session['username'], d.get('adjust_time') or now_str()[:10], now_str()))
    conn.commit()
    conn.close()
    log(session['username'], '备件库存调整', 'id=%s %s qty=%s' % (cur.lastrowid, atype, qty))
    return jsonify({'ok': True, 'id': cur.lastrowid})


@app.route('/api/spare/adjusts/<int:aid>', methods=['DELETE'])
@require_login
def delete_spare_adjust(aid):
    """删除调整记录并反向回补库存：增加调整→库存-数量；减少调整→库存+数量"""
    conn = get_conn()
    row = conn.execute("SELECT * FROM spare_adjusts WHERE id=?", (aid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '调整记录不存在'}), 404
    qty, atype = num(row['qty']), row['type']
    stock = conn.execute(
        "SELECT * FROM spare_stocks WHERE item_id=? AND warehouse=? AND location=? LIMIT 1",
        (row['item_id'], row['warehouse'], row['location'])).fetchone()
    if atype == 'in':
        if not stock or num(stock['qty']) < qty:
            conn.close()
            return jsonify({'ok': False, 'error': '无法删除：当前库存不足（%s），请先调整库存' % (stock['qty'] if stock else 0)})
        conn.execute("UPDATE spare_stocks SET qty=qty-? WHERE id=?", (qty, stock['id']))
    else:
        if not stock:
            conn.close()
            return jsonify({'ok': False, 'error': '无法删除：对应库位已不存在'})
        conn.execute("UPDATE spare_stocks SET qty=qty+? WHERE id=?", (qty, stock['id']))
    conn.execute("DELETE FROM spare_adjusts WHERE id=?", (aid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除备件调整记录', '%s qty=%s' % (aid, qty))
    return jsonify({'ok': True})


@app.route('/api/flows/export', methods=['GET'])
@require_login
def export_flows():
    """出入库流水导出：明细 Sheet + 月度汇总 Sheet"""
    q = request.args.get('q', '').strip()
    ftype = request.args.get('type', '').strip()
    start = request.args.get('start', '').strip()
    end = request.args.get('end', '').strip()
    sql = "SELECT * FROM spare_flows WHERE 1=1"
    args = []
    if q:
        sql += " AND (name LIKE ? OR spec LIKE ? OR brand LIKE ? OR applicant LIKE ? OR remark LIKE ?)"
        like = f"%{q}%"
        args += [like] * 5
    if ftype:
        sql += " AND type=?"
        args.append(ftype)
    if start:
        sql += " AND record_time >= ?"
        args.append(start)
    if end:
        sql += " AND record_time <= ?"
        args.append(end)
    ids = _ids_arg()
    if ids:
        sql += " AND id IN (%s)" % ','.join('?' * len(ids))
        args += ids
    sql += " ORDER BY record_time DESC, id DESC"
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    # 月度汇总（同筛选条件下按月×类型统计）
    sum_sql = ("SELECT substr(record_time,1,7) ym, type, COUNT(*) c, COALESCE(SUM(qty),0) q "
               "FROM spare_flows WHERE 1=1")
    sum_args = []
    if q:
        sum_sql += " AND (name LIKE ? OR spec LIKE ? OR brand LIKE ? OR applicant LIKE ? OR remark LIKE ?)"
        sum_args += [like] * 5
    if ftype:
        sum_sql += " AND type=?"
        sum_args.append(ftype)
    if start:
        sum_sql += " AND record_time >= ?"
        sum_args.append(start)
    if end:
        sum_sql += " AND record_time <= ?"
        sum_args.append(end)
    ids = _ids_arg()
    if ids:
        sum_sql += " AND id IN (%s)" % ','.join('?' * len(ids))
        sum_args += ids
    sum_sql += " GROUP BY ym, type ORDER BY ym"
    sums = conn.execute(sum_sql, sum_args).fetchall()
    conn.close()
    import openpyxl, io as _io
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '流水明细'
    headers = ['类型', '时间', '品牌', '物资名称', '规格型号', '数量', '库房', '库位', '领用人', '审核人', '经办人', '备注']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='D9E2F3')
        cell.alignment = Alignment(horizontal='center')
    for r in rows:
        ws.append([('入库' if r['type'] == 'in' else ('出库' if r['type'] == 'out' else '移库')),
                   r['record_time'], r['brand'], r['name'], r['spec'], num(r['qty']),
                   r['warehouse'], r['location'], r['applicant'], r['reviewer'], r['operator'], r['remark']])
    ws2 = wb.create_sheet('月度汇总')
    ws2.append(['月份', '类型', '笔数', '数量合计'])
    for cell in ws2[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='D9E2F3')
    for r in sums:
        ws2.append([r['ym'], ('入库' if r['type'] == 'in' else ('出库' if r['type'] == 'out' else '移库')),
                    r['c'], num(r['q'])])
    ws2.append(['合计', '', len(rows), sum(num(r['qty']) for r in rows)])
    ws2.cell(row=ws2.max_row, column=4).font = Font(bold=True)
    for wsx, widths in ((ws, [8, 12, 12, 18, 14, 8, 12, 14, 10, 10, 10, 24]),
                        (ws2, [10, 8, 8, 12])):
        from openpyxl.utils import get_column_letter
        for i, w in enumerate(widths, 1):
            wsx.column_dimensions[get_column_letter(i)].width = w
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = '备件出入库流水_%s.xlsx' % datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    log(session['username'], '导出出入库流水', f"rows={len(rows)}")
    return send_file(buf, as_attachment=True, download_name=fname,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ---------------- 盘库 ----------------
@app.route('/api/stocktakes', methods=['GET'])
@require_login
def list_stocktakes():
    conn = get_conn()
    rows = conn.execute("SELECT * FROM stocktakes ORDER BY id DESC LIMIT 100").fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows]})


@app.route('/api/stocktakes', methods=['POST'])
@require_login
def create_stocktake():
    """新建盘点批次：按库房（可空=全部）生成账面明细"""
    d = request.get_json(force=True, silent=True) or {}
    warehouse = str(d.get('warehouse', '')).strip()
    remark = str(d.get('remark', '')).strip()
    conn = get_conn()
    take_time = str(d.get('take_time', '')).strip() or now_str()[:10]
    prefix = 'PD' + take_time[:4] + take_time[5:7]
    n = conn.execute("SELECT COUNT(*) c FROM stocktakes WHERE batch_no LIKE ?", (prefix + '%',)).fetchone()['c'] + 1
    batch_no = '%s%03d' % (prefix, n)
    cur = conn.execute(
        "INSERT INTO stocktakes(batch_no,take_time,operator,warehouse,status,remark,created_at) VALUES(?,?,?,?,?,?,?)",
        (batch_no, take_time, session['real_name'] or session['username'], warehouse, '进行中', remark, now_str()))
    bid = cur.lastrowid
    sql = ("SELECT s.id stock_id, s.item_id, s.warehouse, s.location, s.qty, "
           "i.brand, i.name, i.spec, i.unit FROM spare_stocks s "
           "LEFT JOIN spare_items i ON i.id = s.item_id WHERE 1=1")
    args = []
    if warehouse:
        sql += " AND s.warehouse=?"
        args.append(warehouse)
    sql += " ORDER BY s.warehouse, s.location, i.name"
    stocks = conn.execute(sql, args).fetchall()
    for st in stocks:
        conn.execute(
            "INSERT INTO stocktake_items(batch_id,stock_id,item_id,brand,name,spec,unit,warehouse,location,book_qty,actual_qty,diff_qty,status) "
            "VALUES(?,?,?,?,?,?,?,?,?,?,NULL,0,'待调整')",
            (bid, st['stock_id'], st['item_id'], st['brand'], st['name'], st['spec'], st['unit'],
             st['warehouse'], st['location'], num(st['qty'])))
    conn.commit()
    conn.close()
    log(session['username'], '新建盘点', f"{batch_no} items={len(stocks)}")
    return jsonify({'ok': True, 'id': bid, 'batch_no': batch_no, 'count': len(stocks)})


@app.route('/api/stocktakes/<int:tid>', methods=['GET'])
@require_login
def get_stocktake(tid):
    conn = get_conn()
    b = conn.execute("SELECT * FROM stocktakes WHERE id=?", (tid,)).fetchone()
    if not b:
        conn.close()
        return jsonify({'ok': False, 'error': '盘点批次不存在'}), 404
    items = conn.execute("SELECT * FROM stocktake_items WHERE batch_id=? ORDER BY warehouse, location, name", (tid,)).fetchall()
    conn.close()
    return jsonify({'ok': True, 'batch': dict(b), 'items': [dict(r) for r in items]})


@app.route('/api/stocktakes/<int:tid>', methods=['PUT'])
@require_login
def save_stocktake(tid):
    """保存实盘数：计算盈亏差异"""
    d = request.get_json(force=True, silent=True) or {}
    updates = d.get('items') or []
    conn = get_conn()
    b = conn.execute("SELECT * FROM stocktakes WHERE id=?", (tid,)).fetchone()
    if not b:
        conn.close()
        return jsonify({'ok': False, 'error': '盘点批次不存在'}), 404
    if b['status'] != '进行中':
        conn.close()
        return jsonify({'ok': False, 'error': '该批次已完成，不可修改'})
    for u in updates:
        iid = u.get('id')
        actual = num(u.get('actual_qty'))
        if actual < 0:
            conn.close()
            return jsonify({'ok': False, 'error': '实盘数不能为负'})
        row = conn.execute("SELECT * FROM stocktake_items WHERE id=? AND batch_id=?", (iid, tid)).fetchone()
        if row:
            diff = actual - num(row['book_qty'])
            conn.execute("UPDATE stocktake_items SET actual_qty=?, diff_qty=?, status=? WHERE id=?",
                         (actual, diff, '待调整' if diff != 0 else '一致', iid))
    conn.commit()
    conn.close()
    log(session['username'], '保存盘点实盘', f"batch={tid} rows={len(updates)}")
    return jsonify({'ok': True})


@app.route('/api/stocktakes/<int:tid>/apply', methods=['POST'])
@require_login
def apply_stocktake(tid):
    """确认调整（平库）：按盈亏差异自动更新库存并生成盘盈/盘亏流水，批次完成"""
    conn = get_conn()
    b = conn.execute("SELECT * FROM stocktakes WHERE id=?", (tid,)).fetchone()
    if not b:
        conn.close()
        return jsonify({'ok': False, 'error': '盘点批次不存在'}), 404
    if b['status'] != '进行中':
        conn.close()
        return jsonify({'ok': False, 'error': '该批次已完成，不可重复调整'})
    items = conn.execute("SELECT * FROM stocktake_items WHERE batch_id=? AND diff_qty != 0", (tid,)).fetchall()
    try:
        for it in items:
            diff = num(it['diff_qty'])
            # 更新库存（diff>0 盘盈增加，diff<0 盘亏减少）
            conn.execute("UPDATE spare_stocks SET qty=qty+? WHERE id=?", (diff, it['stock_id']))
            ftype = 'in' if diff > 0 else 'out'
            conn.execute(
                "INSERT INTO spare_flows(type,item_id,brand,name,spec,unit,qty,region,warehouse,location,applicant,reviewer,operator,record_time,remark,created_at) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (ftype, it['item_id'], it['brand'], it['name'], it['spec'], it['unit'], abs(diff),
                 '', it['warehouse'], it['location'], '', '',
                 session['real_name'] or session['username'], b['take_time'],
                 '盘盈调整' if diff > 0 else '盘亏调整', now_str()))
            conn.execute("UPDATE stocktake_items SET status=? WHERE id=?",
                         ('已盘盈' if diff > 0 else '已盘亏', it['id']))
        conn.execute("UPDATE stocktakes SET status='已完成' WHERE id=?", (tid,))
        conn.commit()
    except Exception as e:
        conn.rollback()
        conn.close()
        return jsonify({'ok': False, 'error': '调整失败：%s' % e})
    conn.close()
    log(session['username'], '盘库确认调整', f"batch={tid} diff_items={len(items)}")
    return jsonify({'ok': True, 'adjusted': len(items)})


@app.route('/api/stocktakes/<int:tid>', methods=['DELETE'])
@require_login
def delete_stocktake(tid):
    conn = get_conn()
    b = conn.execute("SELECT * FROM stocktakes WHERE id=?", (tid,)).fetchone()
    if not b:
        conn.close()
        return jsonify({'ok': False, 'error': '盘点批次不存在'}), 404
    if b['status'] != '进行中':
        conn.close()
        return jsonify({'ok': False, 'error': '已完成批次不可删除（可在流水与库存中查看结果）'})
    conn.execute("DELETE FROM stocktake_items WHERE batch_id=?", (tid,))
    conn.execute("DELETE FROM stocktakes WHERE id=?", (tid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除盘点批次', str(tid))
    return jsonify({'ok': True})


# ---------------- 办公电脑 ----------------
@app.route('/api/computers', methods=['GET'])
@require_login
def list_computers():
    q = request.args.get('q', '').strip()
    status = request.args.get('status', '').strip()
    sql = "SELECT * FROM computers WHERE 1=1"
    args = []
    if q:
        sql += " AND (asset_no LIKE ? OR hostname LIKE ? OR brand LIKE ? OR model LIKE ? OR user_name LIKE ? OR department LIKE ? OR ip LIKE ?)"
        like = f"%{q}%"
        args += [like] * 7
    if status:
        sql += " AND status=?"
        args.append(status)
    sql += " ORDER BY id DESC"
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows]})


@app.route('/api/computers/<int:cid>', methods=['GET'])
@require_login
def get_computer(cid):
    conn = get_conn()
    c = conn.execute("SELECT * FROM computers WHERE id=?", (cid,)).fetchone()
    repairs = conn.execute(
        "SELECT * FROM computer_repairs WHERE computer_id=? ORDER BY record_time DESC, id DESC", (cid,)).fetchall()
    conn.close()
    if not c:
        return jsonify({'ok': False, 'error': '电脑不存在'}), 404
    return jsonify({'ok': True, 'data': dict(c), 'repairs': [dict(r) for r in repairs]})


@app.route('/api/computers', methods=['POST'])
@require_login
def add_computer():
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    cur = conn.execute(
        "INSERT INTO computers(asset_no,hostname,brand,model,sn,config,user_name,department,location,ip,mac,"
        "purchase_date,warranty_until,status,remark,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (str(d.get('asset_no', '')).strip(), str(d.get('hostname', '')).strip(),
         str(d.get('brand', '')).strip(), str(d.get('model', '')).strip(),
         str(d.get('sn', '')).strip(), str(d.get('config', '')).strip(),
         str(d.get('user_name', '')).strip(), str(d.get('department', '')).strip(),
         str(d.get('location', '')).strip(), str(d.get('ip', '')).strip(),
         str(d.get('mac', '')).strip(), d.get('purchase_date') or '',
         d.get('warranty_until') or '', d.get('status') or '在用',
         str(d.get('remark', '')).strip(), now_str(), now_str()))
    conn.commit()
    conn.close()
    log(session['username'], '新增电脑', str(cur.lastrowid))
    return jsonify({'ok': True, 'id': cur.lastrowid})


@app.route('/api/computers/<int:cid>', methods=['PUT'])
@require_login
def update_computer(cid):
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    conn.execute(
        "UPDATE computers SET asset_no=?,hostname=?,brand=?,model=?,sn=?,config=?,user_name=?,department=?,location=?,ip=?,mac=?,"
        "purchase_date=?,warranty_until=?,status=?,remark=?,updated_at=? WHERE id=?",
        (str(d.get('asset_no', '')).strip(), str(d.get('hostname', '')).strip(),
         str(d.get('brand', '')).strip(), str(d.get('model', '')).strip(),
         str(d.get('sn', '')).strip(), str(d.get('config', '')).strip(),
         str(d.get('user_name', '')).strip(), str(d.get('department', '')).strip(),
         str(d.get('location', '')).strip(), str(d.get('ip', '')).strip(),
         str(d.get('mac', '')).strip(), d.get('purchase_date') or '',
         d.get('warranty_until') or '', d.get('status') or '在用',
         str(d.get('remark', '')).strip(), now_str(), cid))
    conn.commit()
    conn.close()
    log(session['username'], '修改电脑', str(cid))
    return jsonify({'ok': True})


@app.route('/api/computers/<int:cid>', methods=['DELETE'])
@require_login
def delete_computer(cid):
    conn = get_conn()
    conn.execute("DELETE FROM computer_repairs WHERE computer_id=?", (cid,))
    conn.execute("DELETE FROM computers WHERE id=?", (cid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除电脑', str(cid))
    return jsonify({'ok': True})


@app.route('/api/computer-repairs', methods=['GET'])
@require_login
def list_computer_repairs():
    cid = request.args.get('computer_id', '').strip()
    sql = ("SELECT r.*, c.hostname, c.user_name, c.department, c.model AS computer_model "
           "FROM computer_repairs r LEFT JOIN computers c ON c.id = r.computer_id WHERE 1=1")
    args = []
    if cid:
        sql += " AND r.computer_id=?"
        args.append(cid)
    sql += " ORDER BY r.record_time DESC, r.id DESC"
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows]})


@app.route('/api/computer-repairs', methods=['POST'])
@require_login
def add_computer_repair():
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    cur = conn.execute(
        "INSERT INTO computer_repairs(computer_id,type,detail,record_time,operator,cost,remark,created_at) "
        "VALUES(?,?,?,?,?,?,?,?)",
        (d.get('computer_id'), d.get('type') or '维修', str(d.get('detail', '')).strip(),
         d.get('record_time') or now_str()[:10], str(d.get('operator', '')).strip(),
         num(d.get('cost')), str(d.get('remark', '')).strip(), now_str()))
    conn.commit()
    conn.close()
    log(session['username'], '新增电脑维修/变更', str(cur.lastrowid))
    return jsonify({'ok': True, 'id': cur.lastrowid})


@app.route('/api/computer-repairs/<int:rid>', methods=['PUT'])
@require_login
def update_computer_repair(rid):
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    conn.execute(
        "UPDATE computer_repairs SET computer_id=?,type=?,detail=?,record_time=?,operator=?,cost=?,remark=? WHERE id=?",
        (d.get('computer_id'), d.get('type') or '维修', str(d.get('detail', '')).strip(),
         d.get('record_time') or now_str()[:10], str(d.get('operator', '')).strip(),
         num(d.get('cost')), str(d.get('remark', '')).strip(), rid))
    conn.commit()
    conn.close()
    log(session['username'], '更新电脑维修/变更', str(rid))
    return jsonify({'ok': True})


@app.route('/api/computer-repairs/<int:rid>', methods=['DELETE'])
@require_login
def delete_computer_repair(rid):
    conn = get_conn()
    conn.execute("DELETE FROM computer_repairs WHERE id=?", (rid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除电脑维修记录', str(rid))
    return jsonify({'ok': True})


# ---------------- 办公电脑优化记录（内存/硬盘升级） ----------------
PC_UP_HEADERS = ['申请人', '流程单号', '资产编号', '8G内存', '16G内存', 'sata固态', 'm2固态', '机械硬盘', '升级日期', '备注', '备件', '剩余数量', 'ZK', 'BG', '采购时间']


def _up_date(v):
    """模板日期可能是 datetime，统一为 yyyy-MM-dd"""
    if v is None:
        return ''
    if hasattr(v, 'strftime'):
        return v.strftime('%Y-%m-%d')
    s = str(v).strip()
    return s[:10] if len(s) >= 10 else s


# ---------------- 办公电脑优化记录（主从结构：主表 + 配件明细） ----------------
UPGRADE_TYPE_NAMES = {3: '8G内存', 4: '16G内存', 5: 'SATA固态', 6: 'M2固态', 7: '机械硬盘'}


def _pc_up_items(conn, up_id):
    rows = conn.execute("SELECT * FROM pc_upgrade_items WHERE upgrade_id=? ORDER BY id", (up_id,)).fetchall()
    return [dict(r) for r in rows]


def _pc_norm_items(raw):
    out = []
    for it in (raw or []):
        if not isinstance(it, dict):
            continue
        qty = int(num(it.get('qty', 1)))
        if qty < 1:
            qty = 1
        name = str(it.get('name', '')).strip()
        if not name and not it.get('part_id'):
            continue
        out.append({'part_id': it.get('part_id'), 'brand': str(it.get('brand', '')).strip(),
                    'name': name, 'spec': str(it.get('spec', '')).strip(), 'qty': qty})
    return out


def _pc_apply_items(conn, up_id, items, operator, rec_time):
    """扣减库存并写出库流水（仅对关联了库存备件的明细）"""
    for it in items:
        pid = it.get('part_id')
        qty = it.get('qty', 1)
        if not pid or qty < 1:
            continue
        part = conn.execute("SELECT * FROM pc_parts WHERE id=?", (pid,)).fetchone()
        if not part:
            raise ValueError('关联备件不存在')
        if int(num(part['stock'])) < qty:
            raise ValueError('备件「%s」库存不足（当前 %s），需要 %s，请先采购入库' % (part['name'], part['stock'], qty))
        conn.execute("UPDATE pc_parts SET stock=stock-? WHERE id=?", (qty, pid))
        conn.execute(
            "INSERT INTO pc_flows(type,part_id,part_name,qty,operator,ref,record_time,remark,created_at) VALUES('out',?,?,?,?,?,?,?,?)",
            (pid, part['name'], qty, operator, '优化记录#' + str(up_id), rec_time, '优化记录发放自动出库', now_str()))


def _pc_revert_items(conn, up_id, items):
    """回补库存并删除关联出库流水（删除记录或编辑旧明细时）"""
    for it in items:
        pid = it.get('part_id')
        qty = it.get('qty', 1)
        if not pid or qty < 1:
            continue
        conn.execute("UPDATE pc_parts SET stock=stock+? WHERE id=?", (qty, pid))
        conn.execute("DELETE FROM pc_flows WHERE ref=? AND type='out' AND part_id=?",
                     ('优化记录#' + str(up_id), pid))


@app.route('/api/pc-upgrades', methods=['GET'])
@require_login
def list_pc_upgrades():
    q = request.args.get('q', '').strip()
    sql = "SELECT * FROM pc_upgrades WHERE 1=1"
    args = []
    if q:
        sql += " AND (applicant LIKE ? OR flow_no LIKE ? OR asset_no LIKE ? OR part_name LIKE ?)"
        like = '%' + q + '%'
        args += [like, like, like, like]
    sql += " ORDER BY id DESC"
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    data = []
    for r in rows:
        d = dict(r)
        d['items'] = _pc_up_items(conn, r['id'])
        data.append(d)
    conn.close()
    return jsonify({'ok': True, 'data': data})


@app.route('/api/pc-upgrades', methods=['POST'])
@require_login
def add_pc_upgrade():
    d = request.get_json(force=True, silent=True) or {}
    if not str(d.get('applicant', '')).strip() and not str(d.get('flow_no', '')).strip():
        return jsonify({'ok': False, 'error': '请至少填写申请人或流程单号'})
    items = _pc_norm_items(d.get('items'))
    if not items:
        return jsonify({'ok': False, 'error': '请至少添加一行配件明细（品牌/名称/数量）'})
    conn = get_conn()
    rec_time = d.get('upgrade_date') or now_str()[:10]
    cur = conn.execute(
        "INSERT INTO pc_upgrades(applicant,flow_no,asset_no,upgrade_date,remark,zk,bg,purchase_date,created_at) "
        "VALUES(?,?,?,?,?,?,?,?,?)",
        (str(d.get('applicant', '')).strip(), str(d.get('flow_no', '')).strip(),
         str(d.get('asset_no', '')).strip(), rec_time, str(d.get('remark', '')).strip(),
         int(num(d.get('zk'))), int(num(d.get('bg'))), _up_date(d.get('purchase_date')), now_str()))
    up_id = cur.lastrowid
    try:
        for it in items:
            brand, name, spec = it['brand'], it['name'], it['spec']
            if it['part_id']:
                part = conn.execute("SELECT * FROM pc_parts WHERE id=?", (it['part_id'],)).fetchone()
                if not part:
                    raise ValueError('关联备件不存在')
                brand, name, spec = part['brand'] or '', part['name'], part['spec'] or ''
            conn.execute(
                "INSERT INTO pc_upgrade_items(upgrade_id,part_id,brand,name,spec,qty,created_at) VALUES(?,?,?,?,?,?,?)",
                (up_id, it['part_id'], brand, name, spec, it['qty'], now_str()))
        _pc_apply_items(conn, up_id, items, session['real_name'] or session['username'], rec_time)
        conn.commit()
    except ValueError as e:
        conn.rollback()
        conn.close()
        return jsonify({'ok': False, 'error': str(e)})
    except Exception as e:
        conn.rollback()
        conn.close()
        return jsonify({'ok': False, 'error': '保存失败：%s' % e})
    conn.close()
    log(session['username'], '新增电脑优化记录', str(up_id))
    return jsonify({'ok': True, 'id': up_id})


@app.route('/api/pc-upgrades/<int:rid>', methods=['PUT'])
@require_login
def update_pc_upgrade(rid):
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    old = conn.execute("SELECT * FROM pc_upgrades WHERE id=?", (rid,)).fetchone()
    if not old:
        conn.close()
        return jsonify({'ok': False, 'error': '记录不存在'}), 404
    items = _pc_norm_items(d.get('items'))
    if not items:
        conn.close()
        return jsonify({'ok': False, 'error': '请至少添加一行配件明细（品牌/名称/数量）'})
    try:
        old_items = _pc_up_items(conn, rid)
        _pc_revert_items(conn, rid, old_items)
        conn.execute("DELETE FROM pc_upgrade_items WHERE upgrade_id=?", (rid,))
        rec_time = d.get('upgrade_date') or now_str()[:10]
        for it in items:
            brand, name, spec = it['brand'], it['name'], it['spec']
            if it['part_id']:
                part = conn.execute("SELECT * FROM pc_parts WHERE id=?", (it['part_id'],)).fetchone()
                if not part:
                    raise ValueError('关联备件不存在')
                brand, name, spec = part['brand'] or '', part['name'], part['spec'] or ''
            conn.execute(
                "INSERT INTO pc_upgrade_items(upgrade_id,part_id,brand,name,spec,qty,created_at) VALUES(?,?,?,?,?,?,?)",
                (rid, it['part_id'], brand, name, spec, it['qty'], now_str()))
        _pc_apply_items(conn, rid, items, session['real_name'] or session['username'], rec_time)
        conn.execute(
            "UPDATE pc_upgrades SET applicant=?,flow_no=?,asset_no=?,upgrade_date=?,remark=?,zk=?,bg=?,purchase_date=? WHERE id=?",
            (str(d.get('applicant', '')).strip(), str(d.get('flow_no', '')).strip(),
             str(d.get('asset_no', '')).strip(), rec_time, str(d.get('remark', '')).strip(),
             int(num(d.get('zk'))), int(num(d.get('bg'))), _up_date(d.get('purchase_date')), rid))
        conn.commit()
    except ValueError as e:
        conn.rollback()
        conn.close()
        return jsonify({'ok': False, 'error': str(e)})
    except Exception as e:
        conn.rollback()
        conn.close()
        return jsonify({'ok': False, 'error': '保存失败：%s' % e})
    conn.close()
    log(session['username'], '更新电脑优化记录', str(rid))
    return jsonify({'ok': True})


@app.route('/api/pc-upgrades/<int:rid>', methods=['DELETE'])
@require_login
def delete_pc_upgrade(rid):
    conn = get_conn()
    row = conn.execute("SELECT * FROM pc_upgrades WHERE id=?", (rid,)).fetchone()
    if row:
        items = _pc_up_items(conn, rid)
        _pc_revert_items(conn, rid, items)
        conn.execute("DELETE FROM pc_upgrade_items WHERE upgrade_id=?", (rid,))
        conn.execute("DELETE FROM pc_upgrades WHERE id=?", (rid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除电脑优化记录', str(rid))
    return jsonify({'ok': True})


@app.route('/api/pc-upgrades/template', methods=['GET'])
@require_login
def pc_upgrade_template():
    import openpyxl, io as _io
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '电脑优化记录导入模板'
    ws.append(PC_UP_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='D9E2F3')
        cell.alignment = Alignment(horizontal='center')
    ws.append(['王宽', 'IT03202605012', '20220727000028', '', '1', '1', '', '', '2026-05-12', '重装系统', '西数/机械硬盘/1T', 6, 4, 2, '2026-05-18'])
    widths = [12, 18, 18, 10, 10, 10, 10, 10, 12, 16, 24, 10, 8, 8, 12]
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='办公电脑优化记录导入模板.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/api/pc-upgrades/import', methods=['POST'])
@require_login
def import_pc_upgrades():
    f = request.files.get('file')
    if not f or not f.filename:
        return jsonify({'ok': False, 'error': '请选择要导入的 Excel 文件'})
    if not (f.filename.endswith('.xlsx') or f.filename.endswith('.xls')):
        return jsonify({'ok': False, 'error': '仅支持 .xlsx / .xls 文件'})
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(f.read()), data_only=True)
    except Exception as e:
        return jsonify({'ok': False, 'error': '文件解析失败：%s' % e})
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    conn = get_conn()
    prepared = []
    errors = []
    for idx, row in enumerate(rows, start=2):
        if row is None or all(v is None or str(v).strip() == '' for v in row):
            continue

        def gv(i):
            return str(row[i] or '').strip() if len(row) > i and row[i] is not None else ''
        applicant = gv(0)
        flow_no = gv(1)
        if not applicant and not flow_no:
            errors.append('第%d行：申请人或流程单号至少填一项' % idx)
            continue
        n = {}
        for i, k in ((3, 'ram_8'), (4, 'ram_16'), (5, 'sata_ssd'), (6, 'm2_ssd'), (7, 'hdd')):
            raw = gv(i)
            try:
                n[k] = int(float(raw)) if raw else 0
            except ValueError:
                errors.append('第%d行：%s 数量「%s」不是数字' % (idx, PC_UP_HEADERS[i], raw))
                break
        else:
            stock_raw = gv(11)
            try:
                part_stock = int(float(stock_raw)) if stock_raw else None
            except ValueError:
                errors.append('第%d行：剩余数量「%s」不是数字' % (idx, stock_raw))
                continue
            prepared.append((applicant, flow_no, gv(2), n['ram_8'], n['ram_16'], n['sata_ssd'], n['m2_ssd'], n['hdd'],
                             _up_date(row[8]) if len(row) > 8 else '', gv(9), gv(10), part_stock,
                             gv(12), gv(13), _up_date(row[14]) if len(row) > 14 else ''))
    if errors:
        conn.close()
        return jsonify({'ok': False, 'error': '存在 %d 行校验失败，已整体取消导入' % len(errors), 'errors': errors[:20]})
    if not prepared:
        conn.close()
        return jsonify({'ok': False, 'error': '模板中没有可导入的数据行'})
    try:
        now = now_str()
        for p in prepared:
            cur = conn.execute(
                "INSERT INTO pc_upgrades(applicant,flow_no,asset_no,ram_8,ram_16,sata_ssd,m2_ssd,hdd,upgrade_date,remark,part_name,part_stock,zk,bg,purchase_date,created_at) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (p[0], p[1], p[2], p[3], p[4], p[5], p[6], p[7], p[8], p[9], p[10], p[11], p[12], p[13], p[14], now))
            up_id = cur.lastrowid
            # 非零升级项自动生成配件明细（品牌待后续手动补充）
            for i in (3, 4, 5, 6, 7):
                qty = p[i]
                if qty and qty > 0:
                    conn.execute(
                        "INSERT INTO pc_upgrade_items(upgrade_id,part_id,brand,name,spec,qty,created_at) VALUES(?,NULL,'',?,?,?,?,?)",
                        (up_id, UPGRADE_TYPE_NAMES[i], '', qty, now))
        conn.commit()
        conn.close()
        log(session['username'], '模板导入电脑优化记录', '共%d条' % len(prepared))
        return jsonify({'ok': True, 'added': len(prepared)})
    except Exception as e:
        conn.rollback()
        conn.close()
        return jsonify({'ok': False, 'error': '导入执行失败：%s' % e})


@app.route('/api/pc-upgrades/export', methods=['GET'])
@require_login
def export_pc_upgrades():
    q = request.args.get('q', '').strip()
    conn = get_conn()
    ids = _ids_arg()
    if ids:
        ph = ','.join('?' * len(ids))
        sql = "SELECT * FROM pc_upgrades WHERE id IN (%s)" % ph
        args = list(ids)
    else:
        sql = "SELECT * FROM pc_upgrades WHERE 1=1"
        args = []
    if q:
        sql += " AND (applicant LIKE ? OR flow_no LIKE ? OR asset_no LIKE ? OR part_name LIKE ?)"
        like = '%' + q + '%'
        args += [like, like, like, like]
    sql += " ORDER BY id DESC"
    rows = conn.execute(sql, args).fetchall()
    import openpyxl, io as _io
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '办公电脑优化记录'
    headers = ['申请人', '流程单号', '资产编号', '配件品牌', '配件名称', '规格', '数量', '升级日期', '备注', '采购时间']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='D9E2F3')
        cell.alignment = Alignment(horizontal='center')
    for r in rows:
        items = _pc_up_items(conn, r['id'])
        if items:
            for it in items:
                ws.append([r['applicant'], r['flow_no'], r['asset_no'], it['brand'], it['name'], it['spec'],
                           it['qty'], r['upgrade_date'], r['remark'], r['purchase_date']])
        else:
            # 旧数据兜底：5 类升级项转行
            for k, nm in (('ram_8', '8G内存'), ('ram_16', '16G内存'), ('sata_ssd', 'SATA固态'), ('m2_ssd', 'M2固态'), ('hdd', '机械硬盘')):
                qty = r[k]
                if qty:
                    ws.append([r['applicant'], r['flow_no'], r['asset_no'], '', nm, '', qty,
                               r['upgrade_date'], r['remark'], r['purchase_date']])
    conn.close()
    widths = [12, 18, 18, 14, 22, 12, 8, 12, 20, 12]
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='办公电脑优化记录导出.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


# ---------------- 电脑优化备件库存体系（物资/出入库/调整，逻辑同备品备件） ----------------
@app.route('/api/pc-parts', methods=['GET'])
@require_login
def list_pc_parts():
    q = request.args.get('q', '').strip()
    sql = "SELECT * FROM pc_parts WHERE 1=1"
    args = []
    if q:
        sql += " AND (name LIKE ? OR brand LIKE ? OR spec LIKE ?)"
        like = '%' + q + '%'
        args += [like, like, like]
    sql += " ORDER BY id DESC"
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows]})


@app.route('/api/pc-parts', methods=['POST'])
@require_login
def add_pc_part():
    d = request.get_json(force=True, silent=True) or {}
    name = str(d.get('name', '')).strip()
    if not name:
        return jsonify({'ok': False, 'error': '物资名称必填'})
    conn = get_conn()
    dup = conn.execute("SELECT id FROM pc_parts WHERE name=? AND brand=? AND spec=? AND unit=?",
                       (name, str(d.get('brand', '')).strip(), str(d.get('spec', '')).strip(),
                        str(d.get('unit', '') or '个').strip())).fetchone()
    if dup:
        conn.close()
        return jsonify({'ok': False, 'error': '该物资已存在，可直接登记出入库'})
    cur = conn.execute(
        "INSERT INTO pc_parts(brand,name,spec,unit,stock,low_threshold,remark,storage_location,created_at) VALUES(?,?,?,?,?,?,?,?,?)",
        (str(d.get('brand', '')).strip(), name, str(d.get('spec', '')).strip(),
         str(d.get('unit', '') or '个').strip(), int(num(d.get('stock'))), int(num(d.get('low_threshold'))),
         str(d.get('remark', '')).strip(), str(d.get('storage_location', '')).strip(), now_str()))
    conn.commit()
    conn.close()
    log(session['username'], '新增电脑优化备件', name)
    return jsonify({'ok': True, 'id': cur.lastrowid})


@app.route('/api/pc-parts/<int:pid>', methods=['PUT'])
@require_login
def update_pc_part(pid):
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    cur = conn.execute(
        "UPDATE pc_parts SET brand=?,name=?,spec=?,unit=?,low_threshold=?,remark=?,storage_location=? WHERE id=?",
        (str(d.get('brand', '')).strip(), str(d.get('name', '')).strip(), str(d.get('spec', '')).strip(),
         str(d.get('unit', '') or '个').strip(), int(num(d.get('low_threshold'))),
         str(d.get('remark', '')).strip(), str(d.get('storage_location', '')).strip(), pid))
    conn.commit()
    conn.close()
    if cur.rowcount == 0:
        return jsonify({'ok': False, 'error': '物资不存在'}), 404
    log(session['username'], '编辑电脑优化备件', str(pid))
    return jsonify({'ok': True})


@app.route('/api/pc-parts/<int:pid>', methods=['DELETE'])
@require_login
def delete_pc_part(pid):
    conn = get_conn()
    row = conn.execute("SELECT * FROM pc_parts WHERE id=?", (pid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '物资不存在'}), 404
    if int(num(row['stock'])) != 0:
        conn.close()
        return jsonify({'ok': False, 'error': '当前库存 %s，请先出库/调整清零后再删除' % row['stock']})
    conn.execute("DELETE FROM pc_parts WHERE id=?", (pid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除电脑优化备件', row['name'])
    return jsonify({'ok': True})


@app.route('/api/pc-flows', methods=['GET'])
@require_login
def list_pc_flows():
    q = request.args.get('q', '').strip()
    ftype = request.args.get('type', '').strip()
    sql = "SELECT * FROM pc_flows WHERE 1=1"
    args = []
    if q:
        sql += " AND (part_name LIKE ? OR operator LIKE ? OR remark LIKE ? OR ref LIKE ?)"
        like = '%' + q + '%'
        args += [like, like, like, like]
    if ftype:
        sql += " AND type=?"
        args.append(ftype)
    sql += " ORDER BY record_time DESC, id DESC"
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows]})


@app.route('/api/pc-flows', methods=['POST'])
@require_login
def add_pc_flow():
    """电脑优化备件出入库：入库加库存、出库扣库存（不足拒绝），事务内完成"""
    d = request.get_json(force=True, silent=True) or {}
    ftype = d.get('type', 'in')
    if ftype not in ('in', 'out'):
        return jsonify({'ok': False, 'error': '类型必须为入库或出库'})
    qty = int(num(d.get('qty')))
    if qty <= 0:
        return jsonify({'ok': False, 'error': '数量必须大于 0'})
    part_id = d.get('part_id')
    conn = get_conn()
    if part_id:
        part = conn.execute("SELECT * FROM pc_parts WHERE id=?", (part_id,)).fetchone()
        if not part:
            conn.close()
            return jsonify({'ok': False, 'error': '物资不存在'}), 404
    else:
        name = str(d.get('name', '')).strip()
        if not name:
            conn.close()
            return jsonify({'ok': False, 'error': '请选择或填写物资名称'})
        part = conn.execute(
            "SELECT * FROM pc_parts WHERE name=? AND brand=? AND spec=? AND unit=? LIMIT 1",
            (name, str(d.get('brand', '')).strip(), str(d.get('spec', '')).strip(),
             str(d.get('unit', '') or '个').strip())).fetchone()
        if not part:
            cur = conn.execute("INSERT INTO pc_parts(brand,name,spec,unit,stock,low_threshold,remark,created_at) VALUES(?,?,?,?,?,?,?,?)",
                               (str(d.get('brand', '')).strip(), name, str(d.get('spec', '')).strip(),
                                str(d.get('unit', '') or '个').strip(), 0, 0, '流水自动创建', now_str()))
            part = {'id': cur.lastrowid, 'name': name, 'brand': d.get('brand', ''), 'spec': d.get('spec', ''), 'unit': d.get('unit', '') or '个'}
        part_id = part['id']
    stock = int(num(part['stock']))
    if ftype == 'out' and stock < qty:
        conn.close()
        return jsonify({'ok': False, 'error': '库存不足！当前库存 %s' % stock})
    conn.execute("UPDATE pc_parts SET stock=stock+? WHERE id=?", (qty if ftype == 'in' else -qty, part_id))
    try:
        record_time = d.get('record_time') or now_str()[:10]
        datetime.datetime.strptime(record_time[:10], '%Y-%m-%d')
    except ValueError:
        record_time = datetime.datetime.now().strftime('%Y-%m-%d')
    cur = conn.execute(
        "INSERT INTO pc_flows(type,part_id,part_name,qty,operator,ref,record_time,remark,created_at) VALUES(?,?,?,?,?,?,?,?,?)",
        (ftype, part_id, part['name'], qty, session['real_name'] or session['username'],
         str(d.get('ref', '')).strip(), record_time, str(d.get('remark', '')).strip(), now_str()))
    conn.commit()
    conn.close()
    log(session['username'], '电脑优化备件' + ('入库' if ftype == 'in' else '出库'), 'id=%s qty=%s' % (cur.lastrowid, qty))
    return jsonify({'ok': True, 'id': cur.lastrowid})


@app.route('/api/pc-flows/<int:fid>', methods=['DELETE'])
@require_login
def delete_pc_flow(fid):
    """删除流水并反向回补库存：入库→库存-数量；出库→库存+数量"""
    conn = get_conn()
    row = conn.execute("SELECT * FROM pc_flows WHERE id=?", (fid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '流水不存在'}), 404
    ftype, qty, part_id = row['type'], int(num(row['qty'])), row['part_id']
    part = conn.execute("SELECT * FROM pc_parts WHERE id=?", (part_id,)).fetchone()
    if ftype == 'in':
        if not part or int(num(part['stock'])) < qty:
            conn.close()
            return jsonify({'ok': False, 'error': '无法删除：当前库存不足（%s），请先调整库存' % (part['stock'] if part else 0)})
        conn.execute("UPDATE pc_parts SET stock=stock-? WHERE id=?", (qty, part_id))
    else:
        if not part:
            conn.close()
            return jsonify({'ok': False, 'error': '无法删除：对应物资已不存在'})
        conn.execute("UPDATE pc_parts SET stock=stock+? WHERE id=?", (qty, part_id))
    conn.execute("DELETE FROM pc_flows WHERE id=?", (fid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除电脑优化备件流水', '%s qty=%s' % (fid, qty))
    return jsonify({'ok': True})


@app.route('/api/pc-flows/template', methods=['GET'])
@require_login
def pc_flow_template():
    import openpyxl, io as _io
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '出入库流水导入模板'
    headers = ['类型*', '物资名称*', '品牌', '规格', '单位', '数量*', '日期', '经办人', '备注']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='D9E2F3')
        cell.alignment = Alignment(horizontal='center')
    ws.append(['入库', '西数/机械硬盘/1T', '西数', '1T', '个', 10, '2026-10-09', '', '采购入库'])
    ws.append(['出库', '金士顿/8G内存', '金士顿', '8G', '条', 2, '2026-10-09', '', ''])
    widths = [10, 24, 12, 10, 8, 10, 12, 12, 20]
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='电脑优化备件出入库流水导入模板.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/api/pc-flows/import', methods=['POST'])
@require_login
def import_pc_flows():
    f = request.files.get('file')
    if not f or not f.filename:
        return jsonify({'ok': False, 'error': '请选择要导入的 Excel 文件'})
    if not (f.filename.endswith('.xlsx') or f.filename.endswith('.xls')):
        return jsonify({'ok': False, 'error': '仅支持 .xlsx / .xls 文件'})
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(f.read()), data_only=True)
    except Exception as e:
        return jsonify({'ok': False, 'error': '文件解析失败：%s' % e})
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    conn = get_conn()
    prepared = []
    errors = []
    for idx, row in enumerate(rows, start=2):
        if row is None or all(v is None or str(v).strip() == '' for v in row):
            continue

        def gv(i):
            return str(row[i] or '').strip() if len(row) > i and row[i] is not None else ''
        ftype = gv(0)
        if ftype not in ('in', 'out', '入库', '出库'):
            errors.append('第%d行：类型必须为 入库/出库' % idx)
            continue
        ftype = 'in' if ftype in ('in', '入库') else 'out'
        name = gv(1)
        if not name:
            errors.append('第%d行：物资名称必填' % idx)
            continue
        try:
            qty = int(float(gv(5)))
        except ValueError:
            errors.append('第%d行：数量「%s」不是数字' % (idx, gv(5)))
            continue
        if qty <= 0:
            errors.append('第%d行：数量必须大于 0' % idx)
            continue
        part = conn.execute("SELECT * FROM pc_parts WHERE name=? AND brand=? AND spec=? AND unit=? LIMIT 1",
                            (name, gv(2), gv(3), gv(4) or '个')).fetchone()
        if not part:
            cur = conn.execute("INSERT INTO pc_parts(brand,name,spec,unit,stock,low_threshold,remark,created_at) VALUES(?,?,?,?,?,?,?,?)",
                               (gv(2), name, gv(3), gv(4) or '个', 0, 0, '流水导入自动创建', now_str()))
            part = {'id': cur.lastrowid, 'name': name, 'brand': gv(2), 'spec': gv(3), 'unit': gv(4) or '个'}
        if ftype == 'out' and int(num(part['stock'])) < qty:
            errors.append('第%d行：库存不足！「%s」当前库存 %s，出库 %s' % (idx, name, part['stock'], qty))
            continue
        prepared.append((ftype, part['id'], name, qty, gv(6), gv(7), gv(8)))
    if errors:
        conn.close()
        return jsonify({'ok': False, 'error': '存在 %d 行校验失败，已整体取消导入' % len(errors), 'errors': errors[:20]})
    if not prepared:
        conn.close()
        return jsonify({'ok': False, 'error': '模板中没有可导入的数据行'})
    try:
        now = now_str()
        for p in prepared:
            ftype, part_id, name, qty, rtime, operator, remark = p
            conn.execute("UPDATE pc_parts SET stock=stock+? WHERE id=?", (qty if ftype == 'in' else -qty, part_id))
            conn.execute(
                "INSERT INTO pc_flows(type,part_id,part_name,qty,operator,ref,record_time,remark,created_at) VALUES(?,?,?,?,?,?,?,?,?)",
                (ftype, part_id, name, qty, operator or session['real_name'] or session['username'], '批量导入',
                 rtime or now_str()[:10], remark, now))
        conn.commit()
        conn.close()
        log(session['username'], '模板导入电脑优化备件流水', '共%d条' % len(prepared))
        return jsonify({'ok': True, 'added': len(prepared)})
    except Exception as e:
        conn.rollback()
        conn.close()
        return jsonify({'ok': False, 'error': '导入执行失败：%s' % e})


@app.route('/api/pc-flows/export', methods=['GET'])
@require_login
def export_pc_flows():
    q = request.args.get('q', '').strip()
    ftype = request.args.get('type', '').strip()
    start = request.args.get('start', '').strip()
    end = request.args.get('end', '').strip()
    sql = "SELECT * FROM pc_flows WHERE 1=1"
    args = []
    if q:
        sql += " AND (part_name LIKE ? OR operator LIKE ? OR remark LIKE ?)"
        like = '%' + q + '%'
        args += [like, like, like]
    if ftype:
        sql += " AND type=?"
        args.append(ftype)
    if start:
        sql += " AND record_time >= ?"
        args.append(start)
    if end:
        sql += " AND record_time <= ?"
        args.append(end)
    sql += " ORDER BY record_time DESC, id DESC"
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    import openpyxl, io as _io
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '出入库流水'
    headers = ['类型', '物资名称', '数量', '日期', '经办人', '关联', '备注']
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill('solid', fgColor='D9E2F3')
        cell.alignment = Alignment(horizontal='center')
    for r in rows:
        ws.append(['入库' if r['type'] == 'in' else '出库', r['part_name'], r['qty'], r['record_time'],
                   r['operator'], r['ref'], r['remark']])
    widths = [10, 30, 10, 12, 12, 16, 24]
    from openpyxl.utils import get_column_letter
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='电脑优化备件出入库流水导出.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@app.route('/api/pc-adjusts', methods=['GET'])
@require_login
def list_pc_adjusts():
    q = request.args.get('q', '').strip()
    sql = "SELECT * FROM pc_adjusts WHERE 1=1"
    args = []
    if q:
        sql += " AND (part_name LIKE ? OR reason LIKE ? OR operator LIKE ?)"
        like = '%' + q + '%'
        args += [like, like, like]
    sql += " ORDER BY adjust_time DESC, id DESC"
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows]})


@app.route('/api/pc-adjusts', methods=['POST'])
@require_login
def add_pc_adjust():
    """电脑优化备件库存调整：in=增加 / out=减少，事务内更新库存并记录"""
    d = request.get_json(force=True, silent=True) or {}
    atype = d.get('type', 'in')
    if atype not in ('in', 'out'):
        return jsonify({'ok': False, 'error': '调整类型必须为 增加/减少'})
    qty = int(num(d.get('qty')))
    if qty <= 0:
        return jsonify({'ok': False, 'error': '数量必须大于 0'})
    part_id = d.get('part_id')
    if not part_id:
        return jsonify({'ok': False, 'error': '请选择物资'})
    conn = get_conn()
    part = conn.execute("SELECT * FROM pc_parts WHERE id=?", (part_id,)).fetchone()
    if not part:
        conn.close()
        return jsonify({'ok': False, 'error': '物资不存在'}), 404
    if atype == 'out' and int(num(part['stock'])) < qty:
        conn.close()
        return jsonify({'ok': False, 'error': '库存不足！当前库存 %s' % part['stock']})
    conn.execute("UPDATE pc_parts SET stock=stock+? WHERE id=?", (qty if atype == 'in' else -qty, part_id))
    cur = conn.execute(
        "INSERT INTO pc_adjusts(type,part_id,part_name,qty,reason,operator,adjust_time,created_at) VALUES(?,?,?,?,?,?,?,?)",
        (atype, part_id, part['name'], qty, str(d.get('reason', '')).strip(),
         session['real_name'] or session['username'], d.get('adjust_time') or now_str()[:10], now_str()))
    conn.commit()
    conn.close()
    log(session['username'], '电脑优化备件库存调整', 'id=%s %s qty=%s' % (cur.lastrowid, atype, qty))
    return jsonify({'ok': True, 'id': cur.lastrowid})


@app.route('/api/pc-adjusts/<int:aid>', methods=['DELETE'])
@require_login
def delete_pc_adjust(aid):
    """删除调整记录并反向回补库存：增加调整→库存-数量；减少调整→库存+数量"""
    conn = get_conn()
    row = conn.execute("SELECT * FROM pc_adjusts WHERE id=?", (aid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '调整记录不存在'}), 404
    qty, atype = int(num(row['qty'])), row['type']
    part = conn.execute("SELECT * FROM pc_parts WHERE id=?", (row['part_id'],)).fetchone()
    if atype == 'in':
        if not part or int(num(part['stock'])) < qty:
            conn.close()
            return jsonify({'ok': False, 'error': '无法删除：当前库存不足（%s），请先调整库存' % (part['stock'] if part else 0)})
        conn.execute("UPDATE pc_parts SET stock=stock-? WHERE id=?", (qty, row['part_id']))
    else:
        if not part:
            conn.close()
            return jsonify({'ok': False, 'error': '无法删除：对应物资已不存在'})
        conn.execute("UPDATE pc_parts SET stock=stock+? WHERE id=?", (qty, row['part_id']))
    conn.execute("DELETE FROM pc_adjusts WHERE id=?", (aid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除电脑优化备件调整记录', '%s qty=%s' % (aid, qty))
    return jsonify({'ok': True})


# ---------------- 用户管理 ----------------
@app.route('/api/users', methods=['GET'])
@require_login
def list_users():
    conn = get_conn()
    rows = conn.execute("SELECT id,username,real_name,role,active,perms,created_at FROM users ORDER BY id").fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows]})


@app.route('/api/users', methods=['POST'])
@require_login
def add_user():
    if session.get('role') != 'admin':
        return jsonify({'ok': False, 'error': '仅管理员可新增账号'}), 403
    d = request.get_json(force=True, silent=True) or {}
    username = str(d.get('username', '')).strip()
    password = str(d.get('password', ''))
    if not username or len(password) < 4:
        return jsonify({'ok': False, 'error': '用户名不能为空，密码至少4位'})
    conn = get_conn()
    if conn.execute("SELECT id FROM users WHERE username=?", (username,)).fetchone():
        conn.close()
        return jsonify({'ok': False, 'error': '用户名已存在'})
    perms = json.dumps(d.get('perms') or {}, ensure_ascii=False)
    cur = conn.execute(
        "INSERT INTO users(username,password_hash,real_name,role,active,perms,created_at) VALUES(?,?,?,?,?,?,?)",
        (username, generate_password_hash(password), str(d.get('real_name', '')).strip(),
         d.get('role') or 'admin', 1, perms, now_str()))
    conn.commit()
    conn.close()
    log(session['username'], '新增用户', username)
    return jsonify({'ok': True, 'id': cur.lastrowid})


@app.route('/api/users/batch', methods=['POST'])
@require_login
def batch_add_users():
    """批量创建用户：每行 username/password/real_name/role；viewer 默认全部模块仅查看"""
    if session.get('role') != 'admin':
        return jsonify({'ok': False, 'error': '仅管理员可批量创建账号'}), 403
    items = (request.get_json(force=True, silent=True) or {}).get('users', [])
    if not items or not isinstance(items, list):
        return jsonify({'ok': False, 'error': '没有要创建的用户'})
    results = []
    ok_count = 0
    conn = get_conn()
    for it in items:
        username = str(it.get('username', '')).strip()
        row = {'username': username}
        if not username:
            row.update(ok=False, msg='用户名为空'); results.append(row); continue
        if conn.execute("SELECT id FROM users WHERE username=?", (username,)).fetchone():
            row.update(ok=False, msg='用户名已存在'); results.append(row); continue
        password = str(it.get('password', '') or '123456')
        if len(password) < 4:
            row.update(ok=False, msg='密码至少4位'); results.append(row); continue
        role = it.get('role') or 'viewer'
        if role not in ('admin', 'viewer'):
            role = 'viewer'
        perms = {}
        if role == 'viewer':
            perms = {m: 'view' for m in PERM_MODULES}
        cur = conn.execute(
            "INSERT INTO users(username,password_hash,real_name,role,active,perms,created_at) VALUES(?,?,?,?,?,?,?)",
            (username, generate_password_hash(password), str(it.get('real_name', '')).strip(),
             role, 1, json.dumps(perms, ensure_ascii=False), now_str()))
        row.update(ok=True, msg='创建成功', id=cur.lastrowid)
        ok_count += 1
        results.append(row)
    conn.commit()
    conn.close()
    log(session['username'], '批量创建用户', '成功 %d/%d' % (ok_count, len(results)))
    return jsonify({'ok': True, 'total': len(results), 'success': ok_count, 'data': results})


@app.route('/api/users/<int:uid>', methods=['PUT'])
@require_login
def update_user(uid):
    if session.get('role') != 'admin':
        return jsonify({'ok': False, 'error': '仅管理员可修改账号'}), 403
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    row = conn.execute("SELECT * FROM users WHERE id=?", (uid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '用户不存在'}), 404
    if d.get('password'):
        conn.execute("UPDATE users SET password_hash=? WHERE id=?",
                     (generate_password_hash(str(d['password'])), uid))
    if 'perms' in d:
        try:
            conn.execute("UPDATE users SET perms=? WHERE id=?",
                         (json.dumps(d['perms'], ensure_ascii=False), uid))
        except Exception:
            pass
    conn.execute("UPDATE users SET real_name=?, role=?, active=? WHERE id=?",
                 (str(d.get('real_name', row['real_name'] or '')).strip(),
                  d.get('role') or row['role'],
                  1 if d.get('active', row['active']) in (1, '1', True) else 0, uid))
    conn.commit()
    conn.close()
    log(session['username'], '修改用户', row['username'])
    return jsonify({'ok': True})


@app.route('/api/users/<int:uid>', methods=['DELETE'])
@require_login
def delete_user(uid):
    if session.get('role') != 'admin':
        return jsonify({'ok': False, 'error': '仅管理员可删除账号'}), 403
    if uid == session.get('uid'):
        return jsonify({'ok': False, 'error': '不能删除当前登录账号'})
    conn = get_conn()
    row = conn.execute("SELECT username FROM users WHERE id=?", (uid,)).fetchone()
    # 级联清理该用户的个性化设置（列配置等），避免孤儿数据
    if row:
        conn.execute("DELETE FROM user_settings WHERE username=?", (row['username'],))
    conn.execute("DELETE FROM users WHERE id=?", (uid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除用户', row['username'] if row else str(uid))
    return jsonify({'ok': True})


@app.route('/api/password', methods=['POST'])
@require_login
def change_password():
    d = request.get_json(force=True, silent=True) or {}
    old = str(d.get('old', ''))
    new = str(d.get('new', ''))
    if len(new) < 4:
        return jsonify({'ok': False, 'error': '新密码至少4位'})
    conn = get_conn()
    row = conn.execute("SELECT * FROM users WHERE id=?", (session['uid'],)).fetchone()
    if not row or not check_password_hash(row['password_hash'], old):
        conn.close()
        return jsonify({'ok': False, 'error': '原密码错误'})
    conn.execute("UPDATE users SET password_hash=? WHERE id=?", (generate_password_hash(new), session['uid']))
    conn.commit()
    conn.close()
    log(session['username'], '修改密码')
    return jsonify({'ok': True})


# ---------------- 记事本（富文本笔记 + 待办） ----------------
@app.route('/api/notes', methods=['GET'])
@require_login
def list_notes():
    q = request.args.get('q', '').strip()
    category = request.args.get('category', '').strip()
    todo = request.args.get('todo', '').strip()
    trashed = request.args.get('trashed', '0').strip()
    where = " WHERE 1=1"
    args = []
    if trashed == '1':
        where += " AND deleted_at<>''"
    else:
        where += " AND deleted_at=''"
    if q:
        like = f"%{q}%"
        where += " AND (title LIKE ? OR content LIKE ?)"
        args += [like, like]
    if category:
        where += " AND category=?"
        args.append(category)
    if todo == 'undone':
        where += " AND is_todo=1 AND todo_done=0"
    elif todo == 'done':
        where += " AND todo_done=1"
    elif todo == 'all':
        where += " AND is_todo=1"
    conn = get_conn()
    rows = conn.execute(
        "SELECT id,title,author_id,author_name,category,is_todo,todo_done,todo_due,pinned,created_at,updated_at,deleted_at, "
        "substr(content,1,120) AS preview FROM notes" + where +
        " ORDER BY CASE WHEN deleted_at='' THEN pinned ELSE 1 END DESC, updated_at DESC, id DESC", args).fetchall()
    cats = [r['category'] for r in conn.execute("SELECT DISTINCT category FROM notes WHERE category<>'' AND deleted_at='' ORDER BY category").fetchall()]
    trash_count = conn.execute("SELECT COUNT(*) FROM notes WHERE deleted_at<>''").fetchone()[0]
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows], 'categories': cats, 'trash_count': trash_count})


@app.route('/api/notes/<int:nid>', methods=['GET'])
@require_login
def get_note(nid):
    conn = get_conn()
    row = conn.execute("SELECT * FROM notes WHERE id=?", (nid,)).fetchone()
    conn.close()
    if not row:
        return jsonify({'ok': False, 'error': '笔记不存在'}), 404
    return jsonify({'ok': True, 'data': dict(row)})


@app.route('/api/notes', methods=['POST'])
@require_login
def create_note():
    d = request.get_json(force=True, silent=True) or {}
    title = str(d.get('title') or '').strip() or '未命名笔记'
    now = now_str()
    conn = get_conn()
    cur = conn.execute(
        "INSERT INTO notes(title,content,author_id,author_name,category,is_todo,todo_done,todo_due,pinned,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?)",
        (title, str(d.get('content') or ''), session.get('uid'),
         session.get('real_name') or session.get('username'),
         str(d.get('category') or '默认').strip() or '默认',
         1 if d.get('is_todo') else 0, 0, str(d.get('todo_due') or ''),
         0, now, now))
    conn.commit()
    nid = cur.lastrowid
    conn.close()
    return jsonify({'ok': True, 'id': nid})


@app.route('/api/notes/<int:nid>', methods=['PUT'])
@require_login
def update_note(nid):
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    row = conn.execute("SELECT * FROM notes WHERE id=?", (nid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '笔记不存在'}), 404
    fields = []
    args = []
    for k in ('title', 'content', 'category', 'todo_due'):
        if k in d:
            fields.append(f"{k}=?")
            v = d[k]
            args.append(str(v) if v is not None else '')
    for k in ('is_todo', 'todo_done', 'pinned'):
        if k in d:
            fields.append(f"{k}=?")
            args.append(1 if d[k] in (1, '1', True, 'true') else 0)
    if not fields:
        conn.close()
        return jsonify({'ok': True})
    fields.append("updated_at=?")
    args.append(now_str())
    args.append(nid)
    conn.execute(f"UPDATE notes SET {', '.join(fields)} WHERE id=?", args)
    conn.commit()
    conn.close()
    return jsonify({'ok': True})


@app.route('/api/notes/<int:nid>', methods=['DELETE'])
@require_login
def delete_note(nid):
    """软删除：移入回收站"""
    conn = get_conn()
    row = conn.execute("SELECT author_id FROM notes WHERE id=?", (nid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '笔记不存在'}), 404
    if row['author_id'] != session.get('uid') and session.get('role') != 'admin':
        conn.close()
        return jsonify({'ok': False, 'error': '只能删除自己的笔记（管理员除外）'}), 403
    conn.execute("UPDATE notes SET deleted_at=? WHERE id=?", (now_str(), nid))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})


@app.route('/api/notes/<int:nid>/restore', methods=['POST'])
@require_login
def restore_note(nid):
    conn = get_conn()
    row = conn.execute("SELECT author_id FROM notes WHERE id=?", (nid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '笔记不存在'}), 404
    if row['author_id'] != session.get('uid') and session.get('role') != 'admin':
        conn.close()
        return jsonify({'ok': False, 'error': '只能恢复自己的笔记（管理员除外）'}), 403
    conn.execute("UPDATE notes SET deleted_at='' WHERE id=?", (nid,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})


@app.route('/api/notes/<int:nid>/hard', methods=['DELETE'])
@require_login
def hard_delete_note(nid):
    """彻底删除（回收站里）"""
    conn = get_conn()
    row = conn.execute("SELECT author_id FROM notes WHERE id=?", (nid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '笔记不存在'}), 404
    if row['author_id'] != session.get('uid') and session.get('role') != 'admin':
        conn.close()
        return jsonify({'ok': False, 'error': '只能删除自己的笔记（管理员除外）'}), 403
    conn.execute("DELETE FROM notes WHERE id=?", (nid,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})


# ---------------- 日志 ----------------
@app.route('/api/logs', methods=['GET'])
@require_login
def list_logs():
    page = max(1, int(request.args.get('page', 1)))
    per = max(1, min(100000, int(request.args.get('per', 50))))
    conn = get_conn()
    total = conn.execute("SELECT COUNT(*) c FROM sys_logs").fetchone()['c']
    rows = conn.execute("SELECT * FROM sys_logs ORDER BY id DESC LIMIT ? OFFSET ?",
                        (per, (page - 1) * per)).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows], 'total': total, 'page': page, 'per': per})


# ---------------- 备份 ----------------
def do_backup(manual=False):
    os.makedirs(BACKUP_DIR, exist_ok=True)
    ts = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
    dest = os.path.join(BACKUP_DIR, f'system_{ts}.db')
    try:
        src = sqlite3.connect(DB_PATH)
        dst = sqlite3.connect(dest)
        src.backup(dst)
        dst.close()
        src.close()
        size = os.path.getsize(dest)
        files = sorted(glob.glob(os.path.join(BACKUP_DIR, 'system_*.db')))
        removed = 0
        while len(files) > BACKUP_KEEP:
            os.remove(files.pop(0))
            removed += 1
        conn = get_conn()
        conn.execute("INSERT INTO backup_logs(time,file,size,result) VALUES(?,?,?,?)",
                     (now_str(), os.path.basename(dest), size, 'ok'))
        conn.commit()
        conn.close()
        log('system', '自动备份' if not manual else '手动备份', f"file={os.path.basename(dest)} keep_removed={removed}")
        return {'ok': True, 'file': os.path.basename(dest), 'size': size, 'removed': removed}
    except Exception as e:
        try:
            conn = get_conn()
            conn.execute("INSERT INTO backup_logs(time,file,size,result) VALUES(?,?,?,?)",
                         (now_str(), 'failed', 0, str(e)[:200]))
            conn.commit()
            conn.close()
        except Exception:
            pass
        return {'ok': False, 'error': str(e)}


@app.route('/api/backup', methods=['POST'])
@require_login
def manual_backup():
    return jsonify(do_backup(manual=True))


@app.route('/api/backups', methods=['GET'])
@require_login
def list_backups():
    conn = get_conn()
    rows = conn.execute("SELECT * FROM backup_logs ORDER BY id DESC LIMIT 50").fetchall()
    conn.close()
    files = sorted(glob.glob(os.path.join(BACKUP_DIR, 'system_*.db')), reverse=True)
    file_info = [{'name': os.path.basename(f), 'size': os.path.getsize(f),
                  'mtime': datetime.datetime.fromtimestamp(os.path.getmtime(f)).strftime('%Y-%m-%d %H:%M:%S')}
                 for f in files[:50]]
    return jsonify({'ok': True, 'logs': [dict(r) for r in rows], 'files': file_info})


def backup_loop():
    """后台线程：每天自动备份（当天首次轮询即备份，服务任何时间在线都不漏）"""
    last_day = None
    while True:
        try:
            today = datetime.date.today()
            if last_day != today:
                do_backup()
                last_day = today
        except Exception:
            pass
        time.sleep(300)


# ---------------- AI 助手模块 ----------------
AI_SYSTEM_PROMPT = """你是信息中心的工作记录整理助手。请把用户提供的原始工作信息（可能是口语描述、群聊文字、图片内容等，可能含冗余）整理成规范的工单工作记录。
只输出 JSON（不要输出任何其他文字、不要用代码块包裹），格式：
{"title": "工单名称", "summary": "工单描述", "requester": "需求人"}

规则：
1. 工单名称格式严格如下（注意 " -- " 前后各有一个空格，最后流程编号前用一个空格、不加 --）：
   桌面 -- 区域 -- 部门 -- 事项 [流程编号]
   各部分含义：
   - 桌面：故障/任务对象类型（如 电脑、打印机、网络、系统、投影、电话、即时通 等）；
   - 区域：位置信息（如 中控、三号岛、办公楼、五号岛 等），从原文提取；
   - 部门：需求所属部门（如 安环、检验计量中心、炼油六部、人力资源部 等），从原文提取；
   - 事项：简洁的故障/任务关键词（如 即时通文档无法打开、加装固态硬盘、申请打印机墨盒 等）；
   - 流程编号：原文若有流程编号（如 IT02202610027、TZ20260901 等编号+数字），直接用一个空格接在事项后面；没有则省略。
   示例：
   {"title": "桌面 -- 中控 -- 安环 -- 即时通文档无法打开问题处理"}
   {"title": "桌面 -- 三号岛 -- 炼油六部 -- 办公电脑加装固态硬盘、内存条及申领无线网卡 IT02202610027"}
2. 工单描述：只写核心故障/任务内容（现象与处理要点），简洁通顺；不要包含部门、需求人姓名、具体地点等冗余信息。
3. 需求人：从原文提取需求人姓名（如"需求人：张三"或直接的人名）；无法确定则为空字符串。
4. 统一简体中文；去除口语化冗余（"好的""麻烦""麻烦帮我看下"等客套）。
"""


def get_ai_settings():
    conn = get_conn()
    rows = conn.execute("SELECT k, v FROM ai_settings").fetchall()
    conn.close()
    d = {r['k']: r['v'] for r in rows}
    # 模型列表（JSON）
    ml = []
    try:
        ml = json.loads(d.get('model_list') or '[]')
    except Exception:
        ml = []
    # 兼容旧配置：如果没有模型列表但有单条配置，自动转成一条
    if not ml and d.get('api_base'):
        ml = [{'id': 'default', 'name': '默认模型', 'api_base': d.get('api_base',''), 'api_key': d.get('api_key',''), 'model': d.get('model','')}]
    active = d.get('active_model') or (ml[0]['id'] if ml else '')
    cur = next((m for m in ml if m.get('id') == active), ml[0] if ml else None)
    return {'api_base': (cur or {}).get('api_base',''),
            'api_key': (cur or {}).get('api_key',''),
            'model': (cur or {}).get('model',''),
            'model_list': ml,
            'active_model': active}


def save_ai_model_list(ml, active_id):
    conn = get_conn()
    cur = conn.cursor()
    cur.execute("INSERT INTO ai_settings(k,v) VALUES(?,?) ON CONFLICT(k) DO UPDATE SET v=excluded.v",
                ('model_list', json.dumps(ml, ensure_ascii=False)))
    cur.execute("INSERT INTO ai_settings(k,v) VALUES(?,?) ON CONFLICT(k) DO UPDATE SET v=excluded.v",
                ('active_model', active_id or (ml[0]['id'] if ml else '')))
    conn.commit()
    conn.close()


def save_ai_settings(api_base, api_key, model):
    conn = get_conn()
    cur = conn.cursor()
    for k, v in [('api_base', api_base), ('api_key', api_key), ('model', model)]:
        cur.execute("INSERT INTO ai_settings(k,v) VALUES(?,?) ON CONFLICT(k) DO UPDATE SET v=excluded.v", (k, v))
    conn.commit()
    conn.close()


def call_ai(prompt, text, images):
    """调用 OpenAI 兼容 chat/completions 接口；images 为 base64 字符串列表"""
    st = get_ai_settings()
    if not st['api_base'] or not st['api_key'] or not st['model']:
        raise RuntimeError('AI 接口未配置，请先在「AI 工作台 → AI 设置」填写 API 地址 / Key / 模型')
    url = st['api_base'].rstrip('/') + '/chat/completions'
    content = [{'type': 'text', 'text': (text or '（仅图片）')[:12000]}]
    for b64 in images[:3]:
        content.append({'type': 'image_url', 'image_url': {'url': 'data:image/jpeg;base64,' + b64}})
    body = json.dumps({
        'model': st['model'],
        'messages': [{'role': 'system', 'content': prompt},
                     {'role': 'user', 'content': content}],
        'temperature': 0.2,
        'response_format': {'type': 'json_object'}
    }).encode()
    req = urllib.request.Request(url, data=body, method='POST')
    req.add_header('Content-Type', 'application/json')
    req.add_header('Authorization', 'Bearer ' + st['api_key'])
    try:
        r = urllib.request.urlopen(req, timeout=120)
        resp = json.loads(r.read().decode('utf-8', 'ignore'))
    except urllib.error.HTTPError as e:
        raise RuntimeError('AI 接口错误 %s: %s' % (e.code, e.read().decode('utf-8', 'ignore')[:300]))
    except Exception as e:
        raise RuntimeError('AI 接口调用失败: %s' % e)
    try:
        return resp['choices'][0]['message']['content']
    except Exception:
        raise RuntimeError('AI 返回格式异常: %s' % str(resp)[:300])


def parse_ai_json(raw):
    """剥掉 markdown 代码块等，解析 JSON"""
    t = (raw or '').strip()
    t = t.replace('```json', '').replace('```', '').strip()
    # 找第一个 { 到最后一个 }
    a, b = t.find('{'), t.rfind('}')
    if a >= 0 and b > a:
        t = t[a:b + 1]
    import json as _json
    try:
        return _json.loads(t)
    except Exception:
        return None


@app.route('/api/ai/settings', methods=['GET'])
@require_login
def ai_get_settings():
    st = get_ai_settings()
    def mask(k): return (k[:6] + '****' + k[-4:]) if len(k) > 12 else ('已配置' if k else '')
    return jsonify({'ok': True,
                    'api_base': st['api_base'],
                    'model': st['model'],
                    'api_key_masked': mask(st['api_key']),
                    'model_list': st['model_list'],
                    'active_model': st['active_model']})


@app.route('/api/ai/settings', methods=['PUT'])
@require_login
def ai_put_settings():
    d = request.get_json(force=True, silent=True) or {}
    # 两种模式：兼容旧的单条保存 / 新的多模型列表
    if 'model_list' in d:
        ml = d.get('model_list') or []
        for i, m in enumerate(ml):
            if not m.get('id'): m['id'] = 'm' + str(i+1)
        save_ai_model_list(ml, d.get('active_model') or (ml[0]['id'] if ml else ''))
        log(session['username'], '修改AI模型列表', f"{len(ml)} 个模型")
        return jsonify({'ok': True})
    if 'switch_to' in d:
        st = get_ai_settings()
        save_ai_model_list(st['model_list'], d['switch_to'])
        return jsonify({'ok': True})
    api_base = str(d.get('api_base', '')).strip()
    model = str(d.get('model', '')).strip()
    api_key = str(d.get('api_key', '')).strip()
    cur = get_ai_settings()
    if not api_key:
        api_key = cur['api_key']
    save_ai_settings(api_base, api_key, model)
    log(session['username'], '修改AI设置', f"base={api_base} model={model}")
    return jsonify({'ok': True})


def rule_extract(text):
    """本地规则模式：无在线 AI 配置时的兜底整理（正则提取）"""
    import re
    t = (text or '').strip()
    # 流程编号（如 TZ20260901 / BG20260903）
    no = ''
    m = re.search(r'(?<![A-Za-z0-9])([A-Za-z]{1,4}\d{2,})(?![A-Za-z0-9])', t)
    if m:
        no = m.group(1)
    # 区域（如 3楼 / 四层 / 5栋201）——先提取并从原文移除
    area = ''
    m = re.search(r'([0-9零一二三四五六七八九十]{1,2}(?:楼|层|栋)(?:[0-9零一二三四五六七八九十]{1,3}号)?)', t)
    if m:
        area = m.group(1)
    t2 = t.replace(area, '') if area else t
    # 部门：在去区域后的文本中找「XX部/科/室/中心/组/处/学院」
    dept = ''
    m = re.search(r'([\u4e00-\u9fa5]{1,3}(?:部|科|室|中心|组|处|学院))', t2)
    if m:
        dept = m.group(1)
    # 需求人：优先「需求人：XX」；否则删部门后找「XX 反馈/报修/反映」
    req = ''
    m = re.search(r'需求人[:：]?\s*([\u4e00-\u9fa5]{2,4})', t)
    if m:
        req = m.group(1)
    else:
        t_no_dept = re.sub(r'[\u4e00-\u9fa5]{1,3}(?:部|科|室|中心|组|处|学院)', '', t2)
        m = re.search(r'([\u4e00-\u9fa5]{2,3})(?:反映|反馈|报修|打电话|来电|说|找)', t_no_dept)
        if m:
            req = m.group(1)
    # 桌面对象类型
    desk = ''
    for kw in ['电脑', '打印机', '投影仪', '投影', '网络', '系统', '电话', '扫描仪', '门禁', '监控', '服务器']:
        if kw in t:
            desk = kw
            break
    # 工单描述：从对象词起截取，按顺序删除 需求人/编号 → 客套 → 部门 → 区域
    idx = t.find(desk) if desk else 0
    seg = t[idx:idx + 40] if idx >= 0 else t[:40]
    seg = re.sub(r'需求人[:：]?\s*[\u4e00-\u9fa5]{2,4}', '', seg)
    seg = re.sub(r'(?:编号|流程号|工单号)[:：]?\s*[A-Za-z0-9-]+', '', seg)
    for w in ['报修', '反映', '反馈', '麻烦', '帮忙', '帮看', '请', '好的', '你好', '老师', '师傅', '说', '找']:
        seg = seg.replace(w, '')
    if dept:
        seg = re.sub(re.escape(dept) + '的?', '', seg)
    if area:
        seg = seg.replace(area, '')
    seg = re.sub(r'[，,。、;：:]+$', '', seg.strip())
    for w in ['电话', '在', '的', '说']:
        if seg.endswith(w):
            seg = seg[:-len(w)]
    if req and len(req) >= 2 and seg.startswith(req[0]):
        seg = seg[1:]
    seg = re.sub(r'[，,。、;：:]+$', '', seg.strip())
    seg = seg.strip(' ，,。、;')
    if len(seg) > 40:
        seg = seg[:40] + '…'
    item = (seg[:16] or (desk or '事项')).strip('，,。')
    title = '--'.join([desk or '未知', area, dept, item, no])
    return {'title': title, 'summary': seg, 'requester': req, 'mode': 'rule'}


@app.route('/api/ai/generate', methods=['POST'])
@require_login
def ai_generate():
    d = request.get_json(force=True, silent=True) or {}
    text = str(d.get('text', ''))
    images = d.get('images') or []
    if not text.strip() and not images:
        return jsonify({'ok': False, 'error': '请先输入文字或上传图片'})
    st_cfg = get_ai_settings()
    if not (st_cfg['api_base'] and st_cfg['api_key'] and st_cfg['model']):
        # 未配置在线接口 → 本地规则模式兜底
        if images:
            return jsonify({'ok': False, 'error': '当前为本地规则模式，暂不支持识别图片；请在「AI 设置」配置支持视觉的在线接口，或直接输入文字'})
        obj = rule_extract(text)
        obj.update({'ok': True, 'raw': ''})
        return jsonify(obj)
    try:
        raw = call_ai(AI_SYSTEM_PROMPT, text, images)
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)})
    obj = parse_ai_json(raw)
    if not obj:
        return jsonify({'ok': False, 'error': 'AI 返回无法解析，请重试或检查模型配置', 'raw': raw[:500]})
    return jsonify({'ok': True,
                    'title': str(obj.get('title', '')).strip(),
                    'summary': str(obj.get('summary', '')).strip(),
                    'requester': str(obj.get('requester', '')).strip(),
                    'mode': 'ai',
                    'raw': raw})


@app.route('/api/ai/records', methods=['GET'])
@require_login
def ai_list_records():
    q = request.args.get('q', '').strip()
    start = request.args.get('start', '').strip()
    end = request.args.get('end', '').strip()
    try:
        page = max(1, int(request.args.get('page', 1)))
        size = max(1, min(100000, int(request.args.get('size', 20))))
    except Exception:
        page, size = 1, 20
    is_admin = (session.get('role') == 'admin') or (session.get('username') == 'admin')
    where = " WHERE 1=1" if is_admin else " WHERE operator=?"
    args = [] if is_admin else [session.get('username','')]
    if q:
        like = f"%{q}%"
        where += " AND (title LIKE ? OR summary LIKE ? OR requester LIKE ?)"
        args += [like, like, like]
    if start:
        where += " AND created_at >= ?"
        args.append(start + ' 00:00:00')
    if end:
        where += " AND created_at <= ?"
        args.append(end + ' 23:59:59')
    conn = get_conn()
    total = conn.execute("SELECT COUNT(*) c FROM ai_records" + where, args).fetchone()['c']
    rows = conn.execute(
        "SELECT * FROM ai_records" + where + " ORDER BY id DESC LIMIT ? OFFSET ?",
        args + [size, (page - 1) * size]).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows],
                    'total': total, 'page': page, 'size': size})


@app.route('/api/ai/records', methods=['POST'])
@require_login
def ai_save_record():
    d = request.get_json(force=True, silent=True) or {}
    title = str(d.get('title', '')).strip()
    summary = str(d.get('summary', '')).strip()
    requester = str(d.get('requester', '')).strip()
    raw_text = str(d.get('raw_text', ''))
    if not title and not summary:
        return jsonify({'ok': False, 'error': '工单名称与描述至少填一项'})
    now = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    conn = get_conn()
    cur = conn.execute(
        "INSERT INTO ai_records(title,summary,requester,raw_text,operator,created_at,updated_at) VALUES(?,?,?,?,?,?,?)",
        (title, summary, requester, raw_text, session['username'], now, now))
    conn.commit()
    conn.close()
    log(session['username'], '保存工作记录', f"id={cur.lastrowid}")
    return jsonify({'ok': True, 'id': cur.lastrowid})


@app.route('/api/ai/records/<int:rid>', methods=['DELETE'])
@require_login
def ai_delete_record(rid):
    is_admin = (session.get('role') == 'admin') or (session.get('username') == 'admin')
    conn = get_conn()
    if is_admin:
        row = conn.execute("SELECT id FROM ai_records WHERE id=?", (rid,)).fetchone()
    else:
        row = conn.execute("SELECT id FROM ai_records WHERE id=? AND operator=?", (rid, session.get('username',''))).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '记录不存在或无权删除'}), 404
    conn.execute("DELETE FROM ai_records WHERE id=?", (rid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除工作记录', str(rid))
    return jsonify({'ok': True})


@app.route('/api/ai/knowledge', methods=['GET'])
@require_login
def ai_list_knowledge():
    q = request.args.get('q', '').strip()
    try:
        page = max(1, int(request.args.get('page', 1)))
        size = max(1, min(100000, int(request.args.get('size', 20))))
    except Exception:
        page, size = 1, 20
    where = " WHERE 1=1"
    args = []
    if q:
        like = f"%{q}%"
        where += " AND (title LIKE ? OR content LIKE ? OR tags LIKE ?)"
        args += [like, like, like]
    conn = get_conn()
    total = conn.execute("SELECT COUNT(*) c FROM ai_knowledge" + where, args).fetchone()['c']
    rows = conn.execute(
        "SELECT * FROM ai_knowledge" + where + " ORDER BY id DESC LIMIT ? OFFSET ?",
        args + [size, (page - 1) * size]).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows],
                    'total': total, 'page': page, 'size': size})


@app.route('/api/ai/knowledge', methods=['POST'])
@require_login
def ai_save_knowledge():
    d = request.get_json(force=True, silent=True) or {}
    title = str(d.get('title', '')).strip()
    content = str(d.get('content', '')).strip()
    tags = str(d.get('tags', '')).strip()
    if not title or not content:
        return jsonify({'ok': False, 'error': '标题与内容不能为空'})
    now = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    conn = get_conn()
    cur = conn.execute(
        "INSERT INTO ai_knowledge(title,content,tags,creator,created_at,updated_at) VALUES(?,?,?,?,?,?)",
        (title, content, tags, session['username'], now, now))
    conn.commit()
    conn.close()
    log(session['username'], '新增知识', f"id={cur.lastrowid} title={title}")
    return jsonify({'ok': True, 'id': cur.lastrowid})


@app.route('/api/ai/knowledge/<int:kid>', methods=['DELETE'])
@require_login
def ai_delete_knowledge(kid):
    conn = get_conn()
    row = conn.execute("SELECT id FROM ai_knowledge WHERE id=?", (kid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '知识不存在'}), 404
    conn.execute("DELETE FROM ai_knowledge WHERE id=?", (kid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除知识', str(kid))
    return jsonify({'ok': True})



# ---------------- AI 系统数据分析（本地规则 + 只读SQL，离线可用） ----------------
def _parse_range(t):
    """解析 年份/月份/区间：返回 (start, end, label)"""
    import re as _re
    now = datetime.date.today()
    y = now.year
    m = now.month
    t = (t or '').strip()
    mm = _re.search(r'(20\d{2})\s*年\s*(\d{1,2})\s*月', t)
    if mm:
        y = int(mm.group(1)); m = int(mm.group(2))
        if 1 <= m <= 12:
            s = '%04d-%02d-01' % (y, m)
            import calendar as _cal
            last = _cal.monthrange(y, m)[1]
            return s, '%04d-%02d-%02d' % (y, m, last), '%d年%d月' % (y, m)
    mm = _re.search(r'(20\d{2})\s*年', t)
    if mm:
        y = int(mm.group(1))
        return '%04d-01-01' % y, '%04d-12-31' % y, '%d年' % y
    mm = _re.search(r'(20\d{2})[-/年](\d{1,2})(?:[-/月]|$)', t)
    if mm:
        y = int(mm.group(1)); m = int(mm.group(2))
        if 1 <= m <= 12:
            import calendar as _cal
            s = '%04d-%02d-01' % (y, m)
            last = _cal.monthrange(y, m)[1]
            return s, '%04d-%02d-%02d' % (y, m, last), '%d年%d月' % (y, m)
    if '本月' in t or '这个月' in t:
        import calendar as _cal
        s = '%04d-%02d-01' % (y, m)
        last = _cal.monthrange(y, m)[1]
        return s, '%04d-%02d-%02d' % (y, m, last), '本月'
    if '上月' in t or '上个月' in t:
        import calendar as _cal
        y2, m2 = (y - 1, 12) if m == 1 else (y, m - 1)
        s = '%04d-%02d-01' % (y2, m2)
        last = _cal.monthrange(y2, m2)[1]
        return s, '%04d-%02d-%02d' % (y2, m2, last), '上月'
    if '今年' in t or '本年度' in t:
        return '%04d-01-01' % y, '%04d-12-31' % y, '今年'
    if '去年' in t:
        return '%04d-01-01' % (y - 1), '%04d-12-31' % (y - 1), '去年'
    return '', '', ''


def analyze_question(text):
    """本地数据分析：识别意图并查询系统数据，返回结构化结果或 matched=False"""
    t = (text or '').strip()
    if not t:
        return {'matched': False}
    start, end, label = _parse_range(t)
    conn = get_conn()
    try:
        where = "1=1"
        args = []
        if start:
            where += " AND issue_time >= ?"
            args.append(start)
        if end:
            where += " AND issue_time <= ?"
            args.append(end)
        # 意图识别（按优先级）
        has_cost = any(k in t for k in ['费用', '金额', '花了', '花费', '成本', '花了多少'])
        has_stock = any(k in t for k in ['库存', '缺', '预警', '不足', '低于', '还剩', '备件数'])
        has_issue = any(k in t for k in ['发放', '领取', '领用'])
        has_flow = any(k in t for k in ['出入库', '流水', '入库', '出库'])
        has_repair = any(k in t for k in ['报修', '维修', '故障'])
        has_printer = '打印机' in t
        has_computer = any(k in t for k in ['电脑', '计算机', '办公电脑'])
        has_overview = any(k in t for k in ['总览', '整体', '全部', '汇总', '概况', '系统分析', '总共有', '多少台', '几个模块'])
        howto = any(k in t for k in ['步骤', '怎么', '如何', '操作', '方法', '教程', '调试', '怎么办', '帮我写', '给我写', '起草', '生成一条'])
        if howto and not has_cost and not has_stock and not has_overview:
            return {'matched': False}
        period = label or '全部'

        def rows_of(cols, data):
            return {'cols': cols, 'rows': data}

        if has_overview:
            p_cnt = conn.execute("SELECT COUNT(*) c FROM printers").fetchone()['c']
            pc_cnt = conn.execute("SELECT COUNT(*) c FROM computers").fetchone()['c']
            con_cnt = conn.execute("SELECT COUNT(*) c FROM consumables").fetchone()['c']
            con_sum = conn.execute("SELECT COALESCE(SUM(stock),0) q FROM consumables").fetchone()['q']
            sp_cnt = conn.execute("SELECT COUNT(*) c FROM spare_stocks").fetchone()['c']
            sp_sum = conn.execute("SELECT COALESCE(SUM(qty),0) q FROM spare_stocks").fetchone()['q']
            low = conn.execute("SELECT COUNT(*) c FROM consumables WHERE stock <= low_threshold").fetchone()['c']
            r_total = conn.execute(
                "SELECT COUNT(*) c, COALESCE(SUM(qty),0) q, COALESCE(SUM(qty*price),0) a "
                "FROM consumable_issues WHERE %s" % where, args).fetchone()
            return {'matched': True, 'type': '系统总览', 'answer':
                '系统总览：打印机 %s 台、办公电脑 %s 台；耗材 %s 种共 %s 件（其中低库存预警 %s 种）；备品备件 %s 种共 %s 件。%s发放记录 %s 次，共 %s 件，耗材费用 ¥%s。' % (
                    p_cnt, pc_cnt, con_cnt, int(con_sum), low, sp_cnt, int(sp_sum), period,
                    r_total['c'], int(r_total['q']), round(r_total['a'], 2)),
                **rows_of(['项目', '数量'], [['打印机', p_cnt], ['办公电脑', pc_cnt], ['耗材种类', con_cnt],
                                            ['耗材总件数', con_sum], ['低库存预警耗材', low],
                                            ['备品备件种类', sp_cnt], ['备品备件总件数', sp_sum],
                                            ['%s发放次数' % period, r_total['c']],
                                            ['%s耗材费用(元)' % period, round(r_total['a'], 2)]])}
        if has_cost:
            total = conn.execute(
                "SELECT COUNT(*) c, COALESCE(SUM(qty),0) q, COALESCE(SUM(qty*price),0) a "
                "FROM consumable_issues WHERE %s" % where, args).fetchone()
            depts = conn.execute(
                "SELECT department, COUNT(*) c, COALESCE(SUM(qty*price),0) a FROM consumable_issues "
                "WHERE %s AND department != '' GROUP BY department ORDER BY a DESC LIMIT 8" % where,
                args).fetchall()
            tops = conn.execute(
                "SELECT consumable, model, COUNT(*) c, COALESCE(SUM(qty*price),0) a FROM consumable_issues "
                "WHERE %s GROUP BY consumable, model ORDER BY a DESC LIMIT 5" % where, args).fetchall()
            answer = '%s耗材费用合计 ¥%s，发放 %s 次、共 %s 件。费用最高的部门：' % (
                period, round(total['a'], 2), total['c'], total['q'])
            answer += '、'.join('%s ¥%s' % (d['department'] or '未知', round(d['a'], 2)) for d in depts[:3]) if depts else '（无记录）'
            return {'matched': True, 'type': '费用统计', 'answer': answer,
                    **rows_of(['部门', '次数', '金额(元)'], [[d['department'], d['c'], round(d['a'], 2)] for d in depts]),
                    'extra': {'title': '耗材TOP', 'cols': ['耗材', '打印机型号', '次数', '金额(元)'],
                              'rows': [[x['consumable'], x['model'], x['c'], round(x['a'], 2)] for x in tops]}}
        if has_stock:
            lows = conn.execute(
                "SELECT model, category, stock, low_threshold FROM consumables "
                "WHERE stock <= low_threshold ORDER BY (stock - low_threshold) LIMIT 15").fetchall()
            cats = conn.execute(
                "SELECT category, COUNT(*) c, COALESCE(SUM(stock),0) q FROM consumables GROUP BY category").fetchall()
            if lows:
                answer = '低库存预警耗材 %s 种：' % len(lows)
                answer += '、'.join('%s(余%s/阈值%s)' % (r['model'], r['stock'], r['low_threshold']) for r in lows[:6])
                answer += '。' if len(lows) <= 6 else ' 等。'
            else:
                answer = '当前所有耗材库存均在阈值以上，无低库存预警。'
            answer += ' 各类别库存：' + '、'.join('%s %s件' % (r['category'] or '未分类', r['q']) for r in cats) if cats else answer
            return {'matched': True, 'type': '库存分析', 'answer': answer,
                    **rows_of(['耗材型号', '品类', '当前库存', '预警阈值'], [[r['model'], r['category'], r['stock'], r['low_threshold']] for r in lows])}
        if has_flow:
            f_in = conn.execute("SELECT COUNT(*) c, COALESCE(SUM(qty),0) q FROM spare_flows WHERE type IN ('in','入库')").fetchone()
            f_out = conn.execute("SELECT COUNT(*) c, COALESCE(SUM(qty),0) q FROM spare_flows WHERE type IN ('out','出库')").fetchone()
            latest = conn.execute("SELECT * FROM spare_flows ORDER BY id DESC LIMIT 8").fetchall()
            _ft = lambda v: {'in': '入库', 'out': '出库', 'move': '移库'}.get(v, v)
            answer = '备品备件流水：累计入库 %s 条 %s 件、出库 %s 条 %s 件；最近流水：' % (
                f_in['c'], f_in['q'], f_out['c'], f_out['q'])
            answer += '、'.join('%s %s %s×%s' % (_ft(r['type']), r['name'] or r['brand'], r['qty'], r['unit'] or '') for r in latest[:5])
            return {'matched': True, 'type': '流水统计', 'answer': answer,
                    **rows_of(['类型', '物资', '库位', '数量', '时间'], [[_ft(r['type']), r['name'] or r['brand'], r['location'], r['qty'], r['record_time']] for r in latest])}
        if has_computer:
            pc_cnt = conn.execute("SELECT COUNT(*) c FROM computers").fetchone()['c']
            crep = conn.execute(
                "SELECT COALESCE(c.model,'未知') model, COUNT(*) c, COALESCE(SUM(r.cost),0) a "
                "FROM computer_repairs r LEFT JOIN computers c ON r.computer_id = c.id "
                "GROUP BY c.model ORDER BY c DESC LIMIT 8").fetchall()
            answer = '办公电脑共 %s 台。维修记录（按型号）：' % pc_cnt
            answer += '、'.join('%s %s次(费用¥%s)' % (r['model'] or '未知', r['c'], round(r['a'], 2)) for r in crep[:5]) if crep else '（暂无维修记录）'
            return {'matched': True, 'type': '电脑统计', 'answer': answer,
                    **rows_of(['电脑型号', '维修次数', '维修费用(元)'], [[r['model'], r['c'], round(r['a'], 2)] for r in crep])}
        if has_issue:
            total = conn.execute(
                "SELECT COUNT(*) c, COALESCE(SUM(qty),0) q FROM consumable_issues WHERE %s" % where, args).fetchone()
            depts = conn.execute(
                "SELECT department, COUNT(*) c FROM consumable_issues WHERE %s AND department != '' "
                "GROUP BY department ORDER BY c DESC LIMIT 8" % where, args).fetchall()
            answer = '%s共发放 %s 次、%s 件。发放最多的部门：%s' % (
                period, total['c'], int(total['q']),
                '、'.join('%s %s次' % (d['department'], d['c']) for d in depts[:5]) if depts else '（无记录）')
            return {'matched': True, 'type': '发放统计', 'answer': answer,
                    **rows_of(['部门', '发放次数'], [[d['department'], d['c']] for d in depts])}
        if has_repair or has_printer:
            p_cnt = conn.execute("SELECT COUNT(*) c FROM printers").fetchone()['c']
            repairs = conn.execute(
                "SELECT COALESCE(p.model,'未知') model, COUNT(*) c, COALESCE(SUM(r.cost),0) a "
                "FROM printer_repairs r LEFT JOIN printers p ON r.printer_id = p.id "
                "GROUP BY p.model ORDER BY c DESC LIMIT 8").fetchall()
            answer = '打印机共 %s 台。报修记录（按型号）：' % p_cnt
            answer += '、'.join('%s %s次(费用¥%s)' % (r['model'] or '未知', r['c'], round(r['a'], 2)) for r in repairs[:5]) if repairs else '（暂无报修记录）'
            return {'matched': True, 'type': '打印机/维修统计', 'answer': answer,
                    **rows_of(['打印机型号', '报修次数', '维修费用(元)'], [[r['model'], r['c'], round(r['a'], 2)] for r in repairs])}
        return {'matched': False}
    finally:
        conn.close()


@app.route('/api/ai/analyze', methods=['POST'])
@require_login
def ai_analyze():
    d = request.get_json(force=True, silent=True) or {}
    text = str(d.get('text', ''))
    r = analyze_question(text)
    if not r.get('matched'):
        return jsonify({'ok': True, 'matched': False})
    log(session['username'], '系统数据分析', text[:60])
    return jsonify({'ok': True, 'matched': True, 'type': r.get('type'),
                    'answer': r.get('answer'), 'cols': r.get('cols', []),
                    'rows': r.get('rows', []), 'extra': r.get('extra')})


# ---------------- AI 功能修改需求登记 ----------------
@app.route('/api/ai/requests', methods=['GET'])
@require_login
def ai_list_requests():
    status = request.args.get('status', '').strip()
    where = " WHERE 1=1"
    args = []
    if status:
        where += " AND status=?"
        args.append(status)
    conn = get_conn()
    rows = conn.execute("SELECT * FROM ai_requests" + where + " ORDER BY id DESC", args).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows]})


@app.route('/api/ai/requests', methods=['POST'])
@require_login
def ai_save_request():
    d = request.get_json(force=True, silent=True) or {}
    title = str(d.get('title', '')).strip()
    detail = str(d.get('detail', '')).strip()
    if not title:
        return jsonify({'ok': False, 'error': '需求标题不能为空'})
    now = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    conn = get_conn()
    cur = conn.execute(
        "INSERT INTO ai_requests(title,detail,status,created_by,created_at,updated_at) VALUES(?,?,?,?,?,?)",
        (title, detail, '待处理', session['username'], now, now))
    conn.commit()
    conn.close()
    log(session['username'], '登记功能修改需求', f"id={cur.lastrowid} {title[:40]}")
    return jsonify({'ok': True, 'id': cur.lastrowid})


@app.route('/api/ai/requests/<int:rid>/status', methods=['POST'])
@require_login
def ai_request_status(rid):
    d = request.get_json(force=True, silent=True) or {}
    status = str(d.get('status', '')).strip()
    if status not in ('待处理', '处理中', '已完成', '已关闭'):
        return jsonify({'ok': False, 'error': '状态无效'})
    now = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    conn = get_conn()
    row = conn.execute("SELECT id FROM ai_requests WHERE id=?", (rid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '需求不存在'}), 404
    conn.execute("UPDATE ai_requests SET status=?, updated_at=? WHERE id=?", (status, now, rid))
    conn.commit()
    conn.close()
    log(session['username'], '更新功能需求状态', f"id={rid} -> {status}")
    return jsonify({'ok': True})


@app.route('/api/ai/requests/<int:rid>', methods=['DELETE'])
@require_login
def ai_delete_request(rid):
    conn = get_conn()
    row = conn.execute("SELECT id FROM ai_requests WHERE id=?", (rid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '需求不存在'}), 404
    conn.execute("DELETE FROM ai_requests WHERE id=?", (rid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除功能需求', str(rid))
    return jsonify({'ok': True})


# ---------------- AI 对话（像豆包网页版一样一问一答） ----------------
def call_chat(messages):
    """调用 OpenAI 兼容接口进行多轮对话"""
    st = get_ai_settings()
    if not (st['api_base'] and st['api_key'] and st['model']):
        raise RuntimeError('AI 接口未配置，请点击右上角「⚙️ AI 设置」填写接口信息（可一键填入豆包预设）')
    url = st['api_base'].rstrip('/') + '/chat/completions'
    body = json.dumps({
        'model': st['model'],
        'messages': messages,
        'temperature': 0.7,
        'max_tokens': 2048
    }).encode()
    req = urllib.request.Request(url, data=body, method='POST')
    req.add_header('Content-Type', 'application/json')
    req.add_header('Authorization', 'Bearer ' + st['api_key'])
    try:
        r = urllib.request.urlopen(req, timeout=120)
        resp = json.loads(r.read().decode('utf-8', 'ignore'))
    except urllib.error.HTTPError as e:
        raise RuntimeError('AI 接口错误 %s: %s' % (e.code, e.read().decode('utf-8', 'ignore')[:300]))
    except Exception as e:
        raise RuntimeError('AI 接口调用失败: %s' % e)
    try:
        return resp['choices'][0]['message']['content']
    except Exception:
        raise RuntimeError('AI 返回格式异常: %s' % str(resp)[:300])


def collect_system_context():
    """汇总整个台账系统的关键数据，作为 AI 的系统上下文"""
    conn = get_conn()
    c = conn.cursor()
    parts = []
    try:
        # 当前用户身份（跨会话记忆）
        uname = session.get('username', '')
        real = c.execute("SELECT real_name FROM users WHERE username=?", (uname,)).fetchone()
        who = real['real_name'] if real else uname
        parts.append(f"当前登录用户：{who}（账号 {uname}），是信息中心桌面运维人员")

        # 打印机
        n_printer = c.execute("SELECT COUNT(*) FROM printers").fetchone()[0]
        n_printer_bad = c.execute("SELECT COUNT(*) FROM printers WHERE status LIKE '%故障%' OR status LIKE '%维修%'").fetchone()[0]
        parts.append(f"打印机总数 {n_printer} 台（其中故障/维修中 {n_printer_bad} 台）")

        # 耗材
        n_cons = c.execute("SELECT COUNT(*) FROM consumables").fetchone()[0]
        n_issue = c.execute("SELECT COUNT(*) FROM consumable_issues").fetchone()[0]
        n_adjust = c.execute("SELECT COUNT(*) FROM consumable_adjusts").fetchone()[0]
        parts.append(f"耗材种类 {n_cons} 种，累计发放记录 {n_issue} 条，调整记录 {n_adjust} 条")
        low = c.execute("SELECT model, stock FROM consumables WHERE stock < low_threshold AND low_threshold > 0 ORDER BY stock ASC LIMIT 10").fetchall()
        if low:
            parts.append("⚠️ 低库存耗材：" + "；".join(f"{r['model']}(剩{r['stock']})" for r in low))

        # 备品备件
        n_sp = c.execute("SELECT COUNT(*) FROM spare_items").fetchone()[0]
        n_sp_stock = c.execute("SELECT COUNT(*) FROM spare_stocks WHERE qty > 0").fetchone()[0]
        n_sp_flow = c.execute("SELECT COUNT(*) FROM spare_flows").fetchone()[0]
        parts.append(f"备品备件物资 {n_sp} 种，有库存 {n_sp_stock} 种，出入库流水 {n_sp_flow} 条")

        # 电脑优化
        n_pc = c.execute("SELECT COUNT(*) FROM pc_upgrades").fetchone()[0]
        n_parts = c.execute("SELECT COUNT(*) FROM pc_parts").fetchone()[0]
        parts.append(f"办公电脑优化记录 {n_pc} 条，优化备件种类 {n_parts} 种")

        # 办公电脑
        n_comp = c.execute("SELECT COUNT(*) FROM computers").fetchone()[0]
        parts.append(f"办公电脑台账 {n_comp} 台")

        # 最近发放（近30天）
        try:
            recent = c.execute("SELECT COUNT(*) FROM consumable_issues WHERE issue_date >= date('now','-30 day')").fetchone()[0]
            parts.append(f"近30天耗材发放 {recent} 条")
        except Exception:
            pass

        # 价格表
        n_price = c.execute("SELECT COUNT(*) FROM price_list").fetchone()[0]
        parts.append(f"价格库 {n_price} 条")

        # 知识库（业务规则、部门结构、流程规范）
        kb_rows = c.execute("SELECT title, content, tags FROM ai_knowledge ORDER BY id DESC LIMIT 30").fetchall()
        if kb_rows:
            parts.append("")
            parts.append("【业务知识库（必须严格遵守）】")
            for r in kb_rows:
                line = f"· {r['title']}"
                if r['tags']: line += f"［{r['tags']}］"
                line += f"：{r['content']}"
                parts.append(line)
    except Exception as e:
        parts.append(f"（部分数据读取失败：{e}）")
    finally:
        conn.close()
    return "【系统实时数据概览】\n" + "\n".join("· " + p for p in parts)


@app.route('/api/ai/chat', methods=['POST'])
@require_login
def ai_chat():
    d = request.get_json(force=True, silent=True) or {}
    messages = d.get('messages') or []
    if not isinstance(messages, list) or not messages:
        return jsonify({'ok': False, 'error': '消息不能为空'})
    # 只保留最近 20 轮，防止上下文过长
    history = [m for m in messages[-20:] if isinstance(m, dict) and m.get('role') in ('user', 'assistant') and str(m.get('content', '')).strip()]
    if not history:
        return jsonify({'ok': False, 'error': '消息不能为空'})
    sys_ctx = collect_system_context()
    sys_prompt = ("你是信息中心的 AI 运维助手，负责管理「设备台账管理系统」。"
                  "下面是系统实时数据概览，回答时优先基于这些数据，不要编造数字；"
                  "用户问具体明细时可以让他在系统对应模块查看。使用简体中文，回答简洁有条理，必要时用列表/表格。\n\n"
                  + sys_ctx)
    msgs = [{'role': 'system', 'content': sys_prompt}]
    for m in history:
        content = str(m.get('content', ''))[:4000]
        # 附件信息拼到消息里（当前模型是文本模型，看不到图，但让它知道用户发了什么）
        atts = m.get('attachments') or []
        if atts:
            desc = []
            for a in atts:
                if isinstance(a, dict):
                    nm = a.get('name') or a.get('url') or '文件'
                    sz = a.get('size') or 0
                    szs = f"{sz/1024:.0f}KB" if sz < 1048576 else f"{sz/1048576:.1f}MB"
                    if a.get('is_image'):
                        desc.append(f"图片附件：{nm}（{szs}）")
                    else:
                        desc.append(f"文件附件：{nm}（{szs}）")
            if desc:
                content = content + "\n[用户上传了" + "；".join(desc) + "]（当前为文本模型，无法直接查看图片内容，请提示用户用文字描述图片里的关键信息，或根据文件名/上下文推断）"
        msgs.append({'role': m['role'], 'content': content})
    try:
        reply = call_chat(msgs)
    except Exception as e:
        return jsonify({'ok': False, 'error': str(e)})
    return jsonify({'ok': True, 'reply': (reply or '').strip()})


# ---------------- AI 会话服务端存储 + 文件上传 ----------------
AI_UPLOAD_DIR = os.path.join(BASE_DIR, 'data', 'uploads', 'ai')
os.makedirs(AI_UPLOAD_DIR, exist_ok=True)


@app.route('/api/ai/upload', methods=['POST'])
@require_login
def ai_upload():
    f = request.files.get('file')
    if not f:
        return jsonify({'ok': False, 'error': '未收到文件'})
    fname = secure_filename(f.filename or 'upload')
    # 扩展名白名单：仅允许图片与常见文档，拒绝 html/js/svg 等可被浏览器执行的类型（防存储型XSS）
    ext = os.path.splitext(fname)[1].lower()
    if ext not in ('.png', '.jpg', '.jpeg', '.gif', '.bmp', '.webp',
                   '.xlsx', '.xls', '.docx', '.doc', '.pdf', '.txt', '.csv', '.md'):
        return jsonify({'ok': False, 'error': '文件类型不允许（仅图片及常见文档）'}), 400
    ts = datetime.datetime.now().strftime('%Y%m%d_%H%M%S_')
    save_name = ts + fname
    save_path = os.path.join(AI_UPLOAD_DIR, save_name)
    f.save(save_path)
    size = os.path.getsize(save_path)
    is_img = ext in ('.png', '.jpg', '.jpeg', '.gif', '.bmp', '.webp')
    return jsonify({'ok': True, 'url': '/uploads/ai/' + save_name, 'name': fname, 'size': size, 'is_image': is_img})


@app.route('/api/ai/convs', methods=['GET'])
@require_login
def ai_list_convs():
    uname = session.get('username', '')
    is_admin = (session.get('role') == 'admin') or (uname == 'admin')
    conn = get_conn()
    if is_admin:
        rows = conn.execute("SELECT id, title, created_by, created_at, updated_at FROM ai_conversations ORDER BY updated_at DESC").fetchall()
    else:
        rows = conn.execute("SELECT id, title, created_at, updated_at FROM ai_conversations WHERE created_by=? ORDER BY updated_at DESC", (uname,)).fetchall()
    conn.close()
    return jsonify({'ok': True, 'items': [dict(r) for r in rows]})


def _own_conv(conn, cid):
    """校验会话归属：admin 可访问任何人，普通用户只能访问自己的"""
    uname = session.get('username', '')
    is_admin = (session.get('role') == 'admin') or (uname == 'admin')
    r = conn.execute("SELECT * FROM ai_conversations WHERE id=?", (cid,)).fetchone()
    if not r: return None
    if is_admin or r['created_by'] == uname:
        return r
    return None


@app.route('/api/ai/convs', methods=['POST'])
@require_login
def ai_create_conv():
    d = request.get_json(force=True, silent=True) or {}
    title = (d.get('title') or '新对话').strip()
    uname = session.get('username', '')
    now = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    conn = get_conn()
    cur = conn.execute("INSERT INTO ai_conversations(title,created_by,created_at,updated_at) VALUES(?,?,?,?)",
                       (title, uname, now, now))
    conn.commit()
    cid = cur.lastrowid
    conn.close()
    return jsonify({'ok': True, 'id': cid})


@app.route('/api/ai/convs/<int:cid>', methods=['DELETE'])
@require_login
def ai_del_conv(cid):
    conn = get_conn()
    if not _own_conv(conn, cid):
        conn.close(); return jsonify({'ok': False, 'error': '无权删除该对话'}), 403
    conn.execute("DELETE FROM ai_messages WHERE conv_id=?", (cid,))
    conn.execute("DELETE FROM ai_conversations WHERE id=?", (cid,))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})


@app.route('/api/ai/convs/<int:cid>/title', methods=['POST'])
@require_login
def ai_rename_conv(cid):
    d = request.get_json(force=True, silent=True) or {}
    title = (d.get('title') or '').strip()
    if not title:
        return jsonify({'ok': False, 'error': '标题不能为空'})
    now = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    conn = get_conn()
    if not _own_conv(conn, cid):
        conn.close(); return jsonify({'ok': False, 'error': '无权操作该对话'}), 403
    conn.execute("UPDATE ai_conversations SET title=?, updated_at=? WHERE id=?", (title, now, cid))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})


@app.route('/api/ai/convs/<int:cid>/messages', methods=['GET'])
@require_login
def ai_list_messages(cid):
    conn = get_conn()
    if not _own_conv(conn, cid):
        conn.close(); return jsonify({'ok': False, 'error': '无权查看该对话'}), 403
    rows = conn.execute("SELECT id, role, content, attachments, created_at FROM ai_messages WHERE conv_id=? ORDER BY id ASC", (cid,)).fetchall()
    conn.close()
    out = []
    for r in rows:
        att = []
        try:
            att = json.loads(r['attachments'] or '[]')
        except Exception:
            pass
        out.append({'role': r['role'], 'content': r['content'], 'attachments': att, 'ts': r['created_at']})
    return jsonify({'ok': True, 'items': out})


@app.route('/api/ai/convs/<int:cid>/messages', methods=['POST'])
@require_login
def ai_add_message(cid):
    d = request.get_json(force=True, silent=True) or {}
    role = d.get('role', 'user')
    content = d.get('content', '')
    attachments = d.get('attachments') or []
    now = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    conn = get_conn()
    if not _own_conv(conn, cid):
        conn.close(); return jsonify({'ok': False, 'error': '无权操作该对话'}), 403
    conn.execute("INSERT INTO ai_messages(conv_id,role,content,attachments,created_at) VALUES(?,?,?,?,?)",
                 (cid, role, content, json.dumps(attachments, ensure_ascii=False), now))
    conn.execute("UPDATE ai_conversations SET updated_at=? WHERE id=?", (now, cid))
    row = conn.execute("SELECT title FROM ai_conversations WHERE id=?", (cid,)).fetchone()
    if row and (row['title'] or '').strip() in ('', '新对话') and role == 'user':
        conn.execute("UPDATE ai_conversations SET title=? WHERE id=?", (content[:18], cid))
    conn.commit()
    conn.close()
    return jsonify({'ok': True})


# ---------------- AI 模块扩展：批量导入 / 导出 / 统计 ----------------
def _parse_dt(v):
    """解析导入时间：支持 YYYY-MM-DD HH:MM:SS / YYYY-MM-DD / 空"""
    v = (v or '').strip()
    if not v:
        return datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    for fmt in ('%Y-%m-%d %H:%M:%S', '%Y-%m-%d %H:%M', '%Y/%m/%d %H:%M:%S', '%Y/%m/%d'):
        try:
            return datetime.datetime.strptime(v, fmt).strftime('%Y-%m-%d %H:%M:%S')
        except Exception:
            pass
    return datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')


@app.route('/api/ai/records/import', methods=['POST'])
@require_login
def ai_import_records():
    d = request.get_json(force=True, silent=True) or {}
    rows = d.get('rows') or []
    if not isinstance(rows, list) or not rows:
        return jsonify({'ok': False, 'error': '没有可导入的数据'})
    if len(rows) > 5000:
        return jsonify({'ok': False, 'error': '单次最多导入 5000 条'})
    now = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    conn = get_conn()
    n = 0
    bad = 0
    for r in rows:
        if not isinstance(r, dict):
            bad += 1
            continue
        title = str(r.get('title', '')).strip()
        summary = str(r.get('summary', '')).strip()
        requester = str(r.get('requester', '')).strip()
        raw_text = str(r.get('raw_text', ''))
        if not title and not summary:
            bad += 1
            continue
        created_at = _parse_dt(r.get('created_at'))
        conn.execute(
            "INSERT INTO ai_records(title,summary,requester,raw_text,operator,created_at,updated_at) VALUES(?,?,?,?,?,?,?)",
            (title, summary, requester, raw_text, session['username'], created_at, now))
        n += 1
    conn.commit()
    conn.close()
    log(session['username'], '批量导入工作记录', f"成功{n}条 跳过{bad}条")
    return jsonify({'ok': True, 'imported': n, 'skipped': bad})


@app.route('/api/ai/knowledge/import', methods=['POST'])
@require_login
def ai_import_knowledge():
    d = request.get_json(force=True, silent=True) or {}
    rows = d.get('rows') or []
    if not isinstance(rows, list) or not rows:
        return jsonify({'ok': False, 'error': '没有可导入的数据'})
    if len(rows) > 5000:
        return jsonify({'ok': False, 'error': '单次最多导入 5000 条'})
    now = datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')
    conn = get_conn()
    n = 0
    bad = 0
    for r in rows:
        if not isinstance(r, dict):
            bad += 1
            continue
        title = str(r.get('title', '')).strip()
        content = str(r.get('content', '')).strip()
        tags = str(r.get('tags', '')).strip()
        if not title or not content:
            bad += 1
            continue
        created_at = _parse_dt(r.get('created_at'))
        conn.execute(
            "INSERT INTO ai_knowledge(title,content,tags,creator,created_at,updated_at) VALUES(?,?,?,?,?,?)",
            (title, content, tags, session['username'], created_at, now))
        n += 1
    conn.commit()
    conn.close()
    log(session['username'], '批量导入知识', f"成功{n}条 跳过{bad}条")
    return jsonify({'ok': True, 'imported': n, 'skipped': bad})


def _csv_safe(v):
    """CSV 公式注入防御：单元格以 = + - @ 或制表符/回车开头时，前置单引号避免被 Excel 当作公式执行"""
    s = '' if v is None else str(v)
    if s and s[0] in ('=', '+', '-', '@', '\t', '\r'):
        return "'" + s
    return s


def _csv_download(name, header, rows_list):
    import csv
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(header)
    for r in rows_list:
        w.writerow([_csv_safe(c) for c in r])
    data = buf.getvalue().encode('utf-8-sig')  # BOM 便于 Excel 识别
    from flask import Response as _Resp
    return _Resp(data, mimetype='text/csv; charset=utf-8',
                 headers={'Content-Disposition': 'attachment; filename=' + name})


@app.route('/api/ai/records/export')
@require_login
def ai_export_records():
    q = request.args.get('q', '').strip()
    start = request.args.get('start', '').strip()
    end = request.args.get('end', '').strip()
    is_admin = (session.get('role') == 'admin') or (session.get('username') == 'admin')
    where = " WHERE 1=1" if is_admin else " WHERE operator=?"
    args = [] if is_admin else [session.get('username','')]
    if q:
        like = '%' + q + '%'
        where += " AND (title LIKE ? OR summary LIKE ? OR requester LIKE ?)"
        args += [like, like, like]
    if start:
        where += " AND created_at >= ?"
        args.append(start + ' 00:00:00')
    if end:
        where += " AND created_at <= ?"
        args.append(end + ' 23:59:59')
    conn = get_conn()
    rows = conn.execute(
        "SELECT title,summary,requester,raw_text,operator,created_at FROM ai_records" + where + " ORDER BY created_at DESC, id DESC",
        args).fetchall()
    conn.close()
    name = 'ai_records_' + datetime.datetime.now().strftime('%Y%m%d') + '.csv'
    return _csv_download(name, ['工单名称', '工单描述', '需求人', '原始信息', '操作人', '记录时间'],
                         [[r['title'], r['summary'], r['requester'], r['raw_text'], r['operator'], r['created_at']] for r in rows])


@app.route('/api/ai/knowledge/export')
@require_login
def ai_export_knowledge():
    q = request.args.get('q', '').strip()
    where = " WHERE 1=1"
    args = []
    if q:
        like = '%' + q + '%'
        where += " AND (title LIKE ? OR content LIKE ? OR tags LIKE ?)"
        args += [like, like, like]
    conn = get_conn()
    rows = conn.execute(
        "SELECT title,content,tags,creator,created_at FROM ai_knowledge" + where + " ORDER BY created_at DESC, id DESC",
        args).fetchall()
    conn.close()
    name = 'ai_knowledge_' + datetime.datetime.now().strftime('%Y%m%d') + '.csv'
    return _csv_download(name, ['标题', '内容', '标签', '创建人', '创建时间'],
                         [[r['title'], r['content'], r['tags'], r['creator'], r['created_at']] for r in rows])


@app.route('/api/ai/records/stats')
@require_login
def ai_stats():
    from collections import Counter
    period = request.args.get('period', 'monthly')
    if period not in ('daily', 'weekly', 'monthly'):
        period = 'monthly'
    conn = get_conn()
    uname = session.get('username','')
    is_admin = (session.get('role') == 'admin') or (uname == 'admin')
    op_cond = '' if is_admin else ' AND operator=?'
    op_args = [] if is_admin else [uname]
    total = conn.execute("SELECT COUNT(*) c FROM ai_records WHERE 1=1" + op_cond, op_args).fetchone()['c']
    # 趋势
    if period == 'daily':
        rows = conn.execute("SELECT strftime('%Y-%m-%d', created_at) k, COUNT(*) c FROM ai_records WHERE created_at != ''" + op_cond + " GROUP BY k", op_args).fetchall()
        m = {r['k']: r['c'] for r in rows}
        trend = []
        today = datetime.date.today()
        for i in range(29, -1, -1):
            day = (today - datetime.timedelta(days=i)).strftime('%Y-%m-%d')
            trend.append({'label': day[5:], 'count': m.get(day, 0)})
    elif period == 'weekly':
        rows = conn.execute("SELECT strftime('%Y-W%W', created_at) k, COUNT(*) c FROM ai_records WHERE created_at != ''" + op_cond + " GROUP BY k", op_args).fetchall()
        m = {r['k']: r['c'] for r in rows}
        trend = []
        today = datetime.date.today()
        for i in range(11, -1, -1):
            d0 = today - datetime.timedelta(weeks=i)
            y, wk, _ = d0.isocalendar()
            key = '%d-W%02d' % (y, wk)
            trend.append({'label': key, 'count': m.get(key, 0)})
    else:
        rows = conn.execute("SELECT strftime('%Y-%m', created_at) k, COUNT(*) c FROM ai_records WHERE created_at != ''" + op_cond + " GROUP BY k", op_args).fetchall()
        m = {r['k']: r['c'] for r in rows}
        trend = []
        base = datetime.date.today().year * 12 + datetime.date.today().month - 1
        for i in range(11, -1, -1):
            v = base - i
            key = '%04d-%02d' % (v // 12, v % 12 + 1)
            trend.append({'label': key[2:], 'count': m.get(key, 0)})
    # 分布：从 title 拆 桌面--区域--部门--事项--编号
    c_type, c_area, c_dept = Counter(), Counter(), Counter()
    rows = conn.execute("SELECT title FROM ai_records WHERE title IS NOT NULL AND title != ''" + op_cond, op_args).fetchall()
    for r in rows:
        parts = (r['title'] or '').split('--')
        c_type[parts[0] if parts and parts[0] else '未知'] += 1
        c_area[parts[1] if len(parts) > 1 and parts[1] else '未知'] += 1
        c_dept[parts[2] if len(parts) > 2 and parts[2] else '未知'] += 1
    conn.close()
    return jsonify({'ok': True, 'total': total, 'period': period,
                    'trend': trend,
                    'by_type': [{'name': k, 'count': v} for k, v in c_type.most_common(20)],
                    'by_area': [{'name': k, 'count': v} for k, v in c_area.most_common(20)],
                    'by_dept': [{'name': k, 'count': v} for k, v in c_dept.most_common(20)]})


# ---------------- 图片上传（记事本富文本用） ----------------
UPLOAD_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), 'data', 'uploads')
os.makedirs(UPLOAD_DIR, exist_ok=True)
app.config['MAX_CONTENT_LENGTH'] = 20 * 1024 * 1024  # 20MB

@app.route('/api/upload', methods=['POST'])
@require_login
def upload_image():
    f = request.files.get('file') or request.files.get('image')
    if not f:
        return jsonify({'ok': False, 'error': '没有文件'}), 400
    import uuid
    ext = (f.filename or '').rsplit('.', 1)[-1].lower()
    if ext not in ('png', 'jpg', 'jpeg', 'gif', 'webp', 'bmp'):
        return jsonify({'ok': False, 'error': '仅允许图片文件'}), 400
    fname = datetime.datetime.now().strftime('%Y%m%d_%H%M%S_') + uuid.uuid4().hex[:8] + '.' + ext
    f.save(os.path.join(UPLOAD_DIR, fname))
    return jsonify({'ok': True, 'url': '/uploads/' + fname})


@app.route('/uploads/<path:name>')
@require_login
def uploaded_file(name):
    if '..' in name or name.startswith('/'):
        return jsonify({'ok': False, 'error': '非法路径'}), 400
    return send_from_directory(UPLOAD_DIR, name)


# ---------------- 静态页面 ----------------
@app.route('/')
def index():
    resp = send_from_directory(PUBLIC_DIR, 'index.html')
    resp.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0'
    resp.headers['Pragma'] = 'no-cache'
    resp.headers['Expires'] = '0'
    return resp


@app.route('/uploads/ai/<path:name>')
@require_login
def ai_uploaded_file(name):
    if '..' in name or name.startswith('/'):
        return jsonify({'ok': False, 'error': '非法路径'}), 400
    return send_from_directory(AI_UPLOAD_DIR, name)


@app.route('/<path:name>')
def static_files(name):
    resp = send_from_directory(PUBLIC_DIR, name)
    if name.endswith(('.html', '.js', '.css')):
        resp.headers['Cache-Control'] = 'no-store, no-cache, must-revalidate, max-age=0'
        resp.headers['Pragma'] = 'no-cache'
        resp.headers['Expires'] = '0'
    return resp


# ---------------- 行政办公电脑配件更换记录 ----------------
@app.route('/api/pc-replace', methods=['GET'])
@require_login
def list_pc_replace():
    q = request.args.get('q', '').strip()
    sql = "SELECT * FROM pc_replace WHERE 1=1"
    args = []
    if q:
        sql += " AND (user_name LIKE ? OR department LIKE ? OR location LIKE ? OR part_name LIKE ? OR monitor_sn LIKE ? OR remark LIKE ?)"
        like = f"%{q}%"
        args += [like] * 6
    sql += " ORDER BY id DESC"
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows]})


@app.route('/api/pc-replace/<int:rid>', methods=['GET'])
@require_login
def get_pc_replace(rid):
    conn = get_conn()
    r = conn.execute("SELECT * FROM pc_replace WHERE id=?", (rid,)).fetchone()
    conn.close()
    if not r:
        return jsonify({'ok': False, 'error': '记录不存在'}), 404
    return jsonify({'ok': True, 'data': dict(r)})


def parse_part_qty(qty):
    """解析更换记录备件数量：'电源:2'→2、'1'→1；空/横线/非数字串（如 BVT64309004）→0（不联动）"""
    s = str(qty or '').strip()
    if not s or s in ('-', '—'):
        return 0
    if ':' in s:
        tail = s.split(':', 1)[1].strip()
        return int(tail) if tail.isdigit() else 0
    return int(s) if s.isdigit() else 0


def find_part_stock(conn, part_name, storage):
    """按配件名+存放位置精确匹配配件库存；存放位置为空且名称唯一时也匹配，否则返回 None"""
    name = str(part_name or '').strip()
    if not name:
        return None
    st = str(storage or '').strip()
    if st:
        return conn.execute("SELECT * FROM pc_part_stock WHERE name=? AND storage_location=? LIMIT 1", (name, st)).fetchone()
    rows = conn.execute("SELECT * FROM pc_part_stock WHERE name=? LIMIT 2", (name,)).fetchall()
    return rows[0] if len(rows) == 1 else None


@app.route('/api/pc-replace', methods=['POST'])
@require_login
def add_pc_replace():
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    part_name = str(d.get('part_name', '')).strip()
    storage = str(d.get('storage', '')).strip()
    qty = parse_part_qty(d.get('part_qty'))
    warn = ''
    if qty > 0:
        part = find_part_stock(conn, part_name, storage)
        if not part:
            warn = '未匹配到配件库存条目，本次未扣减库存'
        elif int(num(part['stock'])) < qty:
            conn.close()
            return jsonify({'ok': False, 'error': '配件库存不足：%s 当前库存 %s，需 %s 个' % (part['name'], part['stock'], qty)})
        else:
            conn.execute("UPDATE pc_part_stock SET stock=stock-?,updated_at=? WHERE id=?", (qty, now_str(), part['id']))
            warn = '已自动扣减库存：%s ×%s（%s）' % (part['name'], qty, part['storage_location'])
    cur = conn.execute(
        "INSERT INTO pc_replace(user_name,department,location,part_name,replaced,returned,monitor_sn,part_qty,storage,remark,record_time,operator,created_at,updated_at) "
        "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
        (str(d.get('user_name', '')).strip(), str(d.get('department', '')).strip(),
         str(d.get('location', '')).strip(), part_name,
         str(d.get('replaced', '')).strip() or '是', str(d.get('returned', '')).strip() or '否',
         str(d.get('monitor_sn', '')).strip(), str(d.get('part_qty', '')).strip(),
         storage, str(d.get('remark', '')).strip(),
         d.get('record_time') or now_str(), session['username'], now_str(), now_str()))
    conn.commit()
    conn.close()
    log(session['username'], '新增配件更换记录', str(cur.lastrowid))
    return jsonify({'ok': True, 'id': cur.lastrowid, 'warn': warn})


@app.route('/api/pc-replace/<int:rid>', methods=['PUT'])
@require_login
def update_pc_replace(rid):
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    old = conn.execute("SELECT * FROM pc_replace WHERE id=?", (rid,)).fetchone()
    if not old:
        conn.close()
        return jsonify({'ok': False, 'error': '更换记录不存在'}), 404
    new_name = str(d.get('part_name', '')).strip()
    new_storage = str(d.get('storage', '')).strip()
    new_qty = parse_part_qty(d.get('part_qty'))
    old_qty = parse_part_qty(old['part_qty'])
    old_part = find_part_stock(conn, old['part_name'], old['storage'])
    new_part = find_part_stock(conn, new_name, new_storage)
    warn = ''
    if old_part and old_qty > 0:   # 先回补原扣减
        conn.execute("UPDATE pc_part_stock SET stock=stock+?,updated_at=? WHERE id=?", (old_qty, now_str(), old_part['id']))
        warn = '已回补库存：%s ×%s' % (old_part['name'], old_qty)
        new_part = find_part_stock(conn, new_name, new_storage)   # 同一条目时刷新最新库存再判新扣减
    if new_qty > 0:
        if not new_part:
            warn = (warn + '；' if warn else '') + '未匹配到配件库存条目，本次未扣减库存'
        elif int(num(new_part['stock'])) < new_qty:
            conn.close()   # 未提交 → 回补一并回滚
            return jsonify({'ok': False, 'error': '配件库存不足：%s 当前库存 %s，需 %s 个' % (new_part['name'], new_part['stock'], new_qty)})
        else:
            conn.execute("UPDATE pc_part_stock SET stock=stock-?,updated_at=? WHERE id=?", (new_qty, now_str(), new_part['id']))
            warn = (warn + '；' if warn else '') + '已自动扣减库存：%s ×%s（%s）' % (new_part['name'], new_qty, new_part['storage_location'])
    conn.execute(
        "UPDATE pc_replace SET user_name=?,department=?,location=?,part_name=?,replaced=?,returned=?,monitor_sn=?,part_qty=?,storage=?,remark=?,record_time=?,updated_at=? WHERE id=?",
        (str(d.get('user_name', '')).strip(), str(d.get('department', '')).strip(),
         str(d.get('location', '')).strip(), new_name,
         str(d.get('replaced', '')).strip() or '是', str(d.get('returned', '')).strip() or '否',
         str(d.get('monitor_sn', '')).strip(), str(d.get('part_qty', '')).strip(),
         new_storage, str(d.get('remark', '')).strip(),
         d.get('record_time') or now_str(), now_str(), rid))
    conn.commit()
    conn.close()
    log(session['username'], '修改配件更换记录', str(rid))
    return jsonify({'ok': True, 'warn': warn})


@app.route('/api/pc-replace/<int:rid>', methods=['DELETE'])
@require_login
def delete_pc_replace(rid):
    conn = get_conn()
    row = conn.execute("SELECT * FROM pc_replace WHERE id=?", (rid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '更换记录不存在'}), 404
    qty = parse_part_qty(row['part_qty'])
    warn = ''
    if qty > 0:
        part = find_part_stock(conn, row['part_name'], row['storage'])
        if part:
            conn.execute("UPDATE pc_part_stock SET stock=stock+?,updated_at=? WHERE id=?", (qty, now_str(), part['id']))
            warn = '已回补库存：%s ×%s' % (part['name'], qty)
    conn.execute("DELETE FROM pc_replace WHERE id=?", (rid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除配件更换记录', str(rid))
    return jsonify({'ok': True, 'warn': warn})


@app.route('/api/pc-replace/import', methods=['POST'])
@require_login
def import_pc_replace():
    f = request.files.get('file')
    if not f or not f.filename:
        return jsonify({'ok': False, 'error': '请选择要导入的 Excel 文件'})
    if not (f.filename.endswith('.xlsx') or f.filename.endswith('.xls')):
        return jsonify({'ok': False, 'error': '仅支持 .xlsx / .xls 文件'})
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(f.read()), data_only=True)
    except Exception as e:
        return jsonify({'ok': False, 'error': '文件解析失败：%s' % e})
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    conn = get_conn()
    prepared = []
    errors = []
    for idx, row in enumerate(rows, start=2):
        if row is None or all(v is None or str(v).strip() == '' for v in row):
            continue
        def gv(i):
            return str(row[i] or '').strip() if len(row) > i and row[i] is not None else ''
        name = gv(1)
        if not name:
            errors.append('第%d行：姓名必填' % idx)
            continue
        replaced = gv(5) or '是'
        if replaced not in ('是', '否'):
            errors.append('第%d行：是否更换必须为 是/否' % idx)
            continue
        returned = gv(6) or '否'
        if returned not in ('是', '否'):
            errors.append('第%d行：是否返厂必须为 是/否' % idx)
            continue
        prepared.append((name, gv(2), gv(3), gv(4), replaced, returned, gv(7), gv(9), gv(10), gv(8)))
    if errors:
        conn.close()
        return jsonify({'ok': False, 'error': '存在 %d 行校验失败，已整体取消导入' % len(errors), 'errors': errors[:20]})
    if not prepared:
        conn.close()
        return jsonify({'ok': False, 'error': '模板中没有可导入的数据行'})
    try:
        now = now_str()
        for p in prepared:
            conn.execute(
                "INSERT INTO pc_replace(user_name,department,location,part_name,replaced,returned,monitor_sn,part_qty,storage,remark,record_time,operator,created_at,updated_at) "
                "VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (p[0], p[1], p[2], p[3], p[4], p[5], p[6], p[7], p[8], p[9], now, session['username'], now, now))
        conn.commit()
        conn.close()
        log(session['username'], '模板导入配件更换记录', '共%d条' % len(prepared))
        return jsonify({'ok': True, 'added': len(prepared)})
    except Exception as e:
        conn.rollback()
        conn.close()
        return jsonify({'ok': False, 'error': '导入执行失败：%s' % e})


@app.route('/api/pc-replace/export', methods=['GET'])
@require_login
def export_pc_replace():
    q = request.args.get('q', '').strip()
    conn = get_conn()
    ids = _ids_arg()
    if ids:
        ph = ','.join('?' * len(ids))
        sql = "SELECT * FROM pc_replace WHERE id IN (%s)" % ph
        args = list(ids)
    else:
        sql = "SELECT * FROM pc_replace WHERE 1=1"
        args = []
    if q:
        sql += " AND (user_name LIKE ? OR department LIKE ? OR location LIKE ? OR part_name LIKE ? OR monitor_sn LIKE ? OR remark LIKE ?)"
        like = f"%{q}%"
        args += [like] * 6
    sql += " ORDER BY id DESC"
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    import openpyxl, io as _io
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '更换记录'
    headers = ['姓名', '部门', '位置', '配件', '是否更换', '是否返厂', '显示器序列号', '备注', '备件数量', '存放位置', '记录时间', '录入人']
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='4F46E5')
        c.alignment = Alignment(horizontal='center')
    for r in rows:
        ws.append([r['user_name'], r['department'], r['location'], r['part_name'], r['replaced'], r['returned'],
                   r['monitor_sn'], r['remark'], r['part_qty'], r['storage'], r['record_time'], r['operator']])
    from openpyxl.utils import get_column_letter
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(h) * 2 + 4)
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='行政办公电脑配件更换记录.xlsx')


# ---------------- 配件更换：显示器统计 ----------------
@app.route('/api/pc-monitors', methods=['GET'])
@require_login
def list_pc_monitors():
    q = request.args.get('q', '').strip()
    sql = "SELECT * FROM pc_monitors WHERE 1=1"
    args = []
    if q:
        sql += " AND (area LIKE ? OR department LIKE ? OR location LIKE ?)"
        like = f"%{q}%"
        args += [like] * 3
    sql += " ORDER BY id ASC"
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows]})


@app.route('/api/pc-monitors/<int:rid>', methods=['GET'])
@require_login
def get_pc_monitor(rid):
    conn = get_conn()
    r = conn.execute("SELECT * FROM pc_monitors WHERE id=?", (rid,)).fetchone()
    conn.close()
    if not r:
        return jsonify({'ok': False, 'error': '记录不存在'}), 404
    return jsonify({'ok': True, 'data': dict(r)})


@app.route('/api/pc-monitors', methods=['POST'])
@require_login
def add_pc_monitor():
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    try:
        cnt = int(float(d.get('monitor_count') or 0))
    except (ValueError, TypeError):
        cnt = 0
    cur = conn.execute(
        "INSERT INTO pc_monitors(area,department,location,monitor_count,created_at,updated_at) VALUES(?,?,?,?,?,?)",
        (str(d.get('area', '')).strip(), str(d.get('department', '')).strip(),
         str(d.get('location', '')).strip(), cnt, now_str(), now_str()))
    conn.commit()
    conn.close()
    log(session['username'], '新增显示器统计', str(cur.lastrowid))
    return jsonify({'ok': True, 'id': cur.lastrowid})


@app.route('/api/pc-monitors/<int:rid>', methods=['PUT'])
@require_login
def update_pc_monitor(rid):
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    try:
        cnt = int(float(d.get('monitor_count') or 0))
    except (ValueError, TypeError):
        cnt = 0
    conn.execute(
        "UPDATE pc_monitors SET area=?,department=?,location=?,monitor_count=?,updated_at=? WHERE id=?",
        (str(d.get('area', '')).strip(), str(d.get('department', '')).strip(),
         str(d.get('location', '')).strip(), cnt, now_str(), rid))
    conn.commit()
    conn.close()
    log(session['username'], '修改显示器统计', str(rid))
    return jsonify({'ok': True})


@app.route('/api/pc-monitors/<int:rid>', methods=['DELETE'])
@require_login
def delete_pc_monitor(rid):
    conn = get_conn()
    conn.execute("DELETE FROM pc_monitors WHERE id=?", (rid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除显示器统计', str(rid))
    return jsonify({'ok': True})


@app.route('/api/pc-monitors/import', methods=['POST'])
@require_login
def import_pc_monitors():
    f = request.files.get('file')
    if not f or not f.filename:
        return jsonify({'ok': False, 'error': '请选择要导入的 Excel 文件'})
    if not (f.filename.endswith('.xlsx') or f.filename.endswith('.xls')):
        return jsonify({'ok': False, 'error': '仅支持 .xlsx / .xls 文件'})
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(f.read()), data_only=True)
    except Exception as e:
        return jsonify({'ok': False, 'error': '文件解析失败：%s' % e})
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    conn = get_conn()
    prepared = []
    errors = []
    last_area, last_dept = '', ''
    for idx, row in enumerate(rows, start=2):
        if row is None or all(v is None or str(v).strip() == '' for v in row):
            continue
        def gv(i):
            return str(row[i] or '').strip() if len(row) > i and row[i] is not None else ''
        area = gv(0) or last_area
        dept = gv(1) or last_dept
        loc = gv(2)
        if not loc:
            errors.append('第%d行：位置必填' % idx)
            continue
        try:
            cnt = int(float(gv(3) or 1))
        except (ValueError, TypeError):
            cnt = 1
        last_area, last_dept = area, dept
        prepared.append((area, dept, loc, cnt))
    if errors:
        conn.close()
        return jsonify({'ok': False, 'error': '存在 %d 行校验失败，已整体取消导入' % len(errors), 'errors': errors[:20]})
    if not prepared:
        conn.close()
        return jsonify({'ok': False, 'error': '模板中没有可导入的数据行'})
    try:
        now = now_str()
        for p in prepared:
            conn.execute(
                "INSERT INTO pc_monitors(area,department,location,monitor_count,created_at,updated_at) VALUES(?,?,?,?,?,?)",
                (p[0], p[1], p[2], p[3], now, now))
        conn.commit()
        conn.close()
        log(session['username'], '模板导入显示器统计', '共%d条' % len(prepared))
        return jsonify({'ok': True, 'added': len(prepared)})
    except Exception as e:
        conn.rollback()
        conn.close()
        return jsonify({'ok': False, 'error': '导入执行失败：%s' % e})


@app.route('/api/pc-monitors/export', methods=['GET'])
@require_login
def export_pc_monitors():
    q = request.args.get('q', '').strip()
    conn = get_conn()
    ids = _ids_arg()
    if ids:
        ph = ','.join('?' * len(ids))
        sql = "SELECT * FROM pc_monitors WHERE id IN (%s)" % ph
        args = list(ids)
    else:
        sql = "SELECT * FROM pc_monitors WHERE 1=1"
        args = []
    if q:
        sql += " AND (area LIKE ? OR department LIKE ? OR location LIKE ?)"
        like = f"%{q}%"
        args += [like] * 3
    sql += " ORDER BY id ASC"
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    import openpyxl, io as _io
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '显示器统计'
    headers = ['区域', '部门', '位置', '显示器数量']
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='4F46E5')
        c.alignment = Alignment(horizontal='center')
    for r in rows:
        ws.append([r['area'], r['department'], r['location'], r['monitor_count']])
    from openpyxl.utils import get_column_letter
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(h) * 2 + 4)
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='显示器统计.xlsx')


# ---------------- 配件更换：大厅电脑信息 ----------------
@app.route('/api/pc-hall-computers', methods=['GET'])
@require_login
def list_pc_hall_computers():
    q = request.args.get('q', '').strip()
    sql = "SELECT * FROM pc_hall_computers WHERE 1=1"
    args = []
    if q:
        sql += " AND (area LIKE ? OR department LIKE ? OR location LIKE ? OR ip LIKE ? OR user_name LIKE ?)"
        like = f"%{q}%"
        args += [like] * 5
    sql += " ORDER BY id ASC"
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows]})


@app.route('/api/pc-hall-computers/<int:rid>', methods=['GET'])
@require_login
def get_pc_hall_computer(rid):
    conn = get_conn()
    r = conn.execute("SELECT * FROM pc_hall_computers WHERE id=?", (rid,)).fetchone()
    conn.close()
    if not r:
        return jsonify({'ok': False, 'error': '记录不存在'}), 404
    return jsonify({'ok': True, 'data': dict(r)})


@app.route('/api/pc-hall-computers', methods=['POST'])
@require_login
def add_pc_hall_computer():
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    cur = conn.execute(
        "INSERT INTO pc_hall_computers(area,department,location,ip,user_name,created_at,updated_at) VALUES(?,?,?,?,?,?,?)",
        (str(d.get('area', '')).strip(), str(d.get('department', '')).strip(),
         str(d.get('location', '')).strip(), str(d.get('ip', '')).strip(),
         str(d.get('user_name', '')).strip(), now_str(), now_str()))
    conn.commit()
    conn.close()
    log(session['username'], '新增大厅电脑信息', str(cur.lastrowid))
    return jsonify({'ok': True, 'id': cur.lastrowid})


@app.route('/api/pc-hall-computers/<int:rid>', methods=['PUT'])
@require_login
def update_pc_hall_computer(rid):
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    conn.execute(
        "UPDATE pc_hall_computers SET area=?,department=?,location=?,ip=?,user_name=?,updated_at=? WHERE id=?",
        (str(d.get('area', '')).strip(), str(d.get('department', '')).strip(),
         str(d.get('location', '')).strip(), str(d.get('ip', '')).strip(),
         str(d.get('user_name', '')).strip(), now_str(), rid))
    conn.commit()
    conn.close()
    log(session['username'], '修改大厅电脑信息', str(rid))
    return jsonify({'ok': True})


@app.route('/api/pc-hall-computers/<int:rid>', methods=['DELETE'])
@require_login
def delete_pc_hall_computer(rid):
    conn = get_conn()
    conn.execute("DELETE FROM pc_hall_computers WHERE id=?", (rid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除大厅电脑信息', str(rid))
    return jsonify({'ok': True})


@app.route('/api/pc-hall-computers/import', methods=['POST'])
@require_login
def import_pc_hall_computers():
    f = request.files.get('file')
    if not f or not f.filename:
        return jsonify({'ok': False, 'error': '请选择要导入的 Excel 文件'})
    if not (f.filename.endswith('.xlsx') or f.filename.endswith('.xls')):
        return jsonify({'ok': False, 'error': '仅支持 .xlsx / .xls 文件'})
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(f.read()), data_only=True)
    except Exception as e:
        return jsonify({'ok': False, 'error': '文件解析失败：%s' % e})
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    conn = get_conn()
    prepared = []
    errors = []
    last_area, last_dept = '', ''
    for idx, row in enumerate(rows, start=2):
        if row is None or all(v is None or str(v).strip() == '' for v in row):
            continue
        def gv(i):
            return str(row[i] or '').strip() if len(row) > i and row[i] is not None else ''
        area = gv(0) or last_area
        dept = gv(1) or last_dept
        loc = gv(2)
        if not loc:
            errors.append('第%d行：位置必填' % idx)
            continue
        last_area, last_dept = area, dept
        prepared.append((area, dept, loc, gv(3), gv(4)))
    if errors:
        conn.close()
        return jsonify({'ok': False, 'error': '存在 %d 行校验失败，已整体取消导入' % len(errors), 'errors': errors[:20]})
    if not prepared:
        conn.close()
        return jsonify({'ok': False, 'error': '模板中没有可导入的数据行'})
    try:
        now = now_str()
        for p in prepared:
            conn.execute(
                "INSERT INTO pc_hall_computers(area,department,location,ip,user_name,created_at,updated_at) VALUES(?,?,?,?,?,?,?)",
                (p[0], p[1], p[2], p[3], p[4], now, now))
        conn.commit()
        conn.close()
        log(session['username'], '模板导入大厅电脑信息', '共%d条' % len(prepared))
        return jsonify({'ok': True, 'added': len(prepared)})
    except Exception as e:
        conn.rollback()
        conn.close()
        return jsonify({'ok': False, 'error': '导入执行失败：%s' % e})


@app.route('/api/pc-hall-computers/export', methods=['GET'])
@require_login
def export_pc_hall_computers():
    q = request.args.get('q', '').strip()
    conn = get_conn()
    ids = _ids_arg()
    if ids:
        ph = ','.join('?' * len(ids))
        sql = "SELECT * FROM pc_hall_computers WHERE id IN (%s)" % ph
        args = list(ids)
    else:
        sql = "SELECT * FROM pc_hall_computers WHERE 1=1"
        args = []
    if q:
        sql += " AND (area LIKE ? OR department LIKE ? OR location LIKE ? OR ip LIKE ? OR user_name LIKE ?)"
        like = f"%{q}%"
        args += [like] * 5
    sql += " ORDER BY id ASC"
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    import openpyxl, io as _io
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '大厅电脑信息'
    headers = ['区域', '部门', '位置', 'IP地址', '使用人']
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='4F46E5')
        c.alignment = Alignment(horizontal='center')
    for r in rows:
        ws.append([r['area'], r['department'], r['location'], r['ip'], r['user_name']])
    from openpyxl.utils import get_column_letter
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(h) * 2 + 4)
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='大厅电脑信息.xlsx')


# ---------------- 配件更换：配件库存 ----------------
@app.route('/api/pc-part-stock', methods=['GET'])
@require_login
def list_pc_part_stock():
    q = request.args.get('q', '').strip()
    sql = "SELECT * FROM pc_part_stock WHERE 1=1"
    args = []
    if q:
        sql += " AND (name LIKE ? OR spec LIKE ? OR storage_location LIKE ?)"
        like = '%' + q + '%'
        args += [like, like, like]
    sql += " ORDER BY id DESC"
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows]})


@app.route('/api/pc-part-stock', methods=['POST'])
@require_login
def add_pc_part_stock():
    d = request.get_json(force=True, silent=True) or {}
    name = str(d.get('name', '')).strip()
    if not name:
        return jsonify({'ok': False, 'error': '配件名称必填'})
    spec = str(d.get('spec', '')).strip()
    loc = str(d.get('storage_location', '')).strip()
    conn = get_conn()
    dup = conn.execute("SELECT id FROM pc_part_stock WHERE name=? AND spec=? AND storage_location=?",
                       (name, spec, loc)).fetchone()
    if dup:
        conn.close()
        return jsonify({'ok': False, 'error': '该配件（同规格/同存放位置）已存在，可直接登记入库'})
    cur = conn.execute(
        "INSERT INTO pc_part_stock(name,spec,unit,stock,low_threshold,storage_location,remark,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?)",
        (name, spec, str(d.get('unit', '') or '个').strip(), int(num(d.get('stock'))),
         int(num(d.get('low_threshold'))), loc, str(d.get('remark', '')).strip(), now_str(), now_str()))
    conn.commit()
    conn.close()
    log(session['username'], '新增配件库存', name)
    return jsonify({'ok': True, 'id': cur.lastrowid})


@app.route('/api/pc-part-stock/<int:pid>', methods=['PUT'])
@require_login
def update_pc_part_stock(pid):
    d = request.get_json(force=True, silent=True) or {}
    conn = get_conn()
    cur = conn.execute(
        "UPDATE pc_part_stock SET name=?,spec=?,unit=?,low_threshold=?,storage_location=?,remark=?,updated_at=? WHERE id=?",
        (str(d.get('name', '')).strip(), str(d.get('spec', '')).strip(), str(d.get('unit', '') or '个').strip(),
         int(num(d.get('low_threshold'))), str(d.get('storage_location', '')).strip(),
         str(d.get('remark', '')).strip(), now_str(), pid))
    conn.commit()
    conn.close()
    if cur.rowcount == 0:
        return jsonify({'ok': False, 'error': '配件不存在'}), 404
    log(session['username'], '编辑配件库存', str(pid))
    return jsonify({'ok': True})


@app.route('/api/pc-part-stock/<int:pid>', methods=['DELETE'])
@require_login
def delete_pc_part_stock(pid):
    conn = get_conn()
    row = conn.execute("SELECT * FROM pc_part_stock WHERE id=?", (pid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '配件不存在'}), 404
    if int(num(row['stock'])) != 0:
        conn.close()
        return jsonify({'ok': False, 'error': '当前库存 %s，请先出库/调整清零后再删除' % row['stock']})
    conn.execute("DELETE FROM pc_part_stock WHERE id=?", (pid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除配件库存', row['name'])
    return jsonify({'ok': True})


@app.route('/api/pc-part-stock/import', methods=['POST'])
@require_login
def import_pc_part_stock():
    f = request.files.get('file')
    if not f or not f.filename:
        return jsonify({'ok': False, 'error': '请选择要导入的 Excel 文件'})
    if not (f.filename.endswith('.xlsx') or f.filename.endswith('.xls')):
        return jsonify({'ok': False, 'error': '仅支持 .xlsx / .xls 文件'})
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(f.read()), data_only=True)
    except Exception as e:
        return jsonify({'ok': False, 'error': '文件解析失败：%s' % e})
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    conn = get_conn()
    prepared = []
    errors = []
    for idx, row in enumerate(rows, start=2):
        if row is None or all(v is None or str(v).strip() == '' for v in row):
            continue
        def gv(i):
            return str(row[i] or '').strip() if len(row) > i and row[i] is not None else ''
        name = gv(0)
        if not name:
            errors.append('第%d行：配件名称必填' % idx)
            continue
        prepared.append((name, gv(1), gv(2) or '个', int(num(gv(3))), int(num(gv(4))), gv(5), gv(6)))
    if errors:
        conn.close()
        return jsonify({'ok': False, 'error': '存在 %d 行校验失败，已整体取消导入' % len(errors), 'errors': errors[:20]})
    if not prepared:
        conn.close()
        return jsonify({'ok': False, 'error': '模板中没有可导入的数据行'})
    try:
        now = now_str()
        for p in prepared:
            dup = conn.execute("SELECT id FROM pc_part_stock WHERE name=? AND spec=? AND storage_location=?",
                               (p[0], p[1], p[5])).fetchone()
            if dup:
                conn.execute("UPDATE pc_part_stock SET stock=?,low_threshold=?,unit=?,remark=?,updated_at=? WHERE id=?",
                             (p[3], p[4], p[2], p[6], now, dup[0]))
            else:
                conn.execute(
                    "INSERT INTO pc_part_stock(name,spec,unit,stock,low_threshold,storage_location,remark,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?)",
                    (p[0], p[1], p[2], p[3], p[4], p[5], p[6], now, now))
        conn.commit()
        conn.close()
        log(session['username'], '模板导入配件库存', '共%d条' % len(prepared))
        return jsonify({'ok': True, 'added': len(prepared)})
    except Exception as e:
        conn.rollback()
        conn.close()
        return jsonify({'ok': False, 'error': '导入执行失败：%s' % e})


@app.route('/api/pc-part-stock/export', methods=['GET'])
@require_login
def export_pc_part_stock():
    q = request.args.get('q', '').strip()
    conn = get_conn()
    ids = _ids_arg()
    if ids:
        ph = ','.join('?' * len(ids))
        sql = "SELECT * FROM pc_part_stock WHERE id IN (%s)" % ph
        args = list(ids)
    else:
        sql = "SELECT * FROM pc_part_stock WHERE 1=1"
        args = []
    if q:
        sql += " AND (name LIKE ? OR spec LIKE ? OR storage_location LIKE ?)"
        like = f"%{q}%"
        args += [like, like, like]
    sql += " ORDER BY id DESC"
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    import openpyxl, io as _io
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '配件库存'
    headers = ['配件名称', '规格', '单位', '当前库存', '低库存阈值', '存放位置', '备注']
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='4F46E5')
        c.alignment = Alignment(horizontal='center')
    for r in rows:
        ws.append([r['name'], r['spec'], r['unit'], r['stock'], r['low_threshold'], r['storage_location'], r['remark']])
    from openpyxl.utils import get_column_letter
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(h) * 2 + 4)
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='配件库存.xlsx')


# ---------------- 配件更换：入库记录（登记入库自动更新库存） ----------------
@app.route('/api/pc-part-stock-in', methods=['GET'])
@require_login
def list_pc_part_stock_in():
    q = request.args.get('q', '').strip()
    sql = "SELECT * FROM pc_part_stock_in WHERE 1=1"
    args = []
    if q:
        sql += " AND (part_name LIKE ? OR operator LIKE ? OR storage_location LIKE ? OR remark LIKE ?)"
        like = '%' + q + '%'
        args += [like, like, like, like]
    sql += " ORDER BY record_time DESC, id DESC"
    conn = get_conn()
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    return jsonify({'ok': True, 'data': [dict(r) for r in rows]})


@app.route('/api/pc-part-stock-in', methods=['POST'])
@require_login
def add_pc_part_stock_in():
    """配件入库：写入入库记录并自动增加配件库存（按名称+规格+存放位置匹配，不存在自动创建）"""
    d = request.get_json(force=True, silent=True) or {}
    name = str(d.get('part_name', '')).strip()
    if not name:
        return jsonify({'ok': False, 'error': '配件名称必填'})
    qty = int(num(d.get('qty')))
    if qty <= 0:
        return jsonify({'ok': False, 'error': '入库数量必须大于 0'})
    spec = str(d.get('spec', '')).strip()
    loc = str(d.get('storage_location', '')).strip()
    conn = get_conn()
    part = conn.execute("SELECT * FROM pc_part_stock WHERE name=? AND spec=? AND storage_location=? LIMIT 1",
                        (name, spec, loc)).fetchone()
    if not part:
        cur = conn.execute(
            "INSERT INTO pc_part_stock(name,spec,unit,stock,low_threshold,storage_location,remark,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?)",
            (name, spec, str(d.get('unit', '') or '个').strip(), 0, 0, loc, '入库自动创建', now_str(), now_str()))
        part = {'id': cur.lastrowid, 'stock': 0}
    after = int(num(part['stock'])) + qty
    conn.execute("UPDATE pc_part_stock SET stock=?,updated_at=? WHERE id=?", (after, now_str(), part['id']))
    try:
        record_time = d.get('record_time') or now_str()[:10]
        datetime.datetime.strptime(record_time[:10], '%Y-%m-%d')
    except ValueError:
        record_time = datetime.datetime.now().strftime('%Y-%m-%d')
    cur = conn.execute(
        "INSERT INTO pc_part_stock_in(part_name,spec,qty,after_stock,storage_location,operator,record_time,remark,created_at) VALUES(?,?,?,?,?,?,?,?,?)",
        (name, spec, qty, after, loc, session.get('real_name') or session['username'],
         record_time, str(d.get('remark', '')).strip(), now_str()))
    conn.commit()
    conn.close()
    log(session['username'], '配件入库', '%s qty=%s' % (name, qty))
    return jsonify({'ok': True, 'id': cur.lastrowid, 'after_stock': after})


@app.route('/api/pc-part-stock-in/<int:rid>', methods=['DELETE'])
@require_login
def delete_pc_part_stock_in(rid):
    """删除入库记录并反向扣减库存（不足拒绝）"""
    conn = get_conn()
    row = conn.execute("SELECT * FROM pc_part_stock_in WHERE id=?", (rid,)).fetchone()
    if not row:
        conn.close()
        return jsonify({'ok': False, 'error': '入库记录不存在'}), 404
    qty = int(num(row['qty']))
    part = conn.execute("SELECT * FROM pc_part_stock WHERE name=? AND spec=? AND storage_location=? LIMIT 1",
                        (row['part_name'], row['spec'] or '', row['storage_location'] or '')).fetchone()
    if not part:
        conn.execute("DELETE FROM pc_part_stock_in WHERE id=?", (rid,))
        conn.commit()
        conn.close()
        log(session['username'], '删除配件入库记录', '%s qty=%s' % (rid, qty))
        return jsonify({'ok': True, 'warn': '对应库存条目已不存在，仅删除记录'})
    if int(num(part['stock'])) < qty:
        conn.close()
        return jsonify({'ok': False, 'error': '无法删除：当前库存 %s 不足 %s，请先调整库存' % (part['stock'], qty)})
    conn.execute("UPDATE pc_part_stock SET stock=stock-?,updated_at=? WHERE id=?", (qty, now_str(), part['id']))
    conn.execute("DELETE FROM pc_part_stock_in WHERE id=?", (rid,))
    conn.commit()
    conn.close()
    log(session['username'], '删除配件入库记录', '%s qty=%s' % (rid, qty))
    return jsonify({'ok': True})


@app.route('/api/pc-part-stock-in/import', methods=['POST'])
@require_login
def import_pc_part_stock_in():
    """批量入库：按模板（配件名称/规格/数量/存放位置/经手人/入库时间/备注）逐行入库并联动库存"""
    f = request.files.get('file')
    if not f or not f.filename:
        return jsonify({'ok': False, 'error': '请选择要导入的 Excel 文件'})
    if not (f.filename.endswith('.xlsx') or f.filename.endswith('.xls')):
        return jsonify({'ok': False, 'error': '仅支持 .xlsx / .xls 文件'})
    import openpyxl
    try:
        wb = openpyxl.load_workbook(io.BytesIO(f.read()), data_only=True)
    except Exception as e:
        return jsonify({'ok': False, 'error': '文件解析失败：%s' % e})
    ws = wb.active
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    conn = get_conn()
    prepared = []
    errors = []
    for idx, row in enumerate(rows, start=2):
        if row is None or all(v is None or str(v).strip() == '' for v in row):
            continue
        def gv(i):
            return str(row[i] or '').strip() if len(row) > i and row[i] is not None else ''
        name = gv(0)
        if not name:
            errors.append('第%d行：配件名称必填' % idx)
            continue
        qty = int(num(gv(2)))
        if qty <= 0:
            errors.append('第%d行：入库数量必须大于 0' % idx)
            continue
        prepared.append((name, gv(1), qty, gv(3), gv(4), gv(5), gv(6)))
    if errors:
        conn.close()
        return jsonify({'ok': False, 'error': '存在 %d 行校验失败，已整体取消导入' % len(errors), 'errors': errors[:20]})
    if not prepared:
        conn.close()
        return jsonify({'ok': False, 'error': '模板中没有可导入的数据行'})
    try:
        now = now_str()
        added = 0
        for p in prepared:
            name, spec, qty, loc, op, rt, rm = p
            part = conn.execute("SELECT * FROM pc_part_stock WHERE name=? AND spec=? AND storage_location=? LIMIT 1",
                                (name, spec, loc)).fetchone()
            if not part:
                cur = conn.execute(
                    "INSERT INTO pc_part_stock(name,spec,unit,stock,low_threshold,storage_location,remark,created_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?)",
                    (name, spec, '个', 0, 0, loc, '入库自动创建', now, now))
                part = {'id': cur.lastrowid, 'stock': 0}
            after = int(num(part['stock'])) + qty
            conn.execute("UPDATE pc_part_stock SET stock=?,updated_at=? WHERE id=?", (after, now, part['id']))
            conn.execute(
                "INSERT INTO pc_part_stock_in(part_name,spec,qty,after_stock,storage_location,operator,record_time,remark,created_at) VALUES(?,?,?,?,?,?,?,?,?)",
                (name, spec, qty, after, loc, op or (session.get('real_name') or session['username']),
                 rt or now[:10], rm, now))
            added += 1
        conn.commit()
        conn.close()
        log(session['username'], '模板导入配件入库', '共%d条' % added)
        return jsonify({'ok': True, 'added': added})
    except Exception as e:
        conn.rollback()
        conn.close()
        return jsonify({'ok': False, 'error': '导入执行失败：%s' % e})


@app.route('/api/pc-part-stock-in/export', methods=['GET'])
@require_login
def export_pc_part_stock_in():
    q = request.args.get('q', '').strip()
    conn = get_conn()
    ids = _ids_arg()
    if ids:
        ph = ','.join('?' * len(ids))
        sql = "SELECT * FROM pc_part_stock_in WHERE id IN (%s)" % ph
        args = list(ids)
    else:
        sql = "SELECT * FROM pc_part_stock_in WHERE 1=1"
        args = []
    if q:
        sql += " AND (part_name LIKE ? OR operator LIKE ? OR storage_location LIKE ? OR remark LIKE ?)"
        like = f"%{q}%"
        args += [like, like, like, like]
    sql += " ORDER BY record_time DESC, id DESC"
    rows = conn.execute(sql, args).fetchall()
    conn.close()
    import openpyxl, io as _io
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '配件入库记录'
    headers = ['配件名称', '规格', '入库数量', '入库后库存', '存放位置', '经手人', '入库时间', '备注']
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor='4F46E5')
        c.alignment = Alignment(horizontal='center')
    for r in rows:
        ws.append([r['part_name'], r['spec'], r['qty'], r['after_stock'], r['storage_location'], r['operator'], r['record_time'], r['remark']])
    from openpyxl.utils import get_column_letter
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, len(h) * 2 + 4)
    buf = _io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name='配件入库记录.xlsx')


def ensure_admin():
    conn = get_conn()
    row = conn.execute("SELECT id FROM users WHERE username=?", (DEFAULT_ADMIN[0],)).fetchone()
    if not row:
        conn.execute("INSERT INTO users(username,password_hash,real_name,role,active,created_at) VALUES(?,?,?,?,?,?)",
                     (DEFAULT_ADMIN[0], generate_password_hash(DEFAULT_ADMIN[1]), '系统管理员', 'admin', 1, now_str()))
        conn.commit()
        print(f"已创建默认管理员账号: {DEFAULT_ADMIN[0]} / {DEFAULT_ADMIN[1]}")
    conn.close()


def main():
    from db import init_db
    init_db()
    ensure_admin()
    # access.log 轮转：超过 5MB 归档为 access_日期.log，避免日志无限膨胀
    try:
        if os.path.exists(ACCESS_LOG) and os.path.getsize(ACCESS_LOG) > 5 * 1024 * 1024:
            ts = datetime.datetime.now().strftime('%Y%m%d_%H%M%S')
            os.rename(ACCESS_LOG, os.path.join(DATA_DIR, f'access_{ts}.log'))
    except Exception:
        pass
    th = threading.Thread(target=backup_loop, daemon=True)
    th.start()
    port = int(os.environ.get('TZ_PORT', '8090'))
    print(f"台账管理系统启动: http://127.0.0.1:{port}")
    app.run(host='0.0.0.0', port=port, threaded=True)


if __name__ == '__main__':
    main()
